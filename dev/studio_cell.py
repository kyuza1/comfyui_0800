#@title 🚀 ComfyUI Studio
#@markdown Aperte o **▶** ao lado do título. Vai aparecer um painel para escolher o **pack** (o ⭐ é o recomendado na T4 grátis) e depois é só clicar em **Iniciar**.
#@markdown
#@markdown **SALVAR_NO_DRIVE** guarda tudo o que você criar na pasta `ComfyUI_output` do seu Google Drive (recomendado: assim nada se perde quando o Colab desliga).
SALVAR_NO_DRIVE = True #@param {type:"boolean"}

import base64
import contextlib
import hmac
import html
import io
import json
import os
import random
import re
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from IPython.display import HTML, clear_output, display

# ------------------------------------------------------------------------------
# ESCOLHA DO PACK: painel com os packs prontos e um contador de memoria ao vivo
# ------------------------------------------------------------------------------

# Quanto cada peca pesa na T4 gratis (15 GB de placa, ~12,7 GB de RAM). Os numeros sao aproximados e de proposito
# um pouco pessimistas: o painel so deixa iniciar o que cabe com folga, para o usuario leigo nao travar a sessao.
# vram = memoria da placa quando a peca esta em uso · ram = memoria do sistema · disco = download (GB) · min = instalacao
ITENS = {
    "zimage":   {"nome": "Z Image Turbo (imagem)",            "vram": 10.0, "ram": 6.0, "disco": 20.7, "logo": "Tongyi-MAI", "min": 3},
    "realismo": {"nome": "LoRA UltraReal (fotos realistas)",   "vram": 0.3,  "ram": 0.2, "disco": 0.1,  "logo": "huggingface", "min": 0.2},
    "caricato": {"nome": "LoRA XUPLX (rostos caricatos)",      "vram": 0.3,  "ram": 0.2, "disco": 0.2,  "logo": "huggingface", "min": 0.2},
    "lote":     {"nome": "SDXL Lightning (muitas imagens)",    "vram": 7.0,  "ram": 4.0, "disco": 6.9,  "logo": "bytedance", "min": 1},
    "video":    {"nome": "Wan 2.2 5B + FastWan (vídeo)",       "vram": 12.0, "ram": 7.0, "disco": 18.9, "logo": "Wan-Video", "min": 3},
    "voz":      {"nome": "Voz pt-BR (Chatterbox)",             "vram": 3.0,  "ram": 2.0, "disco": 9.0,  "logo": "resemble-ai", "min": 4},
    "musica":   {"nome": "Música (ACE-Step)",                  "vram": 8.0,  "ram": 4.0, "disco": 7.7,  "logo": "ace-step", "min": 1},
    "efeitos":  {"nome": "Efeitos sonoros (Stable Audio)",     "vram": 6.0,  "ram": 3.0, "disco": 5.6,  "logo": "Stability-AI", "min": 1},
    "fundo":    {"nome": "Removedor de fundo (CPU)",           "vram": 0.0,  "ram": 0.7, "disco": 1.2,  "logo": "danielgatis", "min": 1},
    "webui":    {"nome": "Open WebUI (chat completo)",         "vram": 0.0,  "ram": 1.5, "disco": 7.0,  "logo": "open-webui", "min": 2},
    "comfy":    {"nome": "ComfyUI (base)",                     "vram": 0.0,  "ram": 2.5, "disco": 6.0,  "logo": "Comfy-Org", "min": 2},
}
# IA de conversa (Ollama): opcao do painel -> (modelo no Ollama, rotulo, vram, disco)
CHATS = {
    "qwen35_4b": ("huihui_ai/qwen3.5-abliterated:4B", "Qwen3.5 4B sem censura (leve, enxerga imagens)", 4.5, 3.3),
    "qwen35_9b": ("huihui_ai/qwen3.5-abliterated:9b", "Qwen3.5 9B sem censura (mais inteligente)", 8.0, 6.6),
    "qwen3_4b":  ("huihui_ai/qwen3-abliterated:4b", "Qwen3 4B sem censura (versão anterior, já testada)", 4.0, 2.5),
}
WHISPERS = {"tiny": (0.3, 0.1), "base": (0.5, 0.2), "small": (1.0, 0.5)}   # (ram, disco)
LIMITES = {"vram": 15.0, "ram": 12.0, "disco": 60.0}

# os packs: combinacoes prontas que cabem na T4. "estrela" = recomendado
PACKS = [
    {"id": "engracado", "nome": "Influenciador Engraçado", "logos": ["Tongyi-MAI", "huggingface", "resemble-ai"], "estrela": True,
     "descricao": "Personagens caricatos e esquisitos, no estilo das trends virais, com voz em português.",
     "imagem": "zimage", "lora": "caricato", "video": False, "voz": True, "musica": False, "efeitos": False, "fundo": True,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
    {"id": "hot", "nome": "Influenciador Hot (+18)", "logos": ["Tongyi-MAI", "huggingface", "resemble-ai"], "estrela": True,
     "descricao": "Fotos realistas e sensuais de uma influenciadora fictícia adulta, com voz em português.",
     "imagem": "zimage", "lora": "realismo", "video": False, "voz": True, "musica": False, "efeitos": False, "fundo": True,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
    {"id": "imagem", "nome": "Imagem", "logos": ["Tongyi-MAI", "danielgatis"], "estrela": True,
     "descricao": "Artes, ilustrações e fotos com o Z Image Turbo, o mais rápido de boa qualidade na T4.",
     "imagem": "zimage", "lora": "nenhuma", "video": False, "voz": False, "musica": False, "efeitos": False, "fundo": True,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
    {"id": "video", "nome": "Vídeo", "logos": ["Wan-Video"], "estrela": True,
     "descricao": "Vídeos curtos a partir de um texto ou animando uma foto sua (Wan 2.2).",
     "imagem": "nenhum", "lora": "nenhuma", "video": True, "voz": False, "musica": False, "efeitos": False, "fundo": False,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
    {"id": "audio", "nome": "Áudio e voz", "logos": ["resemble-ai", "ace-step"], "estrela": True,
     "descricao": "Narração em português, clonar uma voz e criar músicas.",
     "imagem": "nenhum", "lora": "nenhuma", "video": False, "voz": True, "musica": True, "efeitos": False, "fundo": False,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
    {"id": "chat", "nome": "Chat sem censura", "logos": ["QwenLM", "open-webui", "ollama"], "estrela": True,
     "descricao": "Conversa sem limite e sem recusas, num chat igual ao ChatGPT. Não instala nada de imagem.",
     "imagem": "nenhum", "lora": "nenhuma", "video": False, "voz": False, "musica": False, "efeitos": False, "fundo": False,
     "chat": "qwen35_9b", "webui": True, "whisper": "tiny"},
    {"id": "lote", "nome": "Muitas imagens", "logos": ["bytedance"], "estrela": False,
     "descricao": "Dezenas de imagens de uma vez, bem rápido (SDXL Lightning, 4 passos). Qualidade menor.",
     "imagem": "lote", "lora": "nenhuma", "video": False, "voz": False, "musica": False, "efeitos": False, "fundo": False,
     "chat": "desligado", "webui": False, "whisper": "desligado"},
]
ESCOLHA_SALVA = "/content/studio_escolha.json"   # a ultima escolha, para o painel ja abrir nela na proxima vez


def conta_memoria(e):
    """{vram, ram, disco, min} de uma escolha. A mesma conta e feita em JavaScript no painel (contaMemoria)."""
    usa_comfy = e["imagem"] != "nenhum" or e["video"] or e["musica"] or e["efeitos"]
    na_placa = ([e["imagem"]] if e["imagem"] != "nenhum" else []) + [k for k in ("video", "voz", "musica", "efeitos") if e[k]]
    if e.get("lora") in ITENS and e["imagem"] == "zimage":   # as LoRAs so existem para o Z Image
        na_placa.append(e["lora"])
    no_sistema = [k for k in ("fundo", "webui") if e[k]] + (["comfy"] if usa_comfy else [])
    total = {"vram": 0.0, "ram": 0.0, "disco": 0.0, "min": 0.0}
    for k in na_placa + no_sistema:
        for campo in total:
            total[campo] += ITENS[k][campo]
    if e["chat"] in CHATS:
        total["vram"] += CHATS[e["chat"]][2]
        total["ram"] += 1.0
        total["disco"] += CHATS[e["chat"]][3]
        total["min"] += 1.5
    if e["whisper"] in WHISPERS:
        total["ram"] += WHISPERS[e["whisper"]][0]
        total["disco"] += WHISPERS[e["whisper"]][1]
        total["min"] += 0.5
    return total


PAINEL_ESCOLHA_HTML = r'''<div id="studio-escolha">
<style>
#studio-escolha{font-family:Cascadia Mono,Cascadia Code,Consolas,DejaVu Sans Mono,Menlo,monospace;background:#fff;color:#1f2328;border:1px solid #57606a;
  max-width:920px;font-size:13px;line-height:1.6;box-shadow:0 6px 24px rgba(0,0,0,.18);margin:16px 0 40px}
#studio-escolha *{box-sizing:border-box}
#studio-escolha .titulo{background:#f3f4f6;display:flex;align-items:center;border-bottom:1px solid #57606a;height:30px;font:12px Segoe UI,system-ui,sans-serif}
#studio-escolha .corpo{padding:22px 26px 24px}
#studio-escolha img.logo{width:16px;height:16px;border-radius:4px;vertical-align:middle;border:1px solid #e5e7eb;background:#fff;margin-right:4px}
#studio-escolha h3{margin:0 0 2px;font-size:15px;color:#172dd7}
#studio-escolha .sub{color:#6e7781;margin:0 0 16px}
#studio-escolha .secao{color:#172dd7;font-weight:bold;margin:18px 0 6px}
#studio-escolha .packs{display:grid;grid-template-columns:repeat(auto-fill,minmax(270px,1fr));gap:10px}
#studio-escolha .pack{border:1px solid #d0d7de;padding:10px 12px;cursor:pointer}
#studio-escolha .pack:hover{border-color:#172dd7}
#studio-escolha .pack.sel{border:2px solid #172dd7;background:#f3f5ff;padding:9px 11px}
#studio-escolha .pack .nome{font-weight:bold}
#studio-escolha .pack .rec{color:#1a7f37;font-size:12px}
#studio-escolha .pack p{margin:4px 0 6px;color:#57606a;font-size:12px;font-family:Segoe UI,system-ui,sans-serif}
#studio-escolha details{margin-top:14px;border:1px solid #d0d7de;padding:8px 12px}
#studio-escolha summary{cursor:pointer;font-weight:bold}
#studio-escolha .opcoes{display:grid;grid-template-columns:repeat(auto-fill,minmax(400px,1fr));gap:2px 18px;margin-top:8px}
#studio-escolha label{display:flex;align-items:center;gap:8px;cursor:pointer}
#studio-escolha input[type=checkbox]{accent-color:#172dd7}
#studio-escolha select{font:inherit;border:1px solid #d0d7de;background:#fff;padding:1px 4px;max-width:300px}
#studio-escolha .barra{display:grid;grid-template-columns:150px auto 1fr;gap:12px;white-space:nowrap}
#studio-escolha .aviso{margin-top:14px;padding:10px 12px;border:2px solid}
#studio-escolha button{margin-top:16px;font:bold 14px Cascadia Mono,Consolas,monospace;background:#172dd7;color:#fff;border:0;padding:10px 26px;cursor:pointer}
#studio-escolha button:disabled{background:#d0d7de;color:#6e7781;cursor:not-allowed}
@media (max-width:640px){#studio-escolha .opcoes{grid-template-columns:1fr}#studio-escolha .barra{grid-template-columns:1fr}}
</style>
<div class="titulo"><span style="padding:0 10px"><img class="logo" src="https://github.com/Comfy-Org.png?size=64" alt=""></span>
<span style="flex:1">Prompt de Comando · comfyui-setup.exe · por __CRIADOR__</span></div>
<div class="corpo">
<h3>Escolha o pack</h3>
<p class="sub">Um pack de cada vez: a placa grátis do Colab é pequena. Os marcados com ★ são os recomendados para a T4.</p>
<div class="packs" id="se-packs"></div>
<details id="se-det"><summary>Personalizar (opcional)</summary><div class="opcoes" id="se-opcoes"></div></details>
<div class="secao">Memória</div>
<div id="se-barras"></div>
<div class="aviso" id="se-aviso"></div>
<button id="se-iniciar">Iniciar &gt;</button>
</div>
</div>
<script>
(function(){
const DADOS = __DADOS__;
let e = Object.assign({}, DADOS.inicial);
const $ = (id) => document.getElementById(id);
const logo = (org) => `<img class="logo" src="https://github.com/${org}.png?size=64" alt="">`;

function contaMemoria(e){   // mesma conta de conta_memoria() no Python
  const I = DADOS.itens, t = {vram:0, ram:0, disco:0, min:0};
  const usaComfy = e.imagem !== "nenhum" || e.video || e.musica || e.efeitos;
  const placa = (e.imagem !== "nenhum" ? [e.imagem] : []).concat(["video","voz","musica","efeitos"].filter(k => e[k]))
    .concat(I[e.lora] && e.imagem === "zimage" ? [e.lora] : []);
  const sistema = ["fundo","webui"].filter(k => e[k]).concat(usaComfy ? ["comfy"] : []);
  for (const k of placa.concat(sistema)) for (const c in t) t[c] += I[k][c];
  const ch = DADOS.chats[e.chat];
  if (ch) { t.vram += ch[2]; t.ram += 1; t.disco += ch[3]; t.min += 1.5; }
  const w = DADOS.whispers[e.whisper];
  if (w) { t.ram += w[0]; t.disco += w[1]; t.min += 0.5; }
  return t;
}

function montarPacks(){
  $("se-packs").innerHTML = "";
  for (const p of DADOS.packs) {
    const d = document.createElement("div");
    d.className = "pack" + (e.pack === p.id ? " sel" : "");
    d.innerHTML = `<div class="nome">${e.pack === p.id ? "&gt; " : ""}${p.nome}</div>` +
      `<div class="rec">${p.estrela ? "★ recomendado na T4" : "&nbsp;"}</div><p>${p.descricao}</p><div>${p.logos.map(logo).join("")}</div>`;
    d.onclick = () => { e = Object.assign({}, p, {pack: p.id, proteger: e.proteger, manutencao: e.manutencao}); montarPacks(); montarOpcoes(); atualizar(); };
    $("se-packs").appendChild(d);
  }
}

function caixa(chave, texto, org){
  return `<label><input type="checkbox" data-k="${chave}" ${e[chave] ? "checked" : ""}>${org ? logo(org) : ""}${texto}</label>`;
}
function lista(chave, texto, opcoes){
  return `<label>${texto} <select data-k="${chave}">` + opcoes.map(([v, t]) => `<option value="${v}" ${e[chave] === v ? "selected" : ""}>${t}</option>`).join("") + "</select></label>";
}
function montarOpcoes(){
  const I = DADOS.itens;
  $("se-opcoes").innerHTML = [
    lista("imagem", "Imagem:", [["nenhum","nenhuma"],["zimage",I.zimage.nome],["lote",I.lote.nome]]),
    lista("lora", "Estilo (LoRA, só no Z Image):", [["nenhuma","nenhum"],["realismo",I.realismo.nome],["caricato",I.caricato.nome]]),
    caixa("video", I.video.nome, I.video.logo),
    caixa("voz", I.voz.nome, I.voz.logo),
    caixa("musica", I.musica.nome, I.musica.logo),
    caixa("efeitos", I.efeitos.nome, I.efeitos.logo),
    caixa("fundo", I.fundo.nome, I.fundo.logo),
    lista("chat", "Chat:", [["desligado","desligado"]].concat(Object.entries(DADOS.chats).map(([k, v]) => [k, v[1]]))),
    caixa("webui", I.webui.nome + " (precisa do chat)", I.webui.logo),
    lista("whisper", "Microfone (Whisper):", [["desligado","desligado"],["tiny","tiny (leve)"],["base","base"],["small","small"]]),
    caixa("proteger", "Proteger o Studio e a API com chave (recomendado)"),
    caixa("manutencao", "Manutenção a distância (só se o criador pedir)"),
  ].join("");
  $("se-opcoes").querySelectorAll("[data-k]").forEach(el => el.onchange = () => {
    e[el.dataset.k] = el.type === "checkbox" ? el.checked : el.value;
    e.pack = "personalizado"; montarPacks(); atualizar();
  });
}

function barra(rotulo, valor, limite, unidade){
  const n = 30, cheio = Math.min(n, Math.round(valor / limite * n));
  const cor = valor > limite ? "#cf222e" : valor / limite > 0.85 ? "#9a6700" : "#172dd7";
  return `<div class="barra"><span>${rotulo}</span><span><span style="color:${cor}">${"█".repeat(cheio)}</span><span style="color:#d0d7de">${"░".repeat(n - cheio)}</span></span>` +
         `<span style="color:${valor > limite ? "#cf222e" : "#6e7781"}">${valor.toFixed(1)} / ${limite} ${unidade}</span></div>`;
}

function atualizar(){
  const t = contaMemoria(e), L = DADOS.limites, problemas = [];
  $("se-barras").innerHTML = barra("Placa (VRAM)", t.vram, L.vram, "GB") + barra("Memória (RAM)", t.ram, L.ram, "GB") +
    barra("Download", t.disco, L.disco, "GB") +
    `<div class="barra"><span>Instalação</span><span style="color:#6e7781">cerca de ${Math.max(1, Math.round(t.min + 2))} min na primeira vez</span><span></span></div>`;
  if (t.vram > L.vram) {
    const I = DADOS.itens, ch = DADOS.chats[e.chat];
    const pecas = (e.imagem !== "nenhum" ? [I[e.imagem]] : []).concat(["video","voz","musica","efeitos"].filter(k => e[k]).map(k => I[k]))
      .map(i => `${i.nome.split(" (")[0]} ${i.vram} GB`).concat(ch ? [`chat ${ch[2]} GB`] : []);
    problemas.push(`não cabe na placa (15 GB): ${pecas.join(" + ")}. Tire uma dessas coisas.`);
  }
  if (t.ram > L.ram) problemas.push("a memória do sistema não aguenta: tire o Open WebUI, o microfone ou um dos modelos.");
  if (t.disco > L.disco) problemas.push("download grande demais para o disco do Colab.");
  const nada = e.imagem === "nenhum" && !e.video && !e.voz && !e.musica && !e.efeitos && !e.fundo && e.chat === "desligado";
  if (nada) problemas.push("escolha pelo menos uma coisa.");
  if (e.webui && e.chat === "desligado") problemas.push("o Open WebUI precisa de um modelo de chat.");
  const av = $("se-aviso");
  if (problemas.length) {
    av.style.cssText = "border-color:#cf222e;background:#fff5f5;color:#cf222e";
    av.innerHTML = "[x] " + problemas.join("<br>[x] ");
  } else {
    av.style.cssText = "border-color:#1a7f37;background:#f0fff4;color:#1a7f37";
    av.innerHTML = "[✓] Cabe na T4. " + (e.pack === "personalizado" ? "Combinação personalizada." : "Pack " + DADOS.packs.find(p => p.id === e.pack).nome + ".");
  }
  $("se-iniciar").disabled = problemas.length > 0;
}

window.studioPromessa = new Promise(resolver => {
  $("se-iniciar").onclick = () => {
    $("se-iniciar").disabled = true; $("se-iniciar").textContent = "Iniciando...";
    resolver(JSON.stringify(e));
  };
});
montarPacks(); montarOpcoes(); atualizar();
})();
</script>'''


def escolher_pack():
    """Mostra o painel e espera o clique em Iniciar. Fora do Colab (nos testes) usa ESCOLHA_TESTE ou o pack Imagem."""
    padrao = dict(next(p for p in PACKS if p["id"] == "imagem"), pack="imagem", proteger=True, manutencao=False)
    try:
        with open(ESCOLHA_SALVA, encoding="utf-8") as f:
            padrao.update(json.load(f))
    except (OSError, ValueError):
        pass
    try:
        from google.colab.output import eval_js
    except ImportError:
        return dict(padrao, **globals().get("ESCOLHA_TESTE", {}))
    dados = {"itens": ITENS, "chats": CHATS, "whispers": WHISPERS, "limites": LIMITES, "packs": PACKS, "inicial": padrao}
    clear_output()
    display(HTML(PAINEL_ESCOLHA_HTML.replace("__DADOS__", json.dumps(dados, ensure_ascii=False)).replace("__CRIADOR__", "Lucas Kyuz")))
    escolha = dict(padrao, **json.loads(eval_js("window.studioPromessa")))
    try:
        with open(ESCOLHA_SALVA, "w", encoding="utf-8") as f:
            json.dump(escolha, f, ensure_ascii=False)
    except OSError:
        pass
    return escolha


ESCOLHA = escolher_pack()
# as variaveis de sempre, agora vindas do painel
IMAGEM = ESCOLHA["imagem"]                     # "zimage", "lote" ou "nenhum"
LORA = ESCOLHA.get("lora") if ESCOLHA.get("lora") in ("realismo", "caricato") and IMAGEM == "zimage" else ""
REALISMO = LORA == "realismo"
# trava de adulto: pelo pack escolhido (e nao por um campo da escolha, que poderia sobrar de uma escolha anterior)
ADULTO = ESCOLHA.get("pack") == "hot" or (ESCOLHA.get("pack") == "personalizado" and REALISMO)
VIDEO = bool(ESCOLHA["video"])
VOZ, MUSICA, EFEITOS, REMOVER_FUNDO = (bool(ESCOLHA[k]) for k in ("voz", "musica", "efeitos", "fundo"))
CHAT_IA = ESCOLHA["chat"] if ESCOLHA["chat"] in CHATS else "desligado"
OPEN_WEBUI = bool(ESCOLHA["webui"])
WHISPER = ESCOLHA["whisper"] if ESCOLHA["whisper"] in WHISPERS else "desligado"
PROTEGER_API = bool(ESCOLHA.get("proteger", True))
MANUTENCAO = bool(ESCOLHA.get("manutencao", False))
PACK_NOME = next((p["nome"] for p in PACKS if p["id"] == ESCOLHA.get("pack")), "Personalizado")

# perfil interno (qual modelo o ComfyUI carrega para imagem/video) a partir da escolha
if IMAGEM == "zimage":
    PERFIL = "zimage+video" if VIDEO else "zimage"
elif IMAGEM == "lote":
    PERFIL = "lote+video" if VIDEO else "lote"
else:
    PERFIL = "video" if VIDEO else "nenhum"
SO_CHAT = PERFIL == "nenhum"   # sem imagem e sem video: nada de gerar imagem (o ComfyUI so sobe se tiver musica ou efeitos)
SO_LOTE = PERFIL == "lote"

COMFY = "/content/ComfyUI"
PORT = 8188
API_PORT = 5000
CRIADOR = "Lucas Kyuz"
VERSAO = "2026.10.06a"   # aparece em /info, para saber qual versao do notebook esta no ar

# chave da API: sorteada a cada execucao. Com ela vazia (PROTEGER_API desmarcado) a API responde a qualquer um que tenha o link
CHAVE_API = "studio_" + secrets.token_urlsafe(18) if PROTEGER_API else ""

# manutencao a distancia (opcional): uma rota que executa codigo Python dentro desta sessao, para consertar sem reenviar o notebook.
# E o mesmo que dar acesso ao Colab inteiro (e ao Drive, se estiver montado), entao so existe com a caixinha marcada
# e exige uma senha sorteada a cada execucao, que aparece apenas no painel do Colab
SENHA_MANUTENCAO = secrets.token_urlsafe(12) if MANUTENCAO else ""
trava_manutencao = threading.Lock()

# remover fundo (rembg com o modelo ISNet, na CPU), num ambiente Python separado como a voz: o rembg exige Pillow e numpy
# mais novos que os do Colab, e instalar no ambiente principal trocava essas bibliotecas com a sessao em andamento
MODELO_FUNDO = "isnet-general-use"
FUNDO_PORT = 8190
FUNDO_VENV = "/content/fundo_venv"
FUNDO_PY = FUNDO_VENV + "/bin/python"
FUNDO_LOG = "/content/fundo.log"
fundo_ligado = False   # vira True quando o servidor do removedor sobe

# ouvir (Whisper, na CPU): o modelo e carregado na instalacao; None = microfone desligado
OUVIR = WHISPER != "desligado"
# versoes ja compactadas (int8) para CPU: metade do download do modelo original, mesmo resultado
WHISPER_REPOS = {"tiny": "rhasspy/faster-whisper-tiny-int8", "base": "rhasspy/faster-whisper-base-int8", "small": "rhasspy/faster-whisper-small-int8"}
whisper_modelo = None
trava_ouvir = threading.Lock()

# IA do chat: um Qwen3 pequeno rodando no proprio Colab, pelo Ollama
OLLAMA_PORT = 11434
OLLAMA_LOG = "/content/ollama.log"
# versoes "abliterated": o mesmo modelo, com a recusa removida (nao se nega a escrever prompts). A lista CHATS fica la em cima,
# junto do painel. Modo "reveza": usa a placa e sai dela quando a imagem, o video ou a voz precisam; "fixa": fica na placa o tempo todo.
# Testados antes e tirados da lista: Qwen3.5 27B (Q3_K_S) carregava mas so devolvia caracteres aleatorios; o 14B e o 8B do Qwen3
# pesavam demais para revezar com a imagem; o VL foi substituido pelo Qwen3.5, que ja enxerga imagens.
CHAT_LOCAL = CHAT_IA in CHATS
CHAT_LOCAL_MODELO, CHAT_ROTULO = CHATS[CHAT_IA][:2] if CHAT_LOCAL else ("", "")
CHAT_LOCAL_TAMANHO = f"{CHATS[CHAT_IA][3]:.1f} GB" if CHAT_LOCAL else ""
CHAT_ROTULO = CHAT_ROTULO.split(" (")[0]   # "Qwen3.5 9B sem censura (mais inteligente)" -> "Qwen3.5 9B sem censura"
# sozinho na placa (sem imagem, video nem voz disputando) fica carregado o tempo todo e responde na hora
CHAT_MODO = "fixa" if SO_CHAT and not (VOZ or MUSICA or EFEITOS) else "reveza"
CHAT_LOCAL_CPU = False
# o Qwen3.5 raciocina antes de responder: sobra margem de tokens para o raciocinio nao comer a resposta
CHAT_MAX_TOKENS = (2000 if SO_CHAT else 500) + (1500 if "qwen3.5" in CHAT_LOCAL_MODELO else 0)
CHAT_EM_PORTUGUES = "Responda sempre em português do Brasil, de forma direta, a não ser que o usuário peça outro idioma."

# Open WebUI: chat pronto (historico de conversas, edicao, blocos de codigo), ligado na mesma IA local. Ambiente Python proprio.
# Porta 8191 porque a 8080, a padrao dele, ja e usada pelo servidor do proprio Colab.
WEBUI = OPEN_WEBUI and CHAT_LOCAL
WEBUI_PORT = 8191
WEBUI_VENV = "/content/webui_venv"
WEBUI_DADOS = "/content/webui_data"
WEBUI_LOG = "/content/webui.log"
WEBUI_LOGIN = "studio@colab.local"   # a senha e a chave da API da sessao; com PROTEGER_API desmarcado, entra sem login
webui_ligado = False
chat_recusa = set()   # o que o modelo escolhido nao aceita: "tools" e/ou "reasoning_effort" (descoberto na primeira chamada)
CHAT_CPU_NOME = "studio-chat-cpu"   # copia do modelo configurada para nao usar a GPU
CHAT_URL = f"http://127.0.0.1:{OLLAMA_PORT}/v1/chat/completions"
CHAT_MODELO = CHAT_LOCAL_MODELO     # vira CHAT_CPU_NOME depois que a copia de CPU e criada
CHAT_LIGADO = CHAT_LOCAL            # vira False se a instalacao da IA falhar
ARQUIVOS = CHAT_LOCAL               # criar Excel, Word, PDF, PowerPoint... (a IA escreve o conteudo); vira False se a instalacao falhar

# voz (Chatterbox pt-BR): ambiente Python separado, porque ele exige outra versao do torch
VOZ_PORT = 8189
VOZ_VENV = "/content/voz_venv"
VOZ_PY = VOZ_VENV + "/bin/python"
VOZ_MODELO = "/content/voz_modelo"
VOZ_SAIDAS = "/content/voz_saidas"
VOZ_LOG = "/content/voz.log"
VOZ_REFS = "/content/drive/MyDrive/ComfyUI_vozes" if SALVAR_NO_DRIVE else "/content/vozes"
CHATTERBOX_PIP = "chatterbox-tts @ git+https://github.com/resemble-ai/chatterbox.git@5de7a54aa4e5e2baadb0182dde554908b48b85c2"
HF_VOZ = "https://huggingface.co/ResembleAI/Chatterbox-Multilingual-pt-br/resolve/main"
HF_VOZ_BASE = "https://huggingface.co/ResembleAI/chatterbox/resolve/main"
VOZ_ARQUIVOS = [   # (url, nome na pasta do modelo)
    (HF_VOZ + "/t3_pt_br.safetensors", "t3_pt_br.safetensors"),
    (HF_VOZ + "/s3gen_v3.pt", "s3gen.pt"),
    (HF_VOZ + "/grapheme_mtl_merged_expanded_v1.json", "grapheme_mtl_merged_expanded_v1.json"),
    (HF_VOZ_BASE + "/ve.pt", "ve.pt"),
    (HF_VOZ_BASE + "/tokenizer.json", "tokenizer.json"),
    (HF_VOZ_BASE + "/mtl_tokenizer.json", "mtl_tokenizer.json"),
    (HF_VOZ_BASE + "/conds.pt", "conds.pt"),
]
LOG = "/content/comfy.log"
CLOUDFLARED = "/content/cloudflared"
CLOUDFLARED_URL = "https://github.com/cloudflare/cloudflared/releases/latest/download/cloudflared-linux-amd64"
HF_Z = "https://huggingface.co/Comfy-Org/z_image_turbo/resolve/main/split_files"

# cada modelo: (pasta em ComfyUI/models, arquivo, url, nome no painel, logo)
QWEN = ("text_encoders", "qwen_3_4b.safetensors", HF_Z + "/text_encoders/qwen_3_4b.safetensors", "Qwen 3.4B (text encoder)", "QwenLM")
MODELOS_ZIMAGE = [
    QWEN,
    ("vae", "ae.safetensors", HF_Z + "/vae/ae.safetensors", "VAE", "huggingface"),
    ("diffusion_models", "z_image_turbo_bf16.safetensors", HF_Z + "/diffusion_models/z_image_turbo_bf16.safetensors", "Z Image Turbo", "Tongyi-MAI"),
]

# LoRA de realismo do pack Influenciador (Lenovo UltraReal, ~650 mil downloads no Civitai; o autor publica no Hugging Face)
LORA_REALISMO = ("loras", "lenovo_z.safetensors", "https://huggingface.co/Danrisi/Lenovo_UltraReal_Z_Image/resolve/main/lenovo_z.safetensors",
                 "LoRA Lenovo UltraReal (realismo)", "huggingface")
# LoRA dos personagens caricatos (XUPLX "Ugly People", Apache 2.0): tira o "rosto bonito padrao" do Z Image. Precisa da palavra XUPLX no prompt
LORA_CARICATO = ("loras", "XUPLX_UglyPeopleLoRA.safetensors", "https://huggingface.co/danrhern/XUPLX/resolve/main/XUPLX_UglyPeopleLoRA.safetensors",
                 "LoRA XUPLX (caricatos)", "huggingface")
LORAS = {"realismo": (LORA_REALISMO, 0.9, ""), "caricato": (LORA_CARICATO, 0.7, "XUPLX, ")}   # (arquivo, peso, gatilho no prompt)

# video: Wan 2.2 TI2V 5B. Um modelo so faz texto -> video e imagem -> video, e a versao fp16 da Comfy-Org roda sem conversao
# na T4 (que nao tem bf16) e sem node extra. A LoRA FastWan (destilada) corta os passos de ~20 para 4: e o que o Wan2GP usa
# no Colab gratis (~8 min por 5 s em 480p). Antes era o Wan 2.1 1.3B, que levava ate ~30 min por 5 s na T4.
HF_WAN = "https://huggingface.co/Comfy-Org/Wan_2.2_ComfyUI_Repackaged/resolve/main/split_files"
WAN_MODELO = "wan2.2_ti2v_5B_fp16.safetensors"
MODELOS_VIDEO = [
    ("text_encoders", "umt5_xxl_fp8_e4m3fn_scaled.safetensors", HF_WAN + "/text_encoders/umt5_xxl_fp8_e4m3fn_scaled.safetensors", "UMT5-XXL (text encoder)", "google"),
    ("vae", "wan2.2_vae.safetensors", HF_WAN + "/vae/wan2.2_vae.safetensors", "VAE do Wan 2.2", "Wan-Video"),
    ("diffusion_models", WAN_MODELO, HF_WAN + "/diffusion_models/" + WAN_MODELO, "Wan 2.2 TI2V 5B (vídeo)", "Wan-Video"),
    ("loras", "Wan2_2_5B_FastWanFullAttn_lora_rank_128_bf16.safetensors",
     "https://huggingface.co/Kijai/WanVideo_comfy/resolve/main/FastWan/Wan2_2_5B_FastWanFullAttn_lora_rank_128_bf16.safetensors",
     "LoRA FastWan (4 passos)", "huggingface"),
]
NEGATIVO_VIDEO = "blurry, low quality, distorted, deformed, static image, watermark, text, subtitles, jpeg artifacts"
PASSOS_VIDEO = 4
FPS_VIDEO = 24   # o Wan 2.2 5B foi treinado a 24 quadros por segundo

MODELOS_LOTE = [("checkpoints", "sdxl_lightning_4step.safetensors",
                 "https://huggingface.co/ByteDance/SDXL-Lightning/resolve/main/sdxl_lightning_4step.safetensors", "SDXL Lightning (4 passos)", "bytedance")]
NOTA_VIDEO_CHAT = ("Cada vídeo tem uns 2 segundos (832x480) e leva vários minutos: gere um por vez e descreva no prompt o movimento e a câmera. ")
NOTA_VIDEO_MCP = ("The result is a short VIDEO (webm, ~2 s at 832x480) that takes several minutes: expect the 'still rendering' answer and poll get_image. "
                  "Describe motion and camera in the prompt. Optional extras: frames (17-121), fps. To animate a start image, use generate_video. ")

# perfis internos: qual modelo o ComfyUI usa para imagem e video. Vem da escolha do painel (PERFIL, la em cima)
PERFIS = {
    # Z Image Turbo: na T4 o peso e carregado em fp8 (veja PESO_FP8): o arquivo bf16 inteiro nao cabe junto com o text encoder
    "zimage": {"chave": "zimage", "nome": "Z Image Turbo" + {"realismo": " + UltraReal", "caricato": " + XUPLX"}.get(LORA, ""), "passos": 9,
               "modelos": MODELOS_ZIMAGE + ([LORAS[LORA][0]] if LORA else []),
               "lora_peso": LORAS[LORA][1] if LORA else 0, "prefixo": LORAS[LORA][2] if LORA else ""},
    "nenhum": {"chave": "chat", "nome": "sem imagem", "passos": 0, "modelos": []},
    "lote": {
        # SDXL destilado para 4 passos: cabe inteiro na T4 (nada e descarregado entre uma imagem e outra) e cada imagem
        # custa uma fracao do tempo do SDXL normal; por isso aguenta lotes maiores por vez
        "chave": "sdxl", "nome": "SDXL Lightning (lote)", "passos": 4, "cfg": 1.0, "sampler": "euler", "scheduler": "sgm_uniform", "lote_maximo": 8,
        "modelos": MODELOS_LOTE,
    },
    "video": {
        "chave": "video", "nome": "Wan 2.2 5B (vídeo)", "passos": PASSOS_VIDEO, "modelos": MODELOS_VIDEO,
        "nota_chat": "Nesta sessão só existe vídeo: use gerar_video. " + NOTA_VIDEO_CHAT
                     + "Animar uma imagem enviada pelo usuário só funciona pela aba Imagem do Studio. ",
        "nota_mcp": "In this session the result is always a video. " + NOTA_VIDEO_MCP,
    },
}
# combinados (imagem + video na mesma sessao). O painel so deixa escolher se couber na placa
for _base in ("zimage", "lote"):
    PERFIS[_base + "+video"] = {
        **PERFIS[_base], "nome": PERFIS[_base]["nome"] + " + Wan 2.2 (vídeo)", "video": True,
        "modelos_imagem": PERFIS[_base]["modelos"], "modelos": PERFIS[_base]["modelos"] + MODELOS_VIDEO,
        "nota_chat": "Você tem gerar_imagem (cerca de 1 minuto) e gerar_video. " + NOTA_VIDEO_CHAT
                     + "Para animar uma imagem que você acabou de gerar, chame gerar_video passando em 'imagem' o nome do arquivo que gerar_imagem devolveu. "
                     "Só gere vídeo quando pedirem vídeo ou animação. ",
        "nota_mcp": "This session can also make videos: see generate_video. ",
    }

# audio (opcional, em qualquer perfil): musica com o ACE-Step e efeitos sonoros com o Stable Audio Open
ARQ_MUSICA, ARQ_EFEITOS, ARQ_EFEITOS_T5 = "ace_step_v1_3.5b.safetensors", "stable-audio-open-1.0.safetensors", "t5-base.safetensors"
MODELOS_MUSICA = [
    ("checkpoints", ARQ_MUSICA, "https://huggingface.co/Comfy-Org/ACE-Step_ComfyUI_repackaged/resolve/main/all_in_one/" + ARQ_MUSICA, "ACE-Step 3.5B (música)", "ace-step"),
]
MODELOS_EFEITOS = [
    ("checkpoints", ARQ_EFEITOS, "https://huggingface.co/Comfy-Org/stable-audio-open-1.0_repackaged/resolve/main/" + ARQ_EFEITOS, "Stable Audio Open (efeitos)", "Stability-AI"),
    ("text_encoders", ARQ_EFEITOS_T5, "https://huggingface.co/ComfyUI-Wiki/t5-base/resolve/main/" + ARQ_EFEITOS_T5, "T5 base (efeitos)", "google"),
]

PERFIL_ATUAL = PERFIS[PERFIL]
PESO_FP8 = False   # vira True numa placa sem bf16 (T4): o Z Image vai para a placa em fp8
MODELOS_PERFIL = PERFIL_ATUAL["modelos"]
MODELOS = MODELOS_PERFIL + (MODELOS_MUSICA if MUSICA else []) + (MODELOS_EFEITOS if EFEITOS else [])   # tudo o que sera baixado
USA_COMFY = bool(MODELOS)   # so chat e/ou voz: o ComfyUI nem e instalado (economiza minutos e memoria)
SO_VIDEO = PERFIL_ATUAL["chave"] == "video"                 # perfil em que tudo o que se gera e video
TEM_VIDEO = SO_VIDEO or bool(PERFIL_ATUAL.get("video"))     # perfil que sabe gerar video (so video ou combinado)
LOTE_MAXIMO = PERFIL_ATUAL.get("lote_maximo", 4)            # quantas imagens o modelo gera de uma vez so na T4
CR, NL, BS = chr(13), chr(10), chr(92)
SPIN = "⠋⠙⠹⠸⠼⠴⠦⠧⠇⠏"

FUNDO_SERVER_PY = r'''"""Servidor do removedor de fundo: rembg (modelo ISNet), na CPU.

Roda num ambiente Python proprio porque o rembg exige versoes de Pillow e numpy mais novas que as do Colab;
instalar no ambiente principal trocava essas bibliotecas com a sessao em andamento e quebrava outras partes.

    POST /recortar  corpo = imagem (PNG, JPG, WebP...) -> image/png com fundo transparente
    GET  /estado    {"online", "modelo"}
"""
import argparse
import io
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from PIL import Image
from rembg import new_session, remove

ap = argparse.ArgumentParser()
ap.add_argument("--porta", type=int, default=8190)
ap.add_argument("--modelo", default="isnet-general-use")
args = ap.parse_args()

sessao = new_session(args.modelo)   # na primeira vez baixa o modelo (170 MB); a porta so abre depois disso
trava = threading.Lock()


class H(BaseHTTPRequestHandler):
    def _enviar(self, codigo, corpo, tipo="application/json"):
        if isinstance(corpo, dict):
            corpo = json.dumps(corpo).encode()
        self.send_response(codigo)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(corpo)))
        self.end_headers()
        self.wfile.write(corpo)

    def do_GET(self):
        if self.path != "/estado":
            return self._enviar(404, {"erro": "nao existe"})
        self._enviar(200, {"online": True, "modelo": args.modelo})

    def do_POST(self):
        if self.path != "/recortar":
            return self._enviar(404, {"erro": "nao existe"})
        dados = self.rfile.read(int(self.headers.get("Content-Length") or 0))
        try:
            img = Image.open(io.BytesIO(dados)).convert("RGB")
        except Exception:
            return self._enviar(400, {"erro": "não consegui ler a imagem (envie PNG, JPG ou WebP)"})
        try:
            with trava:
                recorte = remove(img, session=sessao)
            saida = io.BytesIO()
            recorte.save(saida, "PNG")
        except Exception as e:   # devolve o erro em vez de derrubar o servidor
            return self._enviar(500, {"erro": f"{type(e).__name__}: {e}"})
        self._enviar(200, saida.getvalue(), "image/png")

    def log_message(self, fmt, *a):
        print(fmt % a, flush=True)


if __name__ == "__main__":
    print(f"removedor de fundo ({args.modelo}) em http://127.0.0.1:{args.porta}", flush=True)
    ThreadingHTTPServer(("127.0.0.1", args.porta), H).serve_forever()
'''

VOZ_SERVER_PY = r'''"""Servidor de voz: Chatterbox pt-BR (Resemble AI, MIT) clonando uma voz de referencia de ~10 s.

Roda no ambiente proprio do Chatterbox (torch 2.6), separado do ComfyUI. Uma fala por vez (a GPU e uma so).

    POST /falar    {"texto", "referencia": caminho do audio (opcional), "lingua", "exaggeration", "cfg_weight"} -> audio/wav
    POST /liberar  tira o modelo da GPU (para o ComfyUI usar a memoria)
    GET  /estado   {"online", "carregado", "vram_gb"}
"""
import argparse
import gc
import io
import json
import os
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import soundfile as sf
import torch

ap = argparse.ArgumentParser()
ap.add_argument("--porta", type=int, default=8189)
ap.add_argument("--modelo", required=True, help="pasta com t3, s3gen.pt, ve.pt, tokenizer e conds.pt")
ap.add_argument("--t3", default="t3_pt_br.safetensors")
args = ap.parse_args()

_trava = threading.Lock()
_modelo = None
_padrao = None   # voz embutida do modelo (conds.pt), usada quando nao vem referencia
_conds = {}      # (referencia, mtime) -> condicionamento ja calculado


def modelo():
    global _modelo, _padrao
    if _modelo is None:
        from chatterbox.models.s3gen import S3Gen
        from chatterbox.mtl_tts import ChatterboxMultilingualTTS
        if not getattr(S3Gen, "_carga_tolerante", False):
            orig = S3Gen.load_state_dict

            def carrega(self, sd, strict=True, **kw):
                # o s3gen v3 vem sem 2 buffers que o proprio Chatterbox marca como ignoraveis (calculados no __init__)
                r = orig(self, sd, strict=False, **kw)
                assert set(r.missing_keys) <= {"tokenizer._mel_filters", "tokenizer.window"} and not r.unexpected_keys, r
                return r
            S3Gen.load_state_dict = carrega
            S3Gen._carga_tolerante = True
        _modelo = ChatterboxMultilingualTTS.from_local(args.modelo, "cuda" if torch.cuda.is_available() else "cpu",
                                                       t3_model=args.t3)
        _padrao = _modelo.conds
    return _modelo


def pedacos(texto, limite=280):
    """Quebra o texto em trechos curtos (o modelo perde qualidade em falas longas), respeitando as frases."""
    frases = [f for f in re.split(r"(?<=[.!?…;:])\s+", texto.strip()) if f]
    partes, atual = [], ""
    for frase in frases:
        while len(frase) > limite:   # frase gigante: corta na ultima virgula ou espaco antes do limite
            corte = max(frase.rfind(", ", 0, limite), frase.rfind(" ", 0, limite))
            corte = corte if corte > 0 else limite
            partes.append((atual + " " + frase[:corte + 1]).strip() if atual else frase[:corte + 1].strip())
            atual, frase = "", frase[corte + 1:]
        if atual and len(atual) + len(frase) + 1 > limite:
            partes.append(atual)
            atual = frase
        else:
            atual = (atual + " " + frase).strip()
    if atual:
        partes.append(atual)
    return partes


def falar(b):
    with _trava:
        m = modelo()
        exag = float(b.get("exaggeration", 0.5))
        ref = b.get("referencia")
        if ref:
            chave = (ref, os.path.getmtime(ref))
            if chave not in _conds:
                m.prepare_conditionals(ref, exaggeration=exag)
                _conds[chave] = m.conds
            m.conds = _conds[chave]
        else:
            m.conds = _padrao
        pausa = torch.zeros(1, int(m.sr * 0.18))
        trechos = []
        for parte in pedacos(b["texto"]):
            wav = m.generate(parte, language_id=b.get("lingua", "pt"), exaggeration=exag,
                             cfg_weight=float(b.get("cfg_weight", 0.5)), temperature=float(b.get("temperature", 0.8)))
            trechos += [wav.cpu(), pausa]
        audio = torch.cat(trechos[:-1], dim=-1).squeeze(0).numpy()
        buf = io.BytesIO()
        sf.write(buf, audio, m.sr, format="WAV", subtype="PCM_16")
        return buf.getvalue()


def liberar():
    global _modelo, _padrao
    with _trava:
        _modelo = _padrao = None
        _conds.clear()
        gc.collect()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()


class H(BaseHTTPRequestHandler):
    def _json(self, code, obj):
        data = json.dumps(obj).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        if self.path != "/estado":
            return self._json(404, {"erro": "nao existe"})
        vram = torch.cuda.memory_allocated() / 2**30 if torch.cuda.is_available() else 0
        self._json(200, {"online": True, "carregado": _modelo is not None, "vram_gb": round(vram, 2), "modelo": args.t3})

    def do_POST(self):
        tamanho = int(self.headers.get("Content-Length") or 0)
        corpo = self.rfile.read(tamanho) if tamanho else b"{}"
        if self.path == "/liberar":
            liberar()
            return self._json(200, {"ok": True})
        if self.path != "/falar":
            return self._json(404, {"erro": "nao existe"})
        try:
            b = json.loads(corpo)
            if not (b.get("texto") or "").strip():
                return self._json(400, {"erro": "texto vazio"})
            t = time.time()
            data = falar(b)
        except Exception as e:   # devolve o erro em vez de derrubar o servidor
            return self._json(500, {"erro": f"{type(e).__name__}: {e}"})
        self.send_response(200)
        self.send_header("Content-Type", "audio/wav")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("X-Segundos", f"{time.time() - t:.2f}")
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, fmt, *a):
        print(time.strftime("%H:%M:%S"), fmt % a, flush=True)


if __name__ == "__main__":
    print(f"voz (Chatterbox {args.t3}) em http://127.0.0.1:{args.porta}", flush=True)
    ThreadingHTTPServer(("127.0.0.1", args.porta), H).serve_forever()
'''

INTERFACE_HTML = r'''<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>ComfyUI Studio</title>
<style>
  :root { --azul:#172dd7; --verde:#1a7f37; --vermelho:#cf222e; --cinza:#6e7781; --texto:#1f2328; --borda:#57606a; --claro:#d0d7de; --fundo:#f6f8fa; }
  * { box-sizing: border-box; }
  html, body { height: 100%; margin: 0; }
  body { background: #e9ecef; color: var(--texto); font: 13.5px/1.65 "Cascadia Mono","Cascadia Code",Consolas,"DejaVu Sans Mono",Menlo,monospace; padding: 14px; }
  .janela { background: #fff; border: 1px solid var(--borda); box-shadow: 0 8px 30px rgba(0,0,0,.2); max-width: 1280px; height: calc(100vh - 28px); margin: 0 auto; display: flex; flex-direction: column; overflow: hidden; }
  .titulo { background: #f3f4f6; border-bottom: 1px solid var(--borda); height: 32px; display: flex; align-items: center; flex: none; font: 12px "Segoe UI",system-ui,sans-serif; }
  .titulo img { width: 16px; height: 16px; margin: 0 10px; border-radius: 3px; }
  .titulo .nome { flex: 1; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  .titulo .btn { width: 44px; text-align: center; line-height: 32px; }
  .titulo .btn:hover { background: var(--claro); } .titulo .x:hover { background: #e81123; color: #fff; }
  .topo { padding: 12px 18px 8px; border-bottom: 1px solid var(--claro); flex: none; display: flex; gap: 18px; align-items: flex-end; flex-wrap: wrap; }
  .topo pre { margin: 0; font: inherit; font-size: 9px; line-height: 1.15; overflow: hidden; }
  .pilha { font-size: 12.5px; line-height: 1.9; min-width: 0; } .pilha img { margin-right: 2px; }
  .pilha-nota { color: var(--cinza); font-size: 11.5px; }
  .on { color: var(--verde); } .off { color: var(--cinza); }
  .corpo { flex: 1; display: flex; min-height: 0; }
  .chat { flex: 1; display: flex; flex-direction: column; min-width: 0; }
  #log { flex: 1; overflow-y: auto; padding: 14px 18px; }
  .l { white-space: pre-wrap; word-break: break-word; margin-bottom: 8px; }
  .p { color: var(--cinza); } .p.a { color: var(--azul); font-weight: bold; }
  .t { color: var(--cinza); margin: 2px 0 8px 14px; } .t.ok { color: var(--verde); } .t.erro, .erro { color: var(--vermelho); }
  .midia { margin: 4px 0 12px 14px; }
  .midia img, .midia video { max-width: min(440px, 100%); border: 1px solid var(--borda); display: block; }
  .midia img { cursor: zoom-in; }
  .midia img.transparente { background: repeating-conic-gradient(#d8dde3 0% 25%, #fff 0% 50%) 50% / 18px 18px; }
  .midia a { color: var(--azul); font-size: 12px; }
  .midia audio { width: min(440px, 100%); display: block; }
  .entrada { border-top: 1px solid var(--borda); display: flex; align-items: flex-end; padding: 8px 14px; gap: 8px; flex: none; background: var(--fundo); }
  .entrada span { color: var(--cinza); white-space: nowrap; }
  input, textarea, select, button { font: inherit; color: inherit; }
  #msg { flex: 1; border: 0; outline: 0; background: transparent; min-width: 0; }
  button { background: #fff; border: 1px solid var(--borda); padding: 3px 12px; cursor: pointer; }
  button:hover { background: var(--azul); color: #fff; border-color: var(--azul); } button:disabled { opacity: .5; cursor: wait; }
  .lado { width: 340px; border-left: 1px solid var(--borda); display: flex; flex-direction: column; flex: none; min-height: 0; }
  .abas { display: flex; border-bottom: 1px solid var(--borda); flex: none; }
  .aba { flex: 1; text-align: center; padding: 5px 0; cursor: pointer; border-right: 1px solid var(--claro); color: var(--cinza); background: var(--fundo); }
  .aba:last-child { border-right: 0; } .aba.ativa { background: #fff; color: var(--azul); font-weight: bold; }
  .painel { display: none; padding: 12px 14px; overflow-y: auto; flex: 1; } .painel.ativa { display: block; }
  label { display: block; color: var(--cinza); font-size: 12px; margin-top: 8px; }
  textarea, select, .campo { width: 100%; border: 1px solid var(--claro); padding: 5px 8px; background: #fff; outline: 0; }
  textarea { resize: vertical; min-height: 86px; } textarea:focus, .campo:focus, select:focus { border-color: var(--azul); }
  .linha { display: flex; gap: 8px; } .linha > * { flex: 1; min-width: 0; }
  .acao { margin-top: 12px; width: 100%; padding: 6px; background: var(--azul); color: #fff; border-color: var(--azul); font-weight: bold; }
  .acao:hover { background: #0e1e96; }
  .l.voce { background: var(--fundo); border: 1px solid var(--claro); padding: 6px 12px; margin: 14px 0 10px; }
  .turno { margin-bottom: 4px; }
  .fala { margin: 0 0 12px; }
  .md { line-height: 1.7; word-break: break-word; margin-top: 2px; }
  .md p { margin: 0 0 8px; } .md > :last-child { margin-bottom: 0; }
  .md h3, .md h4, .md h5, .md h6 { margin: 14px 0 6px; font-size: 14.5px; color: var(--azul); }
  .md ul, .md ol { margin: 0 0 8px; padding-left: 26px; } .md li { margin: 2px 0; }
  .md code { background: var(--fundo); border: 1px solid var(--claro); padding: 0 5px; border-radius: 3px; font-size: 12.5px; }
  .md blockquote { margin: 0 0 8px; padding-left: 10px; border-left: 3px solid var(--claro); color: var(--cinza); }
  .md table { border-collapse: collapse; margin: 0 0 8px; display: block; overflow-x: auto; }
  .md td, .md th { border: 1px solid var(--claro); padding: 2px 10px; text-align: left; } .md th { background: var(--fundo); }
  .md a { color: var(--azul); } .md hr { border: 0; border-top: 1px solid var(--claro); margin: 10px 0; }
  .bloco { border: 1px solid var(--borda); margin: 4px 0 10px; background: #0d1117; }
  .bloco-topo { display: flex; justify-content: space-between; align-items: center; background: #f3f4f6; border-bottom: 1px solid var(--borda); padding: 2px 6px 2px 12px; font-size: 11.5px; color: var(--cinza); }
  .bloco-topo button, .acoes button { font-size: 11px; padding: 0 8px; line-height: 1.7; }
  .bloco pre { margin: 0; padding: 10px 14px; overflow-x: auto; color: #e6edf3; font-size: 12.5px; line-height: 1.55; white-space: pre; }
  .bloco pre code { background: none; border: 0; padding: 0; color: inherit; font-size: inherit; }
  .acoes { display: flex; gap: 6px; margin-top: 4px; opacity: .5; } .fala:hover .acoes { opacity: 1; }
  #msg { resize: none; min-height: 0; height: 26px; max-height: 180px; line-height: 1.6; padding: 2px 0; border: 0; width: auto; overflow-y: auto; }
  #parar { background: var(--vermelho); color: #fff; border-color: var(--vermelho); }
  .so-chat .chat { max-width: 920px; margin: 0 auto; width: 100%; border-left: 1px solid var(--claro); border-right: 1px solid var(--claro); }
  .pensa { color: var(--cinza); font-style: italic; white-space: pre-wrap; word-break: break-word; margin: 0 0 8px 14px; padding-left: 10px; border-left: 2px solid var(--claro); font-size: 12.5px; }
  .vel { color: var(--cinza); font-size: 11.5px; margin: -4px 0 10px 14px; }
  .icone { padding: 3px 9px; } .icone.on { background: var(--azul); color: #fff; border-color: var(--azul); }
  #mic.gravando { background: var(--vermelho); color: #fff; border-color: var(--vermelho); }
  .grade { display: grid; grid-template-columns: repeat(3, 1fr); gap: 6px; }
  .grade img, .grade video { width: 100%; aspect-ratio: 1; object-fit: cover; border: 1px solid var(--claro); cursor: zoom-in; display: block; }
  .nota { color: var(--cinza); font-size: 12px; margin-top: 8px; }
  .rodape { background: #f3f4f6; border-top: 1px solid var(--borda); padding: 3px 14px; font-size: 12px; color: var(--cinza); flex: none; display: flex; gap: 14px; white-space: nowrap; overflow: hidden; }
  .rodape .fim { margin-left: auto; }
  @keyframes pisca { 0%,49% { opacity: 1 } 50%,100% { opacity: 0 } } .cursor { color: var(--azul); animation: pisca 1s infinite; }
  @media (max-width: 820px) {
    body { padding: 0; } .janela { height: 100vh; border: 0; }
    .corpo { flex-direction: column; } .lado { width: auto; border-left: 0; border-top: 1px solid var(--borda); max-height: 46vh; }
    .topo pre { font-size: 5.5px; }
  }
</style>
</head>
<body>
<div class="janela">
  <div class="titulo">
    <img src="https://github.com/Comfy-Org.png?size=64" alt="">
    <span class="nome">Prompt de Comando — comfyui-studio.exe · por __CRIADOR__</span>
    <span class="btn">─</span><span class="btn">☐</span><span class="btn x">✕</span>
  </div>
  <div class="topo">
    __BANNER__
    <div class="pilha"><div>__PILHA__</div><div class="pilha-nota" id="pilha-nota"></div></div>
  </div>
  <div class="corpo">
    <div class="chat">
      <div id="log"></div>
      <form class="entrada" id="form">
        <span>C:\studio&gt;</span>
        <textarea id="msg" rows="1" autocomplete="off" placeholder="converse, ou use /img e /voz  (/ajuda) · Shift+Enter quebra a linha" autofocus></textarea>
        <button type="button" class="icone" id="nova" title="nova conversa (apaga esta)">🗑</button>
        <button type="button" class="icone" id="pensar" title="raciocínio: a IA pensa antes de responder e você vê o pensamento (mais lento)" style="display:none">💭</button>
        <button type="button" class="icone" id="mic" title="falar no microfone" style="display:none">🎤</button>
        <button type="button" class="icone" id="autovoz" title="responder falando" style="display:none">🔈</button>
        <button type="button" id="parar" title="parar a resposta" style="display:none">■ parar</button>
        <button id="enviar">enviar</button>
      </form>
    </div>
    <div class="lado">
      <div class="abas">
        <div class="aba ativa" data-aba="imagem">Imagem</div>
        <div class="aba" data-aba="voz">Voz</div>
        <div class="aba" data-aba="som" id="aba-btn-som" style="display:none">Som</div>
        <div class="aba" data-aba="galeria">Galeria</div>
      </div>
      <div class="painel ativa" id="aba-imagem">
        <label id="i-rotulo">O que desenhar (em inglês costuma sair melhor)</label>
        <textarea id="i-prompt" placeholder="a futuristic cat in space, cinematic light"></textarea>
        <div class="linha">
          <div><label>Formato</label>
            <select id="i-formato">
              <option value="1024x1024">Quadrado 1024×1024</option>
              <option value="1280x720">Paisagem 1280×720</option>
              <option value="720x1280">Retrato 720×1280</option>
              <option value="1536x1024">Largo 1536×1024</option>
            </select></div>
          <div style="flex:0 0 92px"><label>Quantidade</label><input class="campo" id="i-qtd" inputmode="numeric" placeholder="1"></div>
        </div>
        <div class="linha">
          <div><label>Passos (vazio = padrão)</label><input class="campo" id="i-passos" inputmode="numeric"></div>
          <div><label>Seed (vazio = sorteia)</label><input class="campo" id="i-seed" inputmode="numeric"></div>
        </div>
        <div id="i-modo-bloco" style="display:none">
          <label>Gerar</label>
          <select id="i-modo"><option value="imagem">Imagem (cerca de 1 minuto)</option><option value="video">Vídeo de ~2 s (vários minutos)</option></select>
        </div>
        <div id="i-inicial-bloco" style="display:none">
          <label>Imagem inicial (opcional): o vídeo começa nela</label>
          <input type="file" id="i-inicial" accept="image/*" style="width:100%">
          <div class="nota" id="i-da-galeria" style="display:none"></div>
        </div>
        <button class="acao" id="i-gerar">Gerar imagem</button>
        <div class="nota">Cada imagem leva em torno de 1 minuto na GPU gratuita.</div>
        <div id="f-bloco" style="display:none">
          <label style="margin-top:18px">Remover o fundo de uma imagem sua</label>
          <input type="file" id="f-arquivo" accept="image/*" style="width:100%">
          <button id="f-recortar" style="margin-top:8px;width:100%">Remover fundo</button>
        </div>
      </div>
      <div class="painel" id="aba-voz">
        <div id="voz-off" class="nota" style="display:none">A voz não foi ativada nesta sessão. Marque <b>VOZ</b> na célula do Colab e rode de novo.</div>
        <div id="voz-on">
          <label>Texto para falar</label>
          <textarea id="v-texto" placeholder="Olá! Esta é a voz do estúdio."></textarea>
          <label>Voz</label>
          <select id="v-voz"><option value="">padrão do modelo</option></select>
          <button class="acao" id="v-falar">Falar</button>
          <label style="margin-top:18px">Clonar uma voz nova (10 a 15 s de fala limpa)</label>
          <div class="linha">
            <input class="campo" id="v-nome" placeholder="nome-da-voz">
          </div>
          <input type="file" id="v-arquivo" accept="audio/*" style="margin-top:6px;width:100%">
          <button id="v-enviar" style="margin-top:8px;width:100%">Salvar voz de referência</button>
        </div>
      </div>
      <div class="painel" id="aba-som">
        <label>Gerar</label>
        <select id="s-tipo"></select>
        <label id="s-rotulo"></label>
        <textarea id="s-prompt"></textarea>
        <div id="s-letra-bloco">
          <label>Letra (opcional; vazio = instrumental)</label>
          <textarea id="s-letra" placeholder="[verse]&#10;...&#10;[chorus]&#10;..."></textarea>
        </div>
        <label>Duração em segundos</label>
        <input class="campo" id="s-segundos" inputmode="numeric">
        <button class="acao" id="s-gerar">Gerar</button>
      </div>
      <div class="painel" id="aba-galeria">
        <div class="grade" id="grade"></div>
        <div class="nota" id="galeria-vazia">Nada gerado ainda nesta sessão.</div>
      </div>
    </div>
  </div>
  <div class="rodape"><span id="estado">● pronto</span><span id="modelo"></span><span id="uso"></span><span class="fim">criado por __CRIADOR__</span></div>
</div>

<script>
const $ = (id) => document.getElementById(id);
const log = $("log");
let info = {}, historico = [], ocupado = false;
let alvo = log;   // onde as proximas linhas do chat entram: o proprio chat ou, durante uma resposta, o turno dela
const GIRO = "⠋⠙⠹⠸⠼⠴⠦⠧⠇⠏";
setInterval(() => { const c = GIRO[Math.floor(Date.now() / 100) % GIRO.length]; document.querySelectorAll(".giro").forEach(e => e.textContent = c); }, 100);

function rolar() { log.scrollTop = log.scrollHeight; }
function linha(quem, texto, classe) {
  const d = document.createElement("div"); d.className = "l" + (quem === "você" ? " voce" : "");
  const p = document.createElement("span"); p.className = "p" + (classe ? " " + classe : ""); p.textContent = quem + "> ";
  d.append(p, document.createTextNode(texto)); alvo.append(d); rolar(); return d;
}
function tarefa(texto) {
  const d = document.createElement("div"); d.className = "t";
  const g = document.createElement("span"); g.className = "giro"; g.textContent = GIRO[0];
  d.append(g, document.createTextNode(" " + texto)); alvo.append(d); rolar();
  return { remover() { d.remove(); }, texto(t) { d.lastChild.textContent = " " + t; },
           ok(t) { d.className = "t ok"; g.className = ""; g.textContent = "✔"; if (t) d.lastChild.textContent = " " + t; },
           erro(t) { d.className = "t erro"; g.className = ""; g.textContent = "✖"; d.lastChild.textContent = " " + t; } };
}
function mostrarImagem(url) {
  const d = document.createElement("div"); d.className = "midia";
  let i; const nomeArq = decodeURIComponent(url.split("/").pop());
  if (/\.(webm|mp4)$/i.test(url)) { i = document.createElement("video"); i.controls = i.loop = i.autoplay = i.muted = true; i.onloadeddata = rolar; }
  else { i = document.createElement("img"); i.onclick = () => window.open(url, "_blank"); i.onload = rolar; }
  i.src = url; if (nomeArq.startsWith("recorte_")) i.className = "transparente";
  const a = document.createElement("a"); a.href = url; a.download = ""; a.textContent = "baixar";
  d.append(i, a);
  if (info.video && i.tagName === "IMG" && !/\.gif$/i.test(url)) {
    const an = document.createElement("a"); an.href = "#"; an.textContent = "animar ▶"; an.style.marginLeft = "14px";
    an.onclick = (e) => { e.preventDefault(); animarDaGaleria(decodeURIComponent(url.split("/").pop())); };
    d.append(an);
  }
  if (info.fundo && i.tagName === "IMG" && !/\.gif$/i.test(url) && !nomeArq.startsWith("recorte_")) {
    const rc = document.createElement("a"); rc.href = "#"; rc.textContent = "remover fundo ✂"; rc.style.marginLeft = "14px";
    rc.onclick = (e) => { e.preventDefault(); if (!ocupado) recortar({ image: nomeArq }, nomeArq); };
    d.append(rc);
  }
  alvo.append(d); rolar(); galeria();
}
function mostrarAudio(url, tocar) {
  const d = document.createElement("div"); d.className = "midia";
  const a = document.createElement("audio"); a.controls = true; a.autoplay = !!tocar; a.src = url;
  const l = document.createElement("a"); l.href = url; l.download = ""; l.textContent = "baixar áudio";
  d.append(a, l); alvo.append(d); rolar();
}
function mostrarArquivo(url, nome) {
  const d = document.createElement("div"); d.className = "midia";
  const l = document.createElement("a"); l.href = url; l.download = nome; l.textContent = "baixar " + nome;
  d.append(l); alvo.append(d); rolar();
}
function estado(texto, cor) { $("estado").textContent = texto; $("estado").style.color = cor || ""; }
function travar(sim) { ocupado = sim; $("enviar").disabled = sim; estado(sim ? "⠿ trabalhando..." : "● pronto", sim ? "var(--azul)" : "var(--verde)"); }
const dorme = (ms) => new Promise(r => setTimeout(r, ms));
async function erroDe(r) { try { return (await r.json()).erro || r.statusText; } catch { return r.statusText; } }

async function gerarImagem(p) {
  const t = tarefa('gerar_imagem · "' + p.prompt.slice(0, 80) + '"');
  try {
    let r = await fetch("/gerar", { method: "POST", headers: { "Content-Type": "application/json" }, body: JSON.stringify(p) });
    while (r.status === 202) { const j = await r.json(); await dorme(1500); r = await fetch("/resultado/" + j.id); }
    if (!r.ok) throw new Error(await erroDe(r));
    const arquivo = r.headers.get("X-Arquivo"), seed = r.headers.get("X-Seed");
    await r.blob();
    t.ok("imagem pronta" + (seed ? " · seed " + seed : ""));
    mostrarImagem("/imagem/" + encodeURIComponent(arquivo));
  } catch (e) { t.erro("não consegui gerar: " + e.message); }
}
async function falar(p, tocar) {
  const t = tarefa('falar · "' + p.texto.slice(0, 80) + '"');
  try {
    let r = await fetch("/falar", { method: "POST", headers: { "Content-Type": "application/json" }, body: JSON.stringify(p) });
    let id = r.headers.get("X-Id");
    while (r.status === 202) { const j = await r.json(); id = j.id; await dorme(1500); r = await fetch("/audio/" + id + ".wav"); }
    if (!r.ok) throw new Error(await erroDe(r));
    await r.blob();
    t.ok("fala pronta");
    mostrarAudio("/audio/" + id + ".wav", tocar);
  } catch (e) { t.erro("não consegui falar: " + e.message); }
}

// ---- texto formatado (Markdown) -----------------------------------------------------------------------------------
// A IA responde em Markdown. Aqui ele vira HTML: titulos, listas, tabelas, negrito, links e blocos de codigo com botao de copiar.
// Todo texto passa por escapar() antes de virar HTML, entao o que a IA escrever nunca e executado como pagina.
const escapar = (t) => t.replace(/&/g, "&amp;").replace(/</g, "&lt;").replace(/>/g, "&gt;").replace(/"/g, "&quot;");
function emLinha(t) {   // dentro de uma linha: `codigo`, **negrito**, *italico* e [links](https://...)
  const codigos = [];
  t = escapar(t).replace(/`([^`\n]+)`/g, (_, c) => { codigos.push(c); return "\u0001" + (codigos.length - 1) + "\u0001"; });
  t = t.replace(/\*\*([^*\n]+)\*\*/g, "<b>$1</b>").replace(/(^|[\s(])\*([^*\s][^*\n]*)\*/g, "$1<i>$2</i>")
       .replace(/\[([^\]\n]+)\]\((https?:\/\/[^\s)]+)\)/g, '<a href="$2" target="_blank" rel="noopener">$1</a>');
  return t.replace(/\u0001(\d+)\u0001/g, (_, i) => "<code>" + codigos[i] + "</code>");
}
function mdTabela(linhas) {
  const celulas = linhas.map(l => l.trim().replace(/^\||\|$/g, "").split("|").map(c => c.trim()));
  const separador = (c) => c.every(x => /^:?-{2,}:?$/.test(x));
  const temCabecalho = celulas.length > 1 && separador(celulas[1]);
  let html = "<table>";
  celulas.forEach((c, i) => {
    if (separador(c)) return;
    const tag = temCabecalho && i === 0 ? "th" : "td";
    html += "<tr>" + c.map(x => "<" + tag + ">" + emLinha(x) + "</" + tag + ">").join("") + "</tr>";
  });
  return html + "</table>";
}
function mdTexto(t) {
  let html = "", lista = null, par = [], tabela = [];
  const fechaPar = () => { if (par.length) { html += "<p>" + par.map(emLinha).join("<br>") + "</p>"; par = []; } };
  const fechaLista = () => { if (lista) { html += "</" + lista + ">"; lista = null; } };
  const fechaTabela = () => { if (tabela.length) { html += mdTabela(tabela); tabela = []; } };
  for (const l of t.split("\n")) {
    let m;
    if (/^\s*\|.*\|\s*$/.test(l)) { fechaPar(); fechaLista(); tabela.push(l); continue; }
    fechaTabela();
    if (!l.trim()) { fechaPar(); fechaLista(); }
    else if ((m = l.match(/^(#{1,4})\s+(.*)/))) { fechaPar(); fechaLista(); const n = m[1].length + 2; html += "<h" + n + ">" + emLinha(m[2]) + "</h" + n + ">"; }
    else if (/^\s*(---+|\*\*\*+)\s*$/.test(l)) { fechaPar(); fechaLista(); html += "<hr>"; }
    else if ((m = l.match(/^\s*([-*•]|\d+[.)])\s+(.*)/))) {
      fechaPar();
      const tipo = /\d/.test(m[1]) ? "ol" : "ul";
      if (lista !== tipo) { fechaLista(); html += "<" + tipo + ">"; lista = tipo; }
      html += "<li>" + emLinha(m[2]) + "</li>";
    }
    else if ((m = l.match(/^>\s?(.*)/))) { fechaPar(); fechaLista(); html += "<blockquote>" + emLinha(m[1]) + "</blockquote>"; }
    else { fechaLista(); par.push(l); }
  }
  fechaPar(); fechaLista(); fechaTabela();
  return html;
}
function md(texto) {
  // as partes impares estao dentro de ``` (bloco de codigo), mesmo que o bloco ainda nao tenha fechado (resposta chegando)
  return texto.split("```").map((parte, i) => {
    if (i % 2 === 0) return mdTexto(parte);
    const quebra = parte.indexOf("\n");
    const linguagem = quebra < 0 ? "" : parte.slice(0, quebra).trim();
    const codigo = (quebra < 0 ? parte : parte.slice(quebra + 1)).replace(/\n$/, "");
    return '<div class="bloco"><div class="bloco-topo"><span>' + escapar(linguagem || "código") + '</span>'
         + '<button type="button" class="copiar">copiar</button></div><pre><code>' + escapar(codigo) + "</code></pre></div>";
  }).join("");
}

// ---- uma resposta da IA: texto formatado + botoes de copiar e refazer ---------------------------------------------
function falaIA(texto) {
  const d = document.createElement("div"); d.className = "fala";
  const rotulo = document.createElement("div"); rotulo.className = "p a"; rotulo.textContent = "studio>";
  const corpo = document.createElement("div"); corpo.className = "md";
  const acoes = document.createElement("div"); acoes.className = "acoes";
  d.append(rotulo, corpo, acoes); alvo.append(d);
  const f = {
    bruto: "",
    definir(t) { f.bruto = t; d.dataset.bruto = t; corpo.innerHTML = md(t); rolar(); },
    finalizar(podeRefazer) {
      acoes.innerHTML = '<button type="button" class="copiar-fala">copiar resposta</button>'
                      + (podeRefazer ? '<button type="button" class="refazer">refazer</button>' : "");
    },
  };
  if (texto !== undefined) f.definir(texto);
  return f;
}
async function copiarTexto(texto, botao) {
  try { await navigator.clipboard.writeText(texto); }
  catch {   // navegador sem permissao para a area de transferencia: o jeito antigo
    const a = document.createElement("textarea"); a.value = texto; document.body.append(a); a.select(); document.execCommand("copy"); a.remove();
  }
  const antes = botao.textContent; botao.textContent = "copiado ✔"; setTimeout(() => { botao.textContent = antes; }, 1400);
}
log.addEventListener("click", (e) => {
  const b = e.target.closest("button"); if (!b) return;
  if (b.classList.contains("copiar")) copiarTexto(b.closest(".bloco").querySelector("code").textContent, b);
  else if (b.classList.contains("copiar-fala")) copiarTexto(b.closest(".fala").dataset.bruto || "", b);
  else if (b.classList.contains("refazer")) refazer();
});

// ---- a conversa ----------------------------------------------------------------------------------------------------
// Cada pergunta abre um "turno": tudo o que a IA fizer para responde-la (pensamento, ferramentas, imagens, texto) fica dentro dele.
// Assim "refazer" consegue apagar a resposta inteira e pedir outra.
let ultimoTurno = null, conversaNoAr = null;
function salvarConversa() { try { localStorage.setItem("studio_conversa", JSON.stringify(historico.slice(-80))); } catch {} }
async function refazer() {
  if (ocupado || !ultimoTurno) return;
  historico.length = ultimoTurno.inicio; ultimoTurno.el.remove();
  travar(true);
  try { await conversar(null, true); } catch (err) { tarefa("").erro("erro: " + err.message); }
  travar(false);
}
$("parar").onclick = () => { if (conversaNoAr) fetch("/chat/" + conversaNoAr + "/parar", { method: "POST" }); };

async function conversar(texto, repetindo) {
  if (!info.chat) { linha("studio", "O chat com IA está desligado nesta sessão (campo CHAT_IA na célula do Colab). Enquanto isso, use /img <prompt> e /voz <texto>, ou os botões ao lado.", "a"); return; }
  if (!repetindo) historico.push({ role: "user", content: texto });
  log.querySelectorAll(".refazer").forEach(b => b.remove());   // so a ultima resposta pode ser refeita
  const turno = document.createElement("div"); turno.className = "turno"; log.append(turno);
  ultimoTurno = { el: turno, inicio: historico.length }; alvo = turno;
  try { await responder(); } finally { alvo = log; conversaNoAr = null; $("parar").style.display = "none"; salvarConversa(); }
}
async function responder() {
  const pensando = tarefa("pensando..."); pensando.pensando = true;
  let r = await fetch("/chat", { method: "POST", headers: { "Content-Type": "application/json" }, body: JSON.stringify({ mensagens: historico, pensar: pensarLigado }) });
  if (!r.ok) { if (historico.length && historico[historico.length - 1].role === "user") historico.pop(); pensando.erro(await erroDe(r)); return; }
  const { id } = await r.json();
  conversaNoAr = id; $("parar").style.display = "";
  let visto = 0, atual = pensando, parcial = null, pensamento = null; const falas = [];
  for (;;) {
    await dorme(300);
    r = await fetch("/chat/" + id + "?desde=" + visto);
    if (!r.ok) { if (atual) atual.erro(await erroDe(r)); return; }
    const j = await r.json();
    for (const ev of j.eventos) {
      visto++;
      if (atual) { atual.pensando ? atual.remover() : atual.ok(); atual = null; }
      if (ev.tipo === "pensamento") {   // o raciocinio da IA, aparecendo enquanto ela pensa
        if (!pensamento) { pensamento = document.createElement("div"); pensamento.className = "pensa"; alvo.append(pensamento); }
        pensamento.textContent = "💭 " + ev.texto; rolar();
      } else if (ev.tipo === "parcial") {   // a resposta sendo escrita
        if (!parcial) parcial = falaIA();
        parcial.definir(ev.texto);
      } else if (ev.tipo === "velocidade") {
        const v = document.createElement("div"); v.className = "vel";
        v.textContent = "⚡ " + (ev.tps ? ev.tps.toLocaleString("pt-BR") + " tokens/s · " : "") + ev.tokens + " tokens · " + ev.segundos.toLocaleString("pt-BR") + " s";
        alvo.append(v); rolar();
      } else if (ev.tipo === "texto") {
        if (!parcial) parcial = falaIA();
        parcial.definir(ev.texto); parcial.finalizar(true);
        parcial = pensamento = null; falas.push(ev.texto);
      }
      else if (ev.tipo === "ferramenta") atual = tarefa(ev.nome + " · " + ev.resumo);
      else if (ev.tipo === "imagem") mostrarImagem(ev.url);
      else if (ev.tipo === "audio") mostrarAudio(ev.url);
      else if (ev.tipo === "arquivo") mostrarArquivo(ev.url, ev.nome);
      else if (ev.tipo === "erro") tarefa("").erro(ev.texto);
    }
    if (j.fim) {
      if (atual) atual.pensando ? atual.remover() : atual.ok();
      historico = historico.concat(j.novas || []);
      if (falarRespostas && info.voz && falas.length) await falar({ texto: falas.join(" "), voz: $("v-voz").value || undefined }, true);
      return;
    }
    if (!atual && !parcial && !pensamento) { atual = tarefa("pensando..."); atual.pensando = true; }
  }
}

async function comando(texto) {
  if (texto === "/ajuda") {
    linha("studio", "Comandos:\n  /img <prompt>    gera uma imagem direto, sem passar pela IA\n  /voz <texto>     fala o texto com a voz padrão\n" + (info.video ? "  /video <prompt>  gera um vídeo curto (leva vários minutos)\n" : "") + (info.musica ? "  /musica <estilo> gera uma música instrumental de 30 s\n" : "") + (info.efeitos ? "  /efeito <som>    gera um efeito sonoro de 8 s\n" : "") + "  /limpar          apaga a conversa\nBotões: 🎤 fala no microfone em vez de digitar; 🔈 faz a IA responder falando.\nQualquer outra coisa vai para o chat com IA" + (info.chat ? "." : " (desligado nesta sessão)."), "a");
  } else if (texto === "/limpar") { log.innerHTML = ""; historico = []; ultimoTurno = null; salvarConversa(); }
  else if (texto.startsWith("/img ")) await gerarImagem({ prompt: texto.slice(5).trim() });
  else if (texto.startsWith("/musica ") && info.musica) await gerarSom("/musica", { prompt: texto.slice(8).trim() });
  else if (texto.startsWith("/efeito ") && info.efeitos) await gerarSom("/efeito", { prompt: texto.slice(8).trim() });
  else if (texto.startsWith("/video ") && info.video) await gerarImagem({ prompt: texto.slice(7).trim(), video: true });
  else if (texto.startsWith("/voz ")) await falar({ texto: texto.slice(5).trim() });
  else await conversar(texto);
}
$("form").onsubmit = async (e) => {
  e.preventDefault();
  const texto = $("msg").value.trim();
  if (!texto || ocupado) return;
  $("msg").value = ""; ajustarCaixa(); linha("você", texto); travar(true);
  try { await comando(texto); } catch (err) { tarefa("").erro("erro: " + err.message); }
  travar(false); $("msg").focus();
};
// caixa de mensagem: cresce com o texto; Enter envia e Shift+Enter quebra a linha (para colar codigo)
function ajustarCaixa() { const m = $("msg"); m.style.height = "auto"; m.style.height = Math.min(m.scrollHeight, 180) + "px"; }
$("msg").addEventListener("input", ajustarCaixa);
$("msg").addEventListener("keydown", (e) => { if (e.key === "Enter" && !e.shiftKey && !e.isComposing) { e.preventDefault(); $("form").requestSubmit(); } });
$("nova").onclick = () => { if (!ocupado) comando("/limpar"); };

// lote: pede tudo de uma vez e vai mostrando as imagens conforme ficam prontas
async function gerarLote(p) {
  const t = tarefa("lote · enviando...");
  try {
    let r = await fetch("/lote", { method: "POST", headers: { "Content-Type": "application/json" }, body: JSON.stringify(p) });
    if (!r.ok) throw new Error(await erroDe(r));
    const { id, total } = await r.json(); let mostradas = 0; const inicio = Date.now();
    for (;;) {
      await dorme(1500);
      r = await fetch("/lote/" + id);
      if (!r.ok) throw new Error(await erroDe(r));
      const j = await r.json();
      for (const arquivo of j.prontas.slice(mostradas)) { mostrarImagem("/imagem/" + encodeURIComponent(arquivo)); mostradas++; }
      const seg = Math.round((Date.now() - inicio) / 1000);
      if (j.fim) { j.erros.length ? t.erro("lote terminou com erro: " + j.erros[0]) : t.ok("lote pronto · " + mostradas + " imagens em " + seg + " s"); return; }
      t.texto("lote · " + mostradas + "/" + total + " imagens · " + seg + " s");
    }
  } catch (e) { t.erro("não consegui gerar o lote: " + e.message); }
}

// remover fundo: de uma imagem do chat (pelo nome) ou de um arquivo enviado
async function recortar(corpo, rotulo) {
  linha("você", "/recortar " + rotulo); travar(true);
  const t = tarefa("remover_fundo · " + rotulo);
  try {
    const json = !(corpo instanceof Blob);
    const r = await fetch("/remover-fundo", { method: "POST", headers: json ? { "Content-Type": "application/json" } : {}, body: json ? JSON.stringify(corpo) : corpo });
    if (!r.ok) throw new Error(await erroDe(r));
    const arquivo = r.headers.get("X-Arquivo"); await r.blob();
    t.ok("fundo removido"); mostrarImagem("/imagem/" + encodeURIComponent(arquivo));
  } catch (e) { t.erro("não consegui remover o fundo: " + e.message); }
  travar(false);
}
$("f-recortar").onclick = () => { const arq = $("f-arquivo").files[0]; if (arq && !ocupado) recortar(arq, arq.name); };

// imagem + video: o modo escolhe o que o botao faz; "animar" numa imagem do chat a usa como imagem inicial
let inicialGaleria = null;
function modoMudou() {
  const v = $("i-modo").value === "video";
  $("i-gerar").textContent = v ? "Gerar vídeo" : "Gerar imagem"; $("i-inicial-bloco").style.display = v ? "block" : "none";
}
$("i-modo").onchange = modoMudou;
$("i-inicial").onchange = () => { inicialGaleria = null; $("i-da-galeria").style.display = "none"; };
function animarDaGaleria(nome) {
  inicialGaleria = nome; $("i-inicial").value = ""; $("i-modo").value = "video"; modoMudou();
  $("i-da-galeria").textContent = "Usando " + nome + ". Descreva o movimento no prompt acima e clique em Gerar vídeo.";
  $("i-da-galeria").style.display = "block";
  document.querySelector('.aba[data-aba="imagem"]').click(); $("i-prompt").value = ""; $("i-prompt").focus();
}

$("i-gerar").onclick = async () => {
  const prompt = $("i-prompt").value.trim(); if (!prompt || ocupado) return;
  const [width, height] = $("i-formato").value.split("x").map(Number);
  const p = { prompt, width, height };
  if ($("i-passos").value) p.steps = Number($("i-passos").value);
  if ($("i-seed").value) p.seed = Number($("i-seed").value);
  const emVideo = $("i-modo").value === "video", inicial = emVideo ? $("i-inicial").files[0] : null;
  if (emVideo) p.video = true;
  // imagem -> video: manda o arquivo em base64 (ou o nome da imagem da galeria) e deixa o servidor seguir a proporcao dela
  if (inicial) p.image = await new Promise((ok, falha) => { const l = new FileReader(); l.onload = () => ok(l.result); l.onerror = falha; l.readAsDataURL(inicial); });
  else if (emVideo && inicialGaleria) p.image = inicialGaleria;
  if (p.image) { delete p.width; delete p.height; }
  // varias imagens de uma vez: a quantidade pedida, e no perfil de lote cada linha do campo e um prompt
  const qtd = Math.max(1, Math.min(64, Number($("i-qtd").value) || 1));
  const prompts = info.lote ? prompt.split("\n").map(x => x.trim()).filter(Boolean) : [prompt];
  if (!emVideo && (qtd > 1 || prompts.length > 1)) {
    linha("você", "/lote " + qtd + " × " + (prompts.length > 1 ? prompts.length + " prompts" : prompt));
    travar(true); await gerarLote({ prompts, quantidade: qtd, width, height, steps: p.steps, seed: p.seed }); travar(false);
    return;
  }
  const comando = !emVideo ? "/img " : p.image ? "/animar " + (inicial ? inicial.name : inicialGaleria) + " · " : "/video ";
  linha("você", comando + prompt); travar(true); await gerarImagem(p); travar(false);
};
$("v-falar").onclick = async () => {
  const texto = $("v-texto").value.trim(); if (!texto || ocupado) return;
  const p = { texto }; if ($("v-voz").value) p.voz = $("v-voz").value;
  linha("você", "/voz " + texto); travar(true); await falar(p); travar(false);
};
$("v-enviar").onclick = async () => {
  const nome = $("v-nome").value.trim(), arq = $("v-arquivo").files[0];
  if (!nome || !arq) { tarefa("").erro("escolha um nome e um arquivo de áudio"); return; }
  const t = tarefa("salvando a voz " + nome);
  const r = await fetch("/vozes/" + encodeURIComponent(nome), { method: "POST", body: arq });
  if (r.ok) { t.ok("voz " + nome + " salva"); vozes(); } else t.erro(await erroDe(r));
};

// conversa por voz: 🎤 grava e transcreve (Whisper); 🔈 faz a resposta da IA ser falada (Chatterbox)
let gravador = null, pedacos = [], falarRespostas = false, pensarLigado = false;
$("pensar").onclick = () => { pensarLigado = !pensarLigado; $("pensar").classList.toggle("on", pensarLigado); };
$("autovoz").onclick = () => {
  falarRespostas = !falarRespostas;
  $("autovoz").classList.toggle("on", falarRespostas); $("autovoz").textContent = falarRespostas ? "🔊" : "🔈";
};
$("mic").onclick = async () => {
  if (gravador) { gravador.stop(); return; }
  if (ocupado) return;
  let fluxo;
  try { fluxo = await navigator.mediaDevices.getUserMedia({ audio: true }); }
  catch (e) { tarefa("").erro("não consegui usar o microfone: " + e.message); return; }
  pedacos = []; gravador = new MediaRecorder(fluxo);
  gravador.ondataavailable = (e) => pedacos.push(e.data);
  gravador.onstop = async () => {
    fluxo.getTracks().forEach(t => t.stop());
    const blob = new Blob(pedacos, { type: gravador.mimeType }); gravador = null;
    $("mic").classList.remove("gravando"); $("mic").textContent = "🎤";
    const t = tarefa("ouvindo o que você disse..."); travar(true);
    try {
      const r = await fetch("/transcrever", { method: "POST", body: blob });
      if (!r.ok) throw new Error(await erroDe(r));
      const texto = ((await r.json()).texto || "").trim();
      if (!texto) t.erro("não entendi nada no áudio");
      else if (info.chat) { t.remover(); linha("você", "🎤 " + texto); await conversar(texto); }
      else { t.remover(); $("msg").value = texto; }
    } catch (e) { t.erro("não consegui transcrever: " + e.message); }
    travar(false);
  };
  gravador.start(); $("mic").classList.add("gravando"); $("mic").textContent = "■";
};

document.querySelectorAll(".aba").forEach(a => a.onclick = () => {
  document.querySelectorAll(".aba, .painel").forEach(e => e.classList.remove("ativa"));
  a.classList.add("ativa"); $("aba-" + a.dataset.aba).classList.add("ativa");
  if (a.dataset.aba === "galeria") galeria();
});
async function galeria() {
  try {
    const j = await (await fetch("/galeria")).json(); const g = $("grade"); g.innerHTML = "";
    for (const nome of j.imagens) { const i = document.createElement(/\.(webm|mp4)$/i.test(nome) ? "video" : "img"); i.loading = "lazy"; i.muted = true; i.src = "/imagem/" + encodeURIComponent(nome); i.onclick = () => window.open(i.src, "_blank"); g.append(i); }
    $("galeria-vazia").style.display = j.imagens.length ? "none" : "block";
  } catch {}
}
// som: musica (ACE-Step) e efeitos (Stable Audio). O mp3 volta na resposta ou, se demorar, e buscado em /resultado
async function gerarSom(rota, p) {
  const t = tarefa((rota === "/musica" ? "gerar_musica" : "gerar_efeito") + ' · "' + p.prompt.slice(0, 80) + '"');
  try {
    let r = await fetch(rota, { method: "POST", headers: { "Content-Type": "application/json" }, body: JSON.stringify(p) });
    while (r.status === 202) { const j = await r.json(); await dorme(1500); r = await fetch("/resultado/" + j.id); }
    if (!r.ok) throw new Error(await erroDe(r));
    const arquivo = r.headers.get("X-Arquivo"); await r.blob();
    t.ok("áudio pronto"); mostrarAudio("/imagem/" + encodeURIComponent(arquivo), true);
  } catch (e) { t.erro("não consegui gerar: " + e.message); }
}
function tipoDeSomMudou() {
  const musica = $("s-tipo").value === "/musica";
  $("s-rotulo").textContent = musica ? "Estilo (em inglês: gênero, instrumentos, clima)" : "O som (em inglês)";
  $("s-prompt").placeholder = musica ? "lofi hip hop, mellow piano, vinyl crackle, 80 bpm" : "heavy wooden door creaking open, then slamming shut";
  $("s-letra-bloco").style.display = musica ? "block" : "none"; $("s-segundos").placeholder = musica ? "30" : "8";
  $("s-gerar").textContent = musica ? "Gerar música" : "Gerar efeito";
}
$("s-tipo").onchange = tipoDeSomMudou;
$("s-gerar").onclick = async () => {
  const prompt = $("s-prompt").value.trim(); if (!prompt || ocupado) return;
  const rota = $("s-tipo").value, p = { prompt };
  if (rota === "/musica" && $("s-letra").value.trim()) p.letra = $("s-letra").value.trim();
  if ($("s-segundos").value) p.segundos = Number($("s-segundos").value);
  linha("você", (rota === "/musica" ? "/musica " : "/efeito ") + prompt); travar(true); await gerarSom(rota, p); travar(false);
};

// rodape: uso de RAM, GPU e disco da maquina do Colab, e quantos pedidos estao na fila
async function uso() {
  try {
    const j = await (await fetch("/status")).json();
    $("uso").textContent = "│ " + j.maquina.join(" │ ") + (j.fila ? " │ fila: " + j.fila : "");
  } catch { $("uso").textContent = "│ sem resposta do Colab"; }
}

async function vozes() {
  if (!info.voz) return;
  const j = await (await fetch("/vozes")).json(); const s = $("v-voz"); s.length = 1;
  for (const v of j.vozes) s.add(new Option(v, v));
}
async function iniciar() {
  info = await (await fetch("/info")).json();
  // o que esta ligado nesta sessao, com a logo de cada parte (vem pronto do servidor, igual ao painel do Colab)
  const modo = info.chat ? " · chat: " + info.chat_modelo : "";
  $("pilha-nota").textContent = (info.versao ? "v" + info.versao : "") + modo;
  if (info.webui) {   // ha um chat completo (Open WebUI) nesta sessao: atalho para ele
    const a = document.createElement("a"); a.href = info.webui; a.target = "_blank"; a.textContent = "abrir o chat completo (Open WebUI) ↗"; a.style.cssText = "margin-left:12px;color:var(--azul)";
    $("pilha-nota").append(a);
  }
  $("modelo").textContent = "│ " + info.perfil;
  uso(); setInterval(uso, 6000);
  $("voz-on").style.display = info.voz ? "block" : "none"; $("voz-off").style.display = info.voz ? "none" : "block";
  if (info.fundo) $("f-bloco").style.display = "block";
  if (info.musica) $("s-tipo").add(new Option("Música", "/musica"));
  if (info.efeitos) $("s-tipo").add(new Option("Efeito sonoro", "/efeito"));
  if (info.musica || info.efeitos) { $("aba-btn-som").style.display = ""; tipoDeSomMudou(); }
  $("i-passos").placeholder = info.passos;
  if (info.tipo === "video") { $("i-formato").value = "1280x720"; $("i-modo").value = "video"; }
  else if (info.video) $("i-modo-bloco").style.display = "block";
  modoMudou();
  $("mic").style.display = info.ouvir ? "" : "none"; $("autovoz").style.display = info.voz ? "" : "none"; $("pensar").style.display = info.chat ? "" : "none";
  if (info.tipo === "chat") { document.querySelector(".lado").style.display = "none"; document.body.classList.add("so-chat"); }   // perfil so de chat
  if (info.lote) { $("i-rotulo").textContent = "Prompts, um por linha (em inglês costuma sair melhor)"; $("i-qtd").placeholder = "4"; $("i-prompt").style.minHeight = "150px"; }
  linha("studio", (info.tipo === "chat" ? "Bem-vindo. Este perfil é só conversa: a IA tem a placa de vídeo inteira para ela.\n" : "Bem-vindo ao ComfyUI Studio. Modelo carregado: " + info.perfil + ".\n") + (info.chat ? "Me diga o que você quer criar, ou digite /ajuda." : "O chat com IA está desligado nesta sessão. Use /img <prompt>, /voz <texto> ou os botões ao lado. /ajuda mostra os comandos."), "a");
  // conversa guardada neste navegador (sobrevive a recarregar a pagina enquanto a sessao do Colab for a mesma)
  try {
    historico = JSON.parse(localStorage.getItem("studio_conversa") || "[]");
    for (const m of historico) {
      if (m.role === "user") linha("você", m.content);
      else if (m.role === "assistant" && m.content) falaIA(m.content).finalizar(false);
    }
  } catch { historico = []; }
  vozes(); galeria();
}
iniciar().catch(e => linha("studio", "não consegui falar com o servidor: " + e.message, "a"));
</script>
</body>
</html>
'''

# ------------------------------------------------------------------------------
# VISUAL
# ------------------------------------------------------------------------------

FONTE = "Cascadia Mono,Cascadia Code,Consolas,DejaVu Sans Mono,Menlo,monospace"  # sem aspas: vai dentro de style='...'
AZUL, VERDE, VERMELHO, CINZA, TEXTO, BORDA, CLARO = "#172dd7", "#1a7f37", "#cf222e", "#6e7781", "#1f2328", "#57606a", "#d0d7de"

# letras grandes estilo "ANSI Shadow"
LETRAS = {
    "C": [" ██████╗", "██╔════╝", "██║     ", "██║     ", "╚██████╗", " ╚═════╝"],
    "O": [" ██████╗ ", "██╔═══██╗", "██║   ██║", "██║   ██║", "╚██████╔╝", " ╚═════╝ "],
    "M": ["███╗   ███╗", "████╗ ████║", "██╔████╔██║", "██║╚██╔╝██║", "██║ ╚═╝ ██║", "╚═╝     ╚═╝"],
    "F": ["███████╗", "██╔════╝", "█████╗  ", "██╔══╝  ", "██║     ", "╚═╝     "],
    "Y": ["██╗   ██╗", "╚██╗ ██╔╝", " ╚████╔╝ ", "  ╚██╔╝  ", "   ██║   ", "   ╚═╝   "],
    "U": ["██╗   ██╗", "██║   ██║", "██║   ██║", "██║   ██║", "╚██████╔╝", " ╚═════╝ "],
    "I": ["██╗", "██║", "██║", "██║", "██║", "╚═╝"],
}
CORES_BANNER = ["#4a63ff", "#3350f5", "#2540e8", "#172dd7", "#1226b8", "#0e1e96"]


def montar_banner(palavra):
    linhas = []
    for i in range(6):
        linha = "".join(LETRAS[c][i].ljust(max(len(x) for x in LETRAS[c])) for c in palavra)
        linhas.append(f"<span style='color:{CORES_BANNER[i]}'>{linha}</span>")
    return (
        f"<pre style='margin:0;font-family:{FONTE};font-size:13px;line-height:1.15;white-space:pre;overflow:hidden;background:none;border:0;padding:0'>"
        + NL.join(linhas) + "</pre>"
    )


BANNER = montar_banner("COMFYUI")


def logo(org, tam=18):
    return (
        f"<img src='https://github.com/{org}.png?size=64' alt='' width='{tam}' height='{tam}' "
        f"style='border-radius:4px;vertical-align:middle;flex:none;border:1px solid #e5e7eb;background:#fff'>"
    )


ETAPAS = [
    ("gpu", "Verificar GPU", "googlecolab"),
    *([("comfy", "Baixar ComfyUI", "Comfy-Org"), ("deps", "Instalar dependências", "pypa")] if USA_COMFY else []),
    ("tunel", "Baixar Cloudflared", "cloudflare"),
    *[(f"m{i}", f"Modelo {i + 1}/{len(MODELOS)} · {m[3]}", m[4]) for i, m in enumerate(MODELOS)],
    *([("nodes", "Instalar nodes extras", "Kosinkadink")] if PERFIL_ATUAL.get("nodes") else []),
    *([("voz_inst", "Instalar Chatterbox (voz)", "resemble-ai"), ("voz_modelo", "Baixar modelo de voz pt-BR", "resemble-ai")] if VOZ else []),
    *([("chat_inst", "Instalar IA de chat (Ollama)", "ollama"), ("chat_modelo", f"Baixar {CHAT_ROTULO} (chat)", "QwenLM")] if CHAT_LOCAL else []),
    *([("arquivos", "Instalar criador de arquivos (Excel, Word, PDF, PowerPoint)", "python-openxml")] if ARQUIVOS else []),
    *([("webui", "Instalar Open WebUI (chat completo)", "open-webui")] if WEBUI else []),
    *([("ouvir", f"Instalar Whisper {WHISPER} (ouvir)", "openai")] if OUVIR else []),
    *([("fundo", "Instalar removedor de fundo (rembg)", "danielgatis")] if REMOVER_FUNDO else []),
    *([("start", "Iniciar ComfyUI", "Comfy-Org")] if USA_COMFY else []),
    *([("voz_start", "Iniciar servidor de voz", "resemble-ai")] if VOZ else []),
    ("api", "Iniciar API + MCP + Studio", "modelcontextprotocol"),
    ("link", "Criar links públicos", "cloudflare"),
]
# quem esta do outro lado da conexao em cada etapa
SERVICOS = [
    ("Google Colab", "googlecolab", ["gpu"]),
    ("GitHub", "github", (["comfy"] if USA_COMFY else []) + ["tunel"] + (["nodes"] if PERFIL_ATUAL.get("nodes") else [])),
    *([("PyPI", "pypa", ["deps"])] if USA_COMFY else []),
    *([("Hugging Face", "huggingface", [f"m{i}" for i in range(len(MODELOS))])] if MODELOS else []),
    *([("ComfyUI :8188", "Comfy-Org", ["start"])] if USA_COMFY else []),
    *([("Voz :8189", "resemble-ai", ["voz_inst", "voz_modelo", "voz_start"])] if VOZ else []),
    *([("Chat IA :11434", "ollama", ["chat_inst", "chat_modelo"])] if CHAT_LOCAL else []),
    *([("Open WebUI :8191", "open-webui", ["webui"])] if WEBUI else []),
    ("API + MCP :5000", "modelcontextprotocol", ["api"]),
    ("Cloudflare", "cloudflare", ["link"]),
]

estado = {
    k: {"nome": nome, "logo": org, "status": "espera", "t0": None, "dt": None, "info": "", "pct": None, "detalhe": ""}
    for k, nome, org in ETAPAS
}
ui = {"link": None, "link_api": None, "link_webui": None, "api_log": [], "link_colab": None, "erro": None, "erro_log": "", "parado": False, "online": None, "t_fim": None}
T0 = time.time()


class Falha(Exception):
    pass


def fmt_t(s):
    s = int(s)
    return f"{s // 60}:{s % 60:02d}"


def fmt_tam(b):
    return f"{b / 1e9:.2f} GB" if b >= 1e9 else f"{b / 1e6:.0f} MB"


def ultimas_linhas(n=12, caminho=None):
    try:
        with open(caminho or LOG, "rb") as f:
            f.seek(0, 2)
            f.seek(max(0, f.tell() - 6000))
            texto = f.read().decode("utf-8", "ignore")
    except OSError:
        return ""
    texto = re.sub(chr(27) + r"\[[0-9;]*[A-Za-z]", "", texto)  # tira os codigos de cor do terminal
    # barras de progresso reescrevem a mesma linha: fica so a versao mais recente
    linhas = [l.rstrip(CR).split(CR)[-1] for l in texto.split(NL)]
    return NL.join([l for l in linhas if l.strip()][-n:])


def barra_txt(pct, n=30, cor=AZUL):
    cheio = int(round(max(0, min(100, pct)) / 100 * n))
    return f"<span style='color:{cor}'>{'█' * cheio}</span><span style='color:{CLARO}'>{'░' * (n - cheio)}</span>"


# copia o texto do atributo data-c (no mousedown, para funcionar mesmo se o painel for redesenhado em seguida)
COPIAR_JS = (
    "var t=this.getAttribute('data-c'),a=document.createElement('textarea');a.value=t;document.body.appendChild(a);"
    "a.select();try{document.execCommand('copy')}catch(e){}a.remove();"
    "if(navigator.clipboard){navigator.clipboard.writeText(t).catch(function(){})}"
    "this.textContent='copiado ✔';event.preventDefault()"
)


SO_O_MAIS_NOVO_JS = (
    "var p=document.querySelectorAll('[data-cfy]'),m=0,i;"
    "for(i=0;i<p.length;i++)m=Math.max(m,+p[i].getAttribute('data-cfy'));"
    "for(i=0;i<p.length;i++)if(+p[i].getAttribute('data-cfy')<m)p[i].remove();this.remove()"
)


def linhas_de_conexao():
    """(logo, rotulo, endereco) do que da para abrir quando esta online. Rotas, exemplos e MCP ficam na pagina /docs."""
    api, linhas = ui["link_api"], []
    if ui["link_webui"]:
        linhas.append(("open-webui", "Chat (Open WebUI)", ui["link_webui"]))
    if api:
        linhas.append(("Comfy-Org", "Studio (criar)", com_chave(api + "/")))
    if ui["link"]:
        linhas.append(("Comfy-Org", "ComfyUI (editor)", ui["link"]))
    elif ui["link_colab"]:
        linhas.append(("googlecolab", "ComfyUI via Colab", ui["link_colab"]) if USA_COMFY
                      else ("googlecolab", "Studio via Colab", com_chave(ui["link_colab"] + "/")))
    if api:
        linhas.append(("modelcontextprotocol", "API + MCP (docs)", com_chave(api + "/docs")))
    return linhas


_metricas = {"quando": 0.0, "valores": []}


def metricas():
    """Uso da maquina para o rodape: [(texto, alerta)]. Medido no maximo a cada 4 s (o nvidia-smi nao e instantaneo)."""
    if time.time() - _metricas["quando"] < 4:
        return _metricas["valores"]
    valores = []
    try:
        with open("/proc/meminfo") as f:
            mem = {linha.split(":")[0]: int(linha.split()[1]) / 1048576 for linha in f if ":" in linha}
        usada = mem["MemTotal"] - mem["MemAvailable"]
        valores.append((f"RAM {usada:.1f}/{mem['MemTotal']:.1f} GB", usada / mem["MemTotal"] > 0.9))
    except (OSError, KeyError, ValueError, IndexError):
        pass
    try:
        smi = subprocess.run(
            ["nvidia-smi", "--query-gpu=memory.used,memory.total,utilization.gpu,temperature.gpu", "--format=csv,noheader,nounits"],
            capture_output=True, text=True, timeout=5,
        )
        usada, total, uso, temperatura = [float(x) for x in smi.stdout.strip().splitlines()[0].split(",")]
        valores.append((f"GPU {usada / 1024:.1f}/{total / 1024:.0f} GB · {uso:.0f}% · {temperatura:.0f}°C", usada / total > 0.97))
    except (OSError, ValueError, IndexError, subprocess.TimeoutExpired):
        pass
    try:
        livre = shutil.disk_usage("/content").free / 1e9
        valores.append((f"disco livre {livre:.0f} GB", livre < 5))
    except OSError:
        pass
    _metricas.update(quando=time.time(), valores=valores)
    return valores


def pilha_html():
    """O que esta sessao roda, numa linha: modelo de imagem + voz + ouvir + chat."""
    principal = [m for m in MODELOS_PERFIL if m[0] in ("diffusion_models", "checkpoints")]
    itens = [(principal[0][4], PERFIL_ATUAL["nome"])] if principal else []
    if MUSICA:
        itens.append(("ace-step", "ACE-Step"))
    if EFEITOS:
        itens.append(("Stability-AI", "Stable Audio"))
    if VOZ:
        itens.append(("resemble-ai", "Chatterbox pt-BR"))
    if OUVIR:
        itens.append(("openai", f"Whisper {WHISPER}"))
    if REMOVER_FUNDO:
        itens.append(("danielgatis", "rembg"))
    if CHAT_LOCAL:
        itens.append(("QwenLM", CHAT_ROTULO))
    separador = f" <span style='color:{CLARO}'>·</span> "
    return separador.join(f"<span style='white-space:nowrap'>{logo(org, 14)} {html.escape(nome)}</span>" for org, nome in itens)


def render():
    agora = time.time()
    spin = SPIN[int(agora * 10) % len(SPIN)]
    esc = html.escape
    feitas = sum(1 for e in estado.values() if e["status"] in ("ok", "aviso"))
    total = len(estado)
    rodando = next((e for e in estado.values() if e["status"] == "rodando"), None)
    decorrido = fmt_t((ui["t_fim"] or agora) - T0)
    ativo = not (ui["erro"] or ui["parado"])

    p = []
    # as animacoes recomecam a cada atualizacao; o delay negativo mantem todas no ritmo
    p.append(
        "<style>"
        "@keyframes cfyflow{from{background-position:0 0}to{background-position:16px 0}}"
        "@keyframes cfyblink{0%,49%{opacity:1}50%,100%{opacity:0}}"
        "@keyframes cfypulse{0%,100%{opacity:1}50%{opacity:.35}}"
        ".cfybtn{width:44px;text-align:center;line-height:30px;color:#1f2328;font-size:13px}"
        ".cfybtn:hover{background:#d0d7de}.cfyx:hover{background:#e81123;color:#fff}"
        "</style>"
        # se sobrar na tela uma copia antiga e congelada do painel (acontece quando o navegador reconecta ao Colab),
        # esta imagem vazia dispara o codigo que apaga todas as copias menos a mais nova
        f"<img src='data:,' alt='' style='display:none' onerror=\"{SO_O_MAIS_NOVO_JS}\">"
    )
    p.append(
        f"<div data-cfy='{int(agora * 1000)}' style='font-family:{FONTE};background:#ffffff;color:{TEXTO};border:1px solid {BORDA};"
        "max-width:920px;overflow:hidden;font-size:13px;line-height:1.7;box-shadow:0 6px 24px rgba(0,0,0,.18);margin:56px 0 360px'>"
    )

    # barra de titulo estilo Prompt de Comando
    p.append(
        f"<div style='background:#f3f4f6;display:flex;align-items:center;border-bottom:1px solid {BORDA};height:30px'>"
        f"<span style='padding:0 10px;display:flex;align-items:center'>{logo('Comfy-Org', 16)}</span>"
        f"<span style='flex:1;font-size:12px;color:{TEXTO};font-family:Segoe UI,system-ui,sans-serif'>"
        f"Prompt de Comando — comfyui-setup.exe · por {CRIADOR}</span>"
        "<span class='cfybtn'>─</span><span class='cfybtn'>☐</span><span class='cfybtn cfyx'>✕</span></div>"
    )
    p.append("<div style='padding:28px 28px 24px'>")

    # banner
    p.append(BANNER)
    p.append(
        f"<div style='margin:12px 0 20px'><span style='float:right;color:{CINZA}'>⏱ {decorrido}</span>{pilha_html()}</div>"
    )

    # conexoes (aparecem em destaque quando tudo fica pronto): uma linha por endereco, com botao de copiar
    conexoes = linhas_de_conexao() if ativo else []
    if conexoes:
        fase = agora % 1.6
        p.append(
            f"<div style='border:2px solid {VERDE};background:#f0fff4;padding:14px 18px;margin-bottom:20px'>"
            f"<div style='color:{VERDE};font-weight:bold;margin-bottom:2px'>"
            f"<span style='animation:cfypulse 1.6s infinite;animation-delay:-{fase:.2f}s'>●</span> "
            f"ONLINE <span style='font-weight:normal;color:{CINZA}'>— clique para abrir</span></div>"
        )
        for org, rotulo, texto in conexoes:
            valor = f"<a href='{esc(texto)}' target='_blank' style='color:{AZUL};font-weight:bold'>{esc(texto)}</a>"
            p.append(
                f"<div style='display:flex;align-items:center;gap:8px;line-height:1.9'>{logo(org, 16)}"
                f"<span style='width:138px;flex:none;color:{CINZA};white-space:nowrap'>{rotulo}</span>"
                f"<span style='flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis'>{valor}</span>"
                f"<span data-c=\"{esc(texto)}\" onmousedown=\"{COPIAR_JS}\" title='copiar' style='flex:none;cursor:pointer;border:1px solid {BORDA};"
                f"background:#fff;padding:0 8px;font-size:11px;line-height:1.7;white-space:nowrap;user-select:none'>copiar</span></div>"
            )
        if ui["link_webui"] and CHAVE_API:
            p.append(
                f"<div style='display:flex;align-items:center;gap:8px;line-height:1.9'><span style='width:18px;text-align:center'>👤</span>"
                f"<span style='width:138px;flex:none;color:{CINZA};white-space:nowrap'>Login do chat</span>"
                f"<span style='flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis'><b>{WEBUI_LOGIN}</b> "
                f"<span style='color:{CINZA}'>— a senha é a chave da API, logo abaixo</span></span>"
                f"<span data-c=\"{WEBUI_LOGIN}\" onmousedown=\"{COPIAR_JS}\" title='copiar' style='flex:none;cursor:pointer;border:1px solid {BORDA};"
                f"background:#fff;padding:0 8px;font-size:11px;line-height:1.7;white-space:nowrap;user-select:none'>copiar</span></div>"
            )
        if CHAVE_API and ui["link_api"]:
            p.append(
                f"<div style='display:flex;align-items:center;gap:8px;line-height:1.9'><span style='width:18px;text-align:center'>🔑</span>"
                f"<span style='width:138px;flex:none;color:{CINZA};white-space:nowrap'>Chave da API</span>"
                f"<span style='flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis'><b>{esc(CHAVE_API)}</b> "
                f"<span style='color:{CINZA}'>— já vai nos links do Studio e da documentação</span></span>"
                f"<span data-c=\"{esc(CHAVE_API)}\" onmousedown=\"{COPIAR_JS}\" title='copiar' style='flex:none;cursor:pointer;border:1px solid {BORDA};"
                f"background:#fff;padding:0 8px;font-size:11px;line-height:1.7;white-space:nowrap;user-select:none'>copiar</span></div>"
            )
        if MANUTENCAO:
            p.append(
                f"<div style='display:flex;align-items:center;gap:8px;line-height:1.9'><span style='width:18px;text-align:center'>🔧</span>"
                f"<span style='width:138px;flex:none;color:{CINZA};white-space:nowrap'>Manutenção (senha)</span>"
                f"<span style='flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis'><b>{esc(SENHA_MANUTENCAO)}</b> "
                f"<span style='color:#9a6700'>— quem tiver esta senha e o link controla este Colab</span></span>"
                f"<span data-c=\"{esc(SENHA_MANUTENCAO)}\" onmousedown=\"{COPIAR_JS}\" title='copiar' style='flex:none;cursor:pointer;border:1px solid {BORDA};"
                f"background:#fff;padding:0 8px;font-size:11px;line-height:1.7;white-space:nowrap;user-select:none'>copiar</span></div>"
            )
        avisos = []
        if not ui["link"]:
            avisos.append("o Cloudflare não devolveu o link do ComfyUI")
        if not ui["link_api"]:
            avisos.append("o Cloudflare não devolveu o link do Studio/API (rode de novo para tentar outra vez)")
        if avisos:
            p.append(f"<div style='font-size:12px;color:#9a6700'>⚠ {'; '.join(avisos)}.</div>")
        p.append(
            f"<div style='font-size:12px;color:{CINZA};margin-top:2px'>Deixe esta célula rodando; para parar, clique no ■ (stop). "
            "Os endereços mudam a cada execução. Rotas, exemplos e MCP estão na página de docs.</div></div>"
        )

    if ui["erro"]:
        p.append(
            f"<div style='border:2px solid {VERMELHO};background:#fff5f5;padding:10px 14px;margin-bottom:12px;color:{VERMELHO}'>"
            f"<b>✖ ERRO</b> — {esc(ui['erro'])}</div>"
        )
    if ui["parado"]:
        p.append(
            f"<div style='border:2px solid {BORDA};background:#f6f8fa;padding:10px 14px;margin-bottom:12px;color:{CINZA}'>"
            "<b>■ PARADO</b> — Rode a célula de novo para continuar. O que já foi baixado não é baixado outra vez.</div>"
        )

    # moldura com titulo na borda, duas colunas: conexoes | etapas
    p.append(
        f"<fieldset style='border:1px solid {BORDA};margin:0;padding:12px 18px 18px;min-width:0'>"
        f"<legend style='margin-left:auto;padding:0 8px;color:{AZUL};font-weight:bold;font-size:12px;width:auto;border:0'>"
        f"comfyui-setup v1.0 · {feitas}/{total} etapas</legend>"
        "<div style='display:flex;flex-wrap:wrap;gap:8px 28px'>"
    )

    # --- coluna 1: conexoes ---
    p.append(f"<div style='flex:1 1 300px;min-width:0'><div style='color:{AZUL};font-weight:bold'>Conexões</div>")
    fase_fluxo = agora % 0.5
    for nome, org, chaves in SERVICOS:
        sts = [estado[k]["status"] for k in chaves]
        com_aviso = "aviso" in sts
        no_ar = ui["online"] and ativo and nome in ("ComfyUI :8188", "Voz :8189", "Chat IA :11434", "Open WebUI :8191", "API + MCP :5000", "Cloudflare")
        if "erro" in sts or com_aviso:
            cor, rotulo, tipo = VERMELHO, "✖ falhou", "solido"
        elif "rodando" in sts:
            cor, rotulo, tipo = AZUL, f"{spin} transferindo", "fluxo"
        elif no_ar:
            cor, rotulo, tipo = VERDE, "● online", "fluxo"
        elif all(s == "ok" for s in sts):
            cor, rotulo, tipo = VERDE, "● conectado", "solido"
        elif "ok" in sts:
            cor, rotulo, tipo = VERDE, "● em espera", "solido"
        else:
            cor, rotulo, tipo = "#8c959f", "○ aguardando", "pontilhado"
        if tipo == "fluxo":
            fio = (
                f"<span style='flex:1;height:2px;background-image:linear-gradient(90deg,{cor} 50%,transparent 50%);"
                f"background-size:16px 2px;animation:cfyflow .5s linear infinite;animation-delay:-{fase_fluxo:.3f}s'></span>"
            )
        elif tipo == "solido":
            fio = f"<span style='flex:1;height:2px;background:{cor}'></span>"
        else:
            fio = f"<span style='flex:1;border-top:2px dotted {CLARO}'></span>"
        p.append(
            f"<div style='display:flex;align-items:center;gap:8px'>{logo(org)}"
            f"<span style='width:124px;white-space:nowrap'>{esc(nome)}</span>{fio}"
            f"<span style='color:{cor};width:118px;white-space:nowrap'>{rotulo}</span></div>"
        )
    p.append("</div>")

    # --- coluna 2: etapas ---
    p.append(f"<div style='flex:1.4 1 400px;min-width:0'><div style='color:{AZUL};font-weight:bold'>Instalação</div>")
    for e in estado.values():
        st = e["status"]
        if st == "ok":
            icone, cor, cor_nome, tempo = "✔", VERDE, TEXTO, fmt_t(e["dt"])
        elif st == "rodando":
            icone, cor, cor_nome, tempo = spin, AZUL, AZUL, fmt_t(agora - e["t0"])
        elif st == "erro":
            icone, cor, cor_nome, tempo = "✖", VERMELHO, VERMELHO, ""
        elif st == "aviso":
            icone, cor, cor_nome, tempo = "!", "#9a6700", "#9a6700", ""
        else:
            icone, cor, cor_nome, tempo = "·", "#8c959f", "#8c959f", ""
        extra = ""
        if st in ("ok", "aviso") and e["info"]:
            extra = f" <span style='color:{CINZA}'>{esc(e['info'][:40])}</span>"
        peso = "bold" if st == "rodando" else "normal"
        opac = "1" if st in ("ok", "rodando", "erro", "aviso") else ".45"
        p.append(
            f"<div style='display:flex;align-items:center;gap:8px'>"
            f"<span style='color:{cor};white-space:nowrap'>[{icone}]</span>"
            f"<span style='opacity:{opac};display:flex'>{logo(e['logo'], 16)}</span>"
            f"<span style='color:{cor_nome};font-weight:{peso};white-space:nowrap;overflow:hidden;text-overflow:ellipsis'>{esc(e['nome'])}{extra}</span>"
            f"<span style='flex:1;border-bottom:1px dotted {CLARO};min-width:12px;height:.8em'></span>"
            f"<span style='color:{CINZA}'>{tempo}</span></div>"
        )
        if st == "rodando":
            if e["pct"] is not None:
                p.append(
                    f"<div style='margin-left:36px;white-space:nowrap;overflow:hidden'>{barra_txt(e['pct'])} "
                    f"<b style='color:{AZUL}'>{e['pct']:.0f}%</b></div>"
                )
            if e["detalhe"]:
                p.append(f"<div style='margin-left:36px;color:{CINZA}'>{esc(e['detalhe'])}</div>")
    p.append("</div></div></fieldset>")

    # linha de comando: o que esta acontecendo agora + cursor piscando
    if ui["erro"]:
        acao = "instalação interrompida por erro"
    elif ui["parado"]:
        acao = "^C"
    elif ui["online"]:
        acao = "tudo no ar, aguardando você abrir o Studio ou o ComfyUI"
    elif rodando:
        acao = rodando["info"] or rodando["nome"].lower() + "..."
    else:
        acao = "preparando..."
    p.append(
        f"<div style='margin-top:10px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis'>"
        f"<span style='color:{CINZA}'>C:{BS}colab{BS}comfyui&gt;</span> {esc(acao[:120])}"
        f"<span style='color:{AZUL};animation:cfyblink 1s infinite;animation-delay:-{agora % 1:.2f}s'>█</span></div>"
    )

    # chamadas recebidas pela API / MCP
    if ui["api_log"] and ativo:
        p.append(
            f"<div style='margin-top:8px;color:{CINZA};font-size:12px'>── chamadas da API / MCP ──</div>"
            f"<pre style='margin:4px 0 0;background:#f0f4ff;border:1px solid {CLARO};color:#24292f;padding:8px 12px;"
            f"font-family:{FONTE};font-size:12px;line-height:1.5;white-space:pre-wrap;word-break:break-all'>"
            f"{esc(NL.join(ui['api_log'][-5:]))}</pre>"
        )

    # log ao vivo do ComfyUI (ou o motivo do erro)
    texto_log = ui["erro_log"] if ui["erro"] else (ultimas_linhas() if ui["online"] else "")
    if texto_log:
        rotulo = "saída do erro" if ui["erro"] else "atividade do ComfyUI (ao vivo)"
        p.append(
            f"<div style='margin-top:8px;color:{CINZA};font-size:12px'>── {rotulo} ──</div>"
            f"<pre style='margin:4px 0 0;background:#f6f8fa;border:1px solid {CLARO};color:#24292f;padding:8px 12px;"
            f"font-family:{FONTE};font-size:12px;line-height:1.5;white-space:pre-wrap;word-break:break-all;min-height:110px'>"
            f"{esc(texto_log)}</pre>"
        )
    p.append("</div>")

    # barra de status no rodape
    if ui["erro"]:
        situacao, cor_sit = "✖ erro", VERMELHO
    elif ui["parado"]:
        situacao, cor_sit = "■ parado", CINZA
    elif ui["online"]:
        situacao, cor_sit = f"● online há {int((agora - ui['online']) // 60)} min", VERDE
    else:
        situacao, cor_sit = f"{spin} instalando", AZUL
    pct_geral = feitas / total * 100
    barra = f"<span style='color:{CLARO}'>│</span>"
    # depois de online a barra de etapas so repetiria "100%": o espaco fica para os medidores
    progresso = "" if ui["online"] else (
        f"<span style='white-space:nowrap'>{feitas}/{total} [{barra_txt(pct_geral, 12)}] {pct_geral:.0f}%</span>{barra}"
    )
    medidores = "".join(
        f"{barra}<span style='color:{VERMELHO if alerta else CINZA};white-space:nowrap'>{esc(texto)}</span>" for texto, alerta in metricas()
    )
    p.append(
        f"<div style='background:#f3f4f6;border-top:1px solid {BORDA};padding:3px 14px;font-size:12px;display:flex;flex-wrap:wrap;gap:0 8px'>"
        f"<b style='color:{cor_sit};white-space:nowrap'>{situacao}</b>{barra}"
        f"{progresso}"
        f"<span style='white-space:nowrap'>⏱ {decorrido}</span>{medidores}"
        # o credito so cabe na mesma linha depois de online (durante a instalacao ele ja esta na barra de titulo)
        + (f"<span style='margin-left:auto;color:{CINZA};white-space:nowrap'>criado por {CRIADOR}</span>" if ui["online"] else "")
        + "</div>"
    )
    p.append("</div>")
    return "".join(p)


_ultima_pintura = [0.0]


def pintar(forcar=False):
    agora = time.time()
    if not forcar and agora - _ultima_pintura[0] < 0.2:
        return
    _ultima_pintura[0] = agora
    # atualiza o mesmo bloco de saida, sem apagar e redesenhar (apagar a cada atualizacao fazia o painel piscar)
    tela.update(HTML(render()))


def comecar(k, info=""):
    estado[k].update(status="rodando", t0=time.time(), info=info)
    pintar(True)


def concluir(k, info=""):
    e = estado[k]
    e.update(status="ok", dt=time.time() - e["t0"], info=info, pct=None, detalhe="")
    pintar(True)


def avisar(k, info):
    """Marca uma etapa opcional como nao concluida, sem parar a instalacao."""
    estado[k].update(status="aviso", info=info, pct=None, detalhe="")
    pintar(True)


def falhar(k, mensagem, log=""):
    estado[k].update(status="erro", info="")
    ui["erro"] = mensagem
    ui["erro_log"] = log
    raise Falha(mensagem)


# ------------------------------------------------------------------------------
# FERRAMENTAS
# ------------------------------------------------------------------------------

def rodar(k, cmd, cwd=None, com_pct=False, obrigatorio=True):
    """Roda um comando mostrando a ultima linha da saida no painel."""
    linhas = []
    proc = subprocess.Popen(cmd, cwd=cwd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, errors="ignore")

    def ler():
        for linha in proc.stdout:
            linha = re.sub(chr(27) + r"\[[0-9;?]*[A-Za-z]", "", linha)
            for parte in linha.replace(CR, NL).split(NL):
                if parte.strip():
                    linhas.append(parte.strip())

    leitor = threading.Thread(target=ler, daemon=True)
    leitor.start()
    while proc.poll() is None:
        if linhas:
            estado[k]["info"] = linhas[-1]
            if com_pct:
                m = re.search(r"(\d+)%", linhas[-1])
                if m:
                    estado[k]["pct"] = float(m.group(1))
        pintar()
        time.sleep(0.1)
    leitor.join(timeout=2)
    if proc.returncode != 0 and obrigatorio:
        falhar(k, "O comando falhou: " + " ".join(cmd[:4]) + " ...", NL.join(linhas[-15:]))
    return linhas


def _baixar(k, url, destino):
    parcial = destino + ".part"
    feito = os.path.getsize(parcial) if os.path.exists(parcial) else 0
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0", "Range": f"bytes={feito}-"})
    try:
        resp = urllib.request.urlopen(req, timeout=60)
    except urllib.error.HTTPError as ex:
        if ex.code == 416 and feito:
            # o arquivo parcial ja estava completo
            os.replace(parcial, destino)
            return
        raise
    if resp.status != 206:
        feito = 0
    total = feito + int(resp.headers.get("Content-Length") or 0)
    inicio, base = time.time(), feito
    e = estado[k]
    e["info"] = "baixando " + os.path.basename(destino)
    with open(parcial, "ab" if feito else "wb") as f:
        while True:
            bloco = resp.read(4 * 1024 * 1024)
            if not bloco:
                break
            f.write(bloco)
            feito += len(bloco)
            dt = time.time() - inicio
            vel = (feito - base) / dt if dt > 0 else 0
            if total:
                e["pct"] = feito / total * 100
                falta = fmt_t((total - feito) / vel) if vel > 0 else "--:--"
                e["detalhe"] = f"{fmt_tam(feito)} de {fmt_tam(total)} · {vel / 1e6:.0f} MB/s · falta {falta}"
            else:
                e["detalhe"] = f"{fmt_tam(feito)} · {vel / 1e6:.0f} MB/s"
            pintar()
    if total and feito < total:
        raise OSError("download incompleto")
    os.replace(parcial, destino)


def baixar(k, url, destino):
    """Download com barra de progresso; continua de onde parou se a conexao cair."""
    for tentativa in range(1, 6):
        try:
            _baixar(k, url, destino)
            return
        except OSError as ex:
            estado[k]["info"] = f"conexão caiu ({ex}), tentando de novo ({tentativa}/5)..."
            pintar(True)
            time.sleep(3)
    falhar(k, "Não consegui baixar " + os.path.basename(destino) + ". Rode a célula de novo para continuar de onde parou.")


def porta_aberta(porta=PORT):
    try:
        socket.create_connection(("127.0.0.1", porta), timeout=1).close()
        return True
    except OSError:
        return False


# ------------------------------------------------------------------------------
# API + MCP (para Claude Code, Codex e outros agentes gerarem imagens)
# ------------------------------------------------------------------------------

def registrar(texto):
    ui["api_log"].append(f"[{fmt_t(time.time() - T0)}] {texto}")
    del ui["api_log"][:-20]


def chamar_comfy(caminho, dados=None):
    req = urllib.request.Request(
        f"http://127.0.0.1:{PORT}{caminho}",
        data=json.dumps(dados).encode() if dados is not None else None,
        headers={"Content-Type": "application/json"},
    )
    return urllib.request.urlopen(req, timeout=60).read()


def montar_workflow(prompt, largura, altura, passos, seed, quadros=16, fps=16, imagem=None, chave=None, lote=1):
    """Grafo do ComfyUI (formato da API). `chave` escolhe o modelo; sem ela, vale o modelo de imagem do perfil."""
    chave = chave or PERFIL_ATUAL["chave"]
    arquivo = {pasta: nome for pasta, nome, *_ in (MODELOS_VIDEO if chave == "video" else PERFIL_ATUAL.get("modelos_imagem", MODELOS_PERFIL))}
    salvar = lambda origem: {"class_type": "SaveImage", "inputs": {"images": [origem, 0], "filename_prefix": "api"}}

    if chave == "video":
        # Wan 2.2 5B + LoRA FastWan: o mesmo modelo faz texto -> video e imagem -> video (a imagem vira o primeiro quadro).
        # Com a LoRA destilada, cfg 1 (o negativo deixa de pesar) e poucos passos. A saida e um .webm
        return {
            "unet": {"class_type": "UNETLoader", "inputs": {"unet_name": arquivo["diffusion_models"], "weight_dtype": "default"}},
            "rapido": {"class_type": "LoraLoaderModelOnly", "inputs": {"model": ["unet", 0], "lora_name": arquivo["loras"], "strength_model": 1.0}},
            "fluxo": {"class_type": "ModelSamplingSD3", "inputs": {"model": ["rapido", 0], "shift": 8.0}},
            "clip": {"class_type": "CLIPLoader", "inputs": {"clip_name": arquivo["text_encoders"], "type": "wan", "device": "default"}},
            "vae": {"class_type": "VAELoader", "inputs": {"vae_name": arquivo["vae"]}},
            "pos": {"class_type": "CLIPTextEncode", "inputs": {"clip": ["clip", 0], "text": prompt}},
            "neg": {"class_type": "CLIPTextEncode", "inputs": {"clip": ["clip", 0], "text": NEGATIVO_VIDEO}},
            **({"inicio": {"class_type": "LoadImage", "inputs": {"image": imagem}}} if imagem else {}),
            "latente": {"class_type": "Wan22ImageToVideoLatent", "inputs": {
                "vae": ["vae", 0], "width": largura, "height": altura, "length": quadros, "batch_size": 1,
                **({"start_image": ["inicio", 0]} if imagem else {})}},
            "gera": {"class_type": "KSampler", "inputs": {
                "model": ["fluxo", 0], "positive": ["pos", 0], "negative": ["neg", 0], "latent_image": ["latente", 0],
                "seed": seed, "steps": passos, "cfg": 1.0,
                "sampler_name": "euler", "scheduler": "simple", "denoise": 1.0,
            }},
            "decod": {"class_type": "VAEDecode", "inputs": {"samples": ["gera", 0], "vae": ["vae", 0]}},
            "salva": {"class_type": "SaveWEBM", "inputs": {
                "images": ["decod", 0], "filename_prefix": "api", "codec": "vp9", "fps": float(fps), "crf": 20.0}},
        }

    if chave == "sdxl":
        g = {"ckpt": {"class_type": "CheckpointLoaderSimple", "inputs": {"ckpt_name": arquivo["checkpoints"]}}}
        modelo, clip = ["ckpt", 0], ["ckpt", 1]
        if "loras" in arquivo:   # perfil com LoRA de estilo
            g["lora"] = {"class_type": "LoraLoader", "inputs": {
                "model": modelo, "clip": clip, "lora_name": arquivo["loras"], "strength_model": 1.0, "strength_clip": 1.0}}
            modelo, clip = ["lora", 0], ["lora", 1]
        return {
            **g,
            "pos": {"class_type": "CLIPTextEncode", "inputs": {"clip": clip, "text": PERFIL_ATUAL.get("prefixo", "") + prompt}},
            "neg": {"class_type": "CLIPTextEncode", "inputs": {"clip": clip, "text": PERFIL_ATUAL.get("negativo", "blurry, low quality, deformed")}},
            "latente": {"class_type": "EmptyLatentImage", "inputs": {"width": largura, "height": altura, "batch_size": lote}},
            "gera": {"class_type": "KSampler", "inputs": {
                "model": modelo, "positive": ["pos", 0], "negative": ["neg", 0], "latent_image": ["latente", 0],
                "seed": seed, "steps": passos, "cfg": PERFIL_ATUAL.get("cfg", 7.0),
                "sampler_name": PERFIL_ATUAL.get("sampler", "dpmpp_2m"), "scheduler": PERFIL_ATUAL.get("scheduler", "karras"), "denoise": 1.0,
            }},
            "decod": {"class_type": "VAEDecode", "inputs": {"samples": ["gera", 0], "vae": ["ckpt", 2]}},
            "salva": salvar("decod"),
        }

    # Z Image Turbo. Na T4 o peso vai para a placa em fp8 (PESO_FP8): o arquivo bf16 (12 GB) mais o text encoder (8 GB)
    # estouram a RAM de 12,7 GB do Colab gratis, e o notebook do Isi-dev recomenda exatamente isso para nao travar
    g = {
        "unet": {"class_type": "UNETLoader", "inputs": {"unet_name": arquivo["diffusion_models"],
                                                         "weight_dtype": "fp8_e4m3fn" if PESO_FP8 else "default"}},
        "clip": {"class_type": "CLIPLoader", "inputs": {"clip_name": arquivo["text_encoders"], "type": "lumina2", "device": "default"}},
        "vae": {"class_type": "VAELoader", "inputs": {"vae_name": arquivo["vae"]}},
        "pos": {"class_type": "CLIPTextEncode", "inputs": {"clip": ["clip", 0], "text": PERFIL_ATUAL.get("prefixo", "") + prompt}},
        "neg": {"class_type": "ConditioningZeroOut", "inputs": {"conditioning": ["pos", 0]}},
    }
    modelo = ["unet", 0]
    if "loras" in arquivo:   # LoRA de estilo dos packs de influenciador. No Z Image, empilhar mais de uma LoRA piora a imagem
        g["lora"] = {"class_type": "LoraLoaderModelOnly", "inputs": {"model": modelo, "lora_name": arquivo["loras"],
                                                                     "strength_model": PERFIL_ATUAL.get("lora_peso", 0.9)}}
        modelo = ["lora", 0]
    g.update({
        "latente": {"class_type": "EmptySD3LatentImage", "inputs": {"width": largura, "height": altura, "batch_size": lote}},
        "fluxo": {"class_type": "ModelSamplingAuraFlow", "inputs": {"model": modelo, "shift": 3.0}},
        "gera": {"class_type": "KSampler", "inputs": {
            "model": ["fluxo", 0], "positive": ["pos", 0], "negative": ["neg", 0], "latent_image": ["latente", 0],
            "seed": seed, "steps": passos, "cfg": 1.0,
            "sampler_name": "res_multistep", "scheduler": "simple", "denoise": 1.0,
        }},
    })
    g["decod"] = {"class_type": "VAEDecode", "inputs": {"samples": ["gera", 0], "vae": ["vae", 0]}}
    g["salva"] = salvar("decod")
    return g


def _inteiro(valor, padrao, minimo, maximo):
    try:
        return max(minimo, min(maximo, int(valor)))
    except (TypeError, ValueError):
        return padrao


def tipo_imagem(nome):
    return {"gif": "image/gif", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp", "webm": "video/webm", "mp4": "video/mp4",
            "mp3": "audio/mpeg", "flac": "audio/flac", "wav": "audio/wav"}.get(nome.rsplit(".", 1)[-1].lower(), "image/png")


def bytes_da_imagem(valor):
    """Imagem que chega pela API: o nome de um arquivo da galeria ou a imagem em base64 (com ou sem 'data:...;base64,')."""
    valor = str(valor)
    da_galeria = os.path.join(COMFY, "output", os.path.basename(valor)) if len(valor) < 300 else ""
    if da_galeria and os.path.isfile(da_galeria):
        with open(da_galeria, "rb") as f:
            return f.read()
    return base64.b64decode(valor.split(",", 1)[-1])


def remover_fundo(dados):
    """Bytes de uma imagem -> (bytes de um PNG com fundo transparente, nome do arquivo salvo na galeria)."""
    if not fundo_ligado:
        raise RuntimeError("o removedor de fundo está desligado nesta sessão (marque REMOVER_FUNDO na célula do Colab)")
    req = urllib.request.Request(f"http://127.0.0.1:{FUNDO_PORT}/recortar", data=dados, headers={"Content-Type": "application/octet-stream"})
    try:
        recorte = urllib.request.urlopen(req, timeout=300).read()
    except urllib.error.HTTPError as ex:
        try:
            motivo = json.loads(ex.read()).get("erro")
        except ValueError:
            motivo = None
        raise ValueError(motivo or "o removedor de fundo falhou")
    nome = f"recorte_{uuid.uuid4().hex[:10]}.png"
    os.makedirs(os.path.join(COMFY, "output"), exist_ok=True)
    with open(os.path.join(COMFY, "output", nome), "wb") as f:
        f.write(recorte)
    return recorte, nome


def imagem_de_entrada(valor):
    """Imagem inicial de um video, gravada em ComfyUI/input (de onde o node LoadImage le). Devolve (arquivo, (largura, altura))."""
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(bytes_da_imagem(valor))).convert("RGB")
    except Exception:
        raise ValueError("não consegui ler a imagem inicial (envie PNG ou JPG em base64, ou o nome de um arquivo da galeria)")
    nome = f"studio_{uuid.uuid4().hex[:10]}.png"
    os.makedirs(os.path.join(COMFY, "input"), exist_ok=True)
    img.save(os.path.join(COMFY, "input", nome))
    return nome, img.size


# No pack Hot (e em qualquer combinacao com a LoRA de realismo feita no Personalizar), pedidos que mencionam menores sao recusados
# e o prompt sempre descreve uma pessoa adulta. Em portugues e em ingles, porque o prompt pode vir nas duas linguas
MENORES = re.compile(r"\b(crian[cç]as?|menin[oa]s?|garot[oa]s?|novinh[oa]s?|adolescentes?|infantil|beb[eê]s?|menor(es)? de idade|"
                     r"([0-9]|1[0-7]) ?anos|child(ren)?|kids?|minors?|teens?|teenagers?|underage|young girl|young boy|little girl|little boy|"
                     r"loli|shota|schoolgirl|schoolboy|preteen|toddler|1[0-7] ?(y/?o|years? old))\b", re.IGNORECASE)
ADULTO_PREFIXO = "adult woman over 25 years old, fictional character, "


# em todos os packs: menor de idade junto com qualquer coisa sexual e sempre recusado (o modelo nao tem filtro proprio)
SEXUAL = re.compile(r"\b(sexy|sensual|nud[ea]s?|naked|nsfw|lingerie|biquini|bikini|topless|seminu[ad]?|pelad[oa]|er[oó]tic[oa]|erotic|"
                    r"porn\w*|sex\w*|provocante|seductive|lewd|hentai|ecchi|underwear|calcinha|suti[aã])\b", re.IGNORECASE)


def checar_prompt_adulto(prompt):
    if MENORES.search(prompt) and (ADULTO or SEXUAL.search(prompt)):
        raise ValueError("pedido recusado: este Studio não cria conteúdo envolvendo menores de idade"
                         + (" (o pack Hot só cria personagens adultos e fictícios)" if ADULTO else ""))


def enfileirar(prompt, largura=None, altura=None, passos=None, seed=0, extras=None):
    """Coloca uma imagem (ou video) na fila do ComfyUI. Devolve (id, seed)."""
    prompt = str(prompt or "").strip()
    if not prompt:
        raise ValueError("prompt vazio")
    checar_prompt_adulto(prompt)
    if ADULTO:
        prompt = ADULTO_PREFIXO + prompt
    preparar_gpu("imagem")
    extras = extras or {}
    if SO_CHAT:
        raise RuntimeError("este perfil é só de chat: não gera imagens (escolha outro perfil na célula do Colab)")
    if (extras.get("video") or extras.get("image")) and not TEM_VIDEO:
        raise ValueError("este perfil não gera vídeo (escolha um perfil com vídeo na célula do Colab)")
    # perfil so de video: tudo e video. Perfil combinado: video quando pedem ("video": true) ou mandam uma imagem inicial
    video = SO_VIDEO or (TEM_VIDEO and bool(extras.get("video") or extras.get("image")))
    imagem = None
    if extras.get("image"):
        imagem, tamanho = imagem_de_entrada(extras["image"])
        if largura is None and altura is None:
            largura, altura = tamanho   # sem tamanho pedido, o video segue a proporcao da imagem
    # sem tamanho pedido: quadrado para imagem, paisagem 480p para video (o formato em que o Wan foi treinado)
    largura, altura = _inteiro(largura, 832 if video else 1024, 256, 2048), _inteiro(altura, 480 if video else 1024, 256, 2048)
    if video:
        # 480p e o que a T4 aguenta: mantem a proporcao pedida, com o lado maior em 832. O VAE do Wan 2.2 pede multiplos de 32
        fator = 832 / max(largura, altura)
        largura, altura = max(256, round(largura * fator / 32) * 32), max(256, round(altura * fator / 32) * 32)
    else:
        largura, altura = largura // 16 * 16, altura // 16 * 16
    passos = _inteiro(passos, PASSOS_VIDEO if video else PERFIL_ATUAL["passos"], 1, 50)
    seed = _inteiro(seed, 0, 0, 2**63) or random.randint(1, 2**32 - 1)
    quadros = (_inteiro(extras.get("frames"), 49, 17, 121) - 1) // 4 * 4 + 1   # o Wan so aceita 4n+1 quadros; 49 = 2 s a 24 fps
    lote = 1 if video else _inteiro(extras.get("batch"), 1, 1, LOTE_MAXIMO)
    workflow = montar_workflow(prompt, largura, altura, passos, seed, quadros, _inteiro(extras.get("fps"), FPS_VIDEO, 8, 30), imagem,
                                "video" if video else None, lote)
    try:
        resposta = json.loads(chamar_comfy("/prompt", {"prompt": workflow}))
    except urllib.error.HTTPError as ex:
        raise RuntimeError("o ComfyUI recusou o pedido: " + ex.read().decode("utf-8", "ignore")[:300])
    return resposta["prompt_id"], seed


# ------------------------------------------------------------------------------
# MUSICA (ACE-Step) E EFEITOS SONOROS (Stable Audio Open): rodam no ComfyUI, como a imagem
# ------------------------------------------------------------------------------

def montar_workflow_audio(tipo, prompt, letra, segundos, passos, seed):
    """Grafo do ComfyUI para "musica" ou "efeito". A saida e um mp3 na pasta output."""
    salvar = {"class_type": "SaveAudioMP3", "inputs": {"audio": ["decod", 0], "filename_prefix": "api_" + tipo, "quality": "V0"}}
    if tipo == "musica":
        return {
            "ckpt": {"class_type": "CheckpointLoaderSimple", "inputs": {"ckpt_name": ARQ_MUSICA}},
            "fluxo": {"class_type": "ModelSamplingSD3", "inputs": {"model": ["ckpt", 0], "shift": 5.0}},
            # tags = estilo/instrumentos/clima; a letra pode ter marcacoes como [verse] e [chorus]
            "pos": {"class_type": "TextEncodeAceStepAudio", "inputs": {
                "clip": ["ckpt", 1], "tags": prompt, "lyrics": letra or "[instrumental]", "lyrics_strength": 0.99}},
            "neg": {"class_type": "ConditioningZeroOut", "inputs": {"conditioning": ["pos", 0]}},
            "latente": {"class_type": "EmptyAceStepLatentAudio", "inputs": {"seconds": float(segundos), "batch_size": 1}},
            "gera": {"class_type": "KSampler", "inputs": {
                "model": ["fluxo", 0], "positive": ["pos", 0], "negative": ["neg", 0], "latent_image": ["latente", 0],
                "seed": seed, "steps": passos, "cfg": 5.0, "sampler_name": "euler", "scheduler": "simple", "denoise": 1.0}},
            "decod": {"class_type": "VAEDecodeAudio", "inputs": {"samples": ["gera", 0], "vae": ["ckpt", 2]}},
            "salva": salvar,
        }
    return {
        "ckpt": {"class_type": "CheckpointLoaderSimple", "inputs": {"ckpt_name": ARQ_EFEITOS}},
        "clip": {"class_type": "CLIPLoader", "inputs": {"clip_name": ARQ_EFEITOS_T5, "type": "stable_audio", "device": "default"}},
        "texto": {"class_type": "CLIPTextEncode", "inputs": {"clip": ["clip", 0], "text": prompt}},
        "vazio": {"class_type": "CLIPTextEncode", "inputs": {"clip": ["clip", 0], "text": "low quality, noise, distortion"}},
        # o Stable Audio precisa saber a duracao tambem no condicionamento, nao so no latente
        "duracao": {"class_type": "ConditioningStableAudio", "inputs": {
            "positive": ["texto", 0], "negative": ["vazio", 0], "seconds_start": 0.0, "seconds_total": float(segundos)}},
        "latente": {"class_type": "EmptyLatentAudio", "inputs": {"seconds": float(segundos), "batch_size": 1}},
        "gera": {"class_type": "KSampler", "inputs": {
            "model": ["ckpt", 0], "positive": ["duracao", 0], "negative": ["duracao", 1], "latent_image": ["latente", 0],
            "seed": seed, "steps": passos, "cfg": 5.0, "sampler_name": "dpmpp_3m_sde_gpu", "scheduler": "exponential", "denoise": 1.0}},
        "decod": {"class_type": "VAEDecodeAudio", "inputs": {"samples": ["gera", 0], "vae": ["ckpt", 2]}},
        "salva": salvar,
    }


def enfileirar_audio(tipo, prompt, letra=None, segundos=None, passos=None, seed=0):
    """Coloca uma musica ou um efeito sonoro na fila do ComfyUI. Devolve (id, seed). O resultado sai pelo mesmo esperar()."""
    if not (MUSICA if tipo == "musica" else EFEITOS):
        campo = "MUSICA" if tipo == "musica" else "EFEITOS"
        raise RuntimeError(f"isto não foi ativado nesta sessão (marque {campo} na célula do Colab e rode de novo)")
    prompt = str(prompt or "").strip()
    if not prompt:
        raise ValueError("prompt vazio")
    preparar_gpu("imagem")
    if tipo == "musica":
        segundos = _inteiro(segundos, 30, 5, 240)
    else:
        segundos = _inteiro(segundos, 8, 1, 47)   # o Stable Audio Open foi treinado em clipes de ate 47 s
    passos = _inteiro(passos, 50, 10, 100)
    seed = _inteiro(seed, 0, 0, 2**63) or random.randint(1, 2**32 - 1)
    workflow = montar_workflow_audio(tipo, prompt, str(letra or "").strip(), segundos, passos, seed)
    try:
        resposta = json.loads(chamar_comfy("/prompt", {"prompt": workflow}))
    except urllib.error.HTTPError as ex:
        raise RuntimeError("o ComfyUI recusou o pedido: " + ex.read().decode("utf-8", "ignore")[:300])
    return resposta["prompt_id"], seed


MCP_TOOLS_AUDIO = {
    "musica": {
        "name": "generate_music",
        "description": (
            "Generate music with ACE-Step on ComfyUI (running on Google Colab). Returns a URL to the mp3. "
            "If it is still rendering after ~80s, returns a job id to pass to get_image."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "prompt": {"type": "string", "description": "Style tags in English: genre, instruments, mood, tempo. E.g. 'lofi hip hop, mellow piano, vinyl crackle, 80 bpm'."},
                "lyrics": {"type": "string", "description": "Optional lyrics, with section marks like [verse] and [chorus]. Omit for an instrumental track."},
                "seconds": {"type": "integer", "default": 30, "description": "5-240."},
                "seed": {"type": "integer", "default": 0, "description": "0 = random"},
            },
            "required": ["prompt"],
        },
    },
    "efeito": {
        "name": "generate_sound_effect",
        "description": (
            "Generate a sound effect or ambience with Stable Audio Open on ComfyUI (running on Google Colab). Returns a URL to the mp3. "
            "If it is still rendering after ~80s, returns a job id to pass to get_image."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "prompt": {"type": "string", "description": "The sound, in English. E.g. 'heavy wooden door creaking open, then slamming shut'."},
                "seconds": {"type": "integer", "default": 8, "description": "1-47."},
                "seed": {"type": "integer", "default": 0, "description": "0 = random"},
            },
            "required": ["prompt"],
        },
    },
}


lotes = {}   # id do lote -> {"pedidos": [ids do ComfyUI], "total": imagens pedidas}


def arquivos_do_pedido(pid):
    """Todos os arquivos que um pedido gerou (um lote devolve varios), ou None se ainda nao terminou."""
    item = json.loads(chamar_comfy("/history/" + pid)).get(pid)
    if not item:
        return None
    if item.get("status", {}).get("status_str") == "error":
        raise RuntimeError("o ComfyUI deu erro ao gerar (veja GET /log)")
    return [arq["filename"] for saida in item.get("outputs", {}).values() for lista in saida.values() if isinstance(lista, list)
            for arq in lista if isinstance(arq, dict) and "filename" in arq]


def iniciar_lote(dados):
    """Enfileira muitas imagens de uma vez. Cada prompt vira pedidos de ate LOTE_MAXIMO imagens, que o ComfyUI gera em sequencia."""
    if TEM_VIDEO and not PERFIL_ATUAL.get("modelos_imagem"):
        raise RuntimeError("lote só funciona em perfis com modelo de imagem")
    prompts = dados.get("prompts") if isinstance(dados.get("prompts"), list) else [dados.get("prompt")]
    prompts = [str(x).strip() for x in prompts if str(x or "").strip()]
    if not prompts:
        raise ValueError("mande 'prompts' (lista) ou 'prompt'")
    quantidade = _inteiro(dados.get("quantidade"), 4, 1, 64)
    if len(prompts) * quantidade > 128:
        raise ValueError("lote grande demais: no máximo 128 imagens por pedido")
    semente = _inteiro(dados.get("seed"), 0, 0, 2**62)
    pedidos = []
    for prompt in prompts:
        falta = quantidade
        while falta:
            agora = min(falta, LOTE_MAXIMO)
            pid, _ = enfileirar(prompt, dados.get("width"), dados.get("height"), dados.get("steps"),
                                semente + len(pedidos) if semente else 0, {"batch": agora})
            pedidos.append(pid)
            falta -= agora
    lid = uuid.uuid4().hex[:12]
    lotes[lid] = {"pedidos": pedidos, "total": len(prompts) * quantidade}
    return lid


def andamento_do_lote(lid):
    lote = lotes.get(lid)
    if not lote:
        raise ValueError("lote não encontrado")
    prontas, pendentes, erros = [], 0, []
    for pid in lote["pedidos"]:
        try:
            arquivos = arquivos_do_pedido(pid)
        except RuntimeError as ex:
            erros.append(str(ex))
            continue
        if arquivos is None:
            pendentes += 1
        else:
            prontas += arquivos
    return {"total": lote["total"], "prontas": prontas, "imagens": [com_chave(f"{base_api()}/imagem/{urllib.parse.quote(a)}") for a in prontas],
            "fim": pendentes == 0, "erros": erros}


def validar_id(pid):
    pid = str(pid or "")
    if not re.fullmatch(r"[0-9a-fA-F-]{8,64}", pid):
        raise ValueError("id inválido")
    return pid


def esperar(pid, limite=80):
    """Espera a imagem ficar pronta. Devolve (bytes_png, arquivo) ou None se ainda estiver gerando.

    O limite existe porque o tunel do Cloudflare corta respostas que demoram mais de ~100s.
    """
    fim = time.time() + limite
    while time.time() < fim:
        item = json.loads(chamar_comfy("/history/" + pid)).get(pid)
        if item:
            if item.get("status", {}).get("status_str") == "error":
                raise RuntimeError("o ComfyUI deu erro ao gerar (veja a atividade no painel do Colab)")
            # imagens e videos chegam em listas de {filename, subfolder, type}, cada node com o seu nome de chave
            imagens = [arq for saida in item.get("outputs", {}).values() for lista in saida.values() if isinstance(lista, list)
                       for arq in lista if isinstance(arq, dict) and "filename" in arq]
            if not imagens:
                raise RuntimeError("a geração terminou sem imagem")
            url = f"http://127.0.0.1:{PORT}/view?" + urllib.parse.urlencode(imagens[0])
            return urllib.request.urlopen(url, timeout=60).read(), imagens[0]["filename"]
        time.sleep(1)
    return None


def base_api():
    return ui["link_api"] or f"http://127.0.0.1:{API_PORT}"


def com_chave(url):
    """Acrescenta a chave da API ao endereco, para abrir direto no navegador ou num cliente que nao manda cabecalhos."""
    return url + ("&" if "?" in url else "?") + "chave=" + CHAVE_API if CHAVE_API else url


def miniatura(dados):
    """JPEG reduzido em base64, para a resposta do MCP nao ficar pesada."""
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(dados)).convert("RGB")
        img.thumbnail((768, 768))
        saida = io.BytesIO()
        img.save(saida, "JPEG", quality=85)
        return base64.b64encode(saida.getvalue()).decode()
    except Exception:
        return None


# ------------------------------------------------------------------------------
# VOZ (Chatterbox pt-BR, opcional)
# ------------------------------------------------------------------------------

trabalhos_voz = {}
trava_voz = threading.Lock()


def preparar_gpu(quem):
    """A T4 nao segura imagem, voz e chat juntos com folga: quem vai usar pede para os outros soltarem a memoria.

    A IA de chat fica carregada entre uma resposta e outra (recarregar a cada mensagem deixava o chat lento)
    e so e descarregada aqui, quando a imagem ou a voz vao usar a placa.
    """
    if quem != "chat" and CHAT_LIGADO and CHAT_MODO == "reveza":
        try:
            # todos os que o Ollama tiver carregado (pelo Open WebUI da para conversar com outro modelo alem do escolhido)
            carregados = json.loads(urllib.request.urlopen(f"http://127.0.0.1:{OLLAMA_PORT}/api/ps", timeout=20).read()).get("models", [])
            for modelo in [m.get("name") for m in carregados] or [CHAT_MODELO]:
                pedido = json.dumps({"model": modelo, "keep_alive": 0}).encode()   # keep_alive 0 = descarregar agora
                urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{OLLAMA_PORT}/api/generate", data=pedido,
                                                              headers={"Content-Type": "application/json"}), timeout=60).read()
        except (OSError, ValueError):
            pass
    if quem != "imagem":
        try:
            chamar_comfy("/free", {"unload_models": True, "free_memory": True})
        except OSError:
            pass
    if VOZ and quem != "voz":
        try:
            urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{VOZ_PORT}/liberar", data=b"{}"), timeout=60).read()
        except OSError:
            pass


def nome_seguro(nome):
    return re.sub(r"[^a-zA-Z0-9_-]", "", str(nome or ""))[:40]


def listar_vozes():
    if not os.path.isdir(VOZ_REFS):
        return []
    return sorted({os.path.splitext(f)[0] for f in os.listdir(VOZ_REFS)})


def caminho_voz(nome):
    for f in sorted(os.listdir(VOZ_REFS)) if os.path.isdir(VOZ_REFS) else []:
        if os.path.splitext(f)[0] == nome:
            return os.path.join(VOZ_REFS, f)
    raise ValueError(f"a voz '{nome}' não existe. Vozes salvas: {', '.join(listar_vozes()) or 'nenhuma'}")


def salvar_voz(nome, dados):
    """Guarda um audio de referencia (10 a 15 s de fala limpa) para clonar depois pelo nome."""
    nome = nome_seguro(nome)
    if not nome:
        raise ValueError("nome de voz inválido (use letras, números, - e _)")
    if dados[:4] == b"RIFF":
        ext = ".wav"
    elif dados[:4] == b"fLaC":
        ext = ".flac"
    elif dados[:4] == b"OggS":
        ext = ".ogg"
    elif dados[4:8] == b"ftyp":
        ext = ".m4a"
    elif dados[:3] == b"ID3" or (len(dados) > 2 and dados[0] == 0xFF and dados[1] & 0xE0 == 0xE0):
        ext = ".mp3"
    else:
        raise ValueError("formato de áudio não reconhecido (envie wav, mp3, flac, ogg ou m4a)")
    os.makedirs(VOZ_REFS, exist_ok=True)
    for f in os.listdir(VOZ_REFS):
        if os.path.splitext(f)[0] == nome:
            os.remove(os.path.join(VOZ_REFS, f))
    with open(os.path.join(VOZ_REFS, nome + ext), "wb") as f:
        f.write(dados)
    return nome


def iniciar_fala(texto, voz=None, exaggeration=None, cfg_weight=None):
    """Comeca a gerar a fala em segundo plano. Devolve o id para esperar_fala."""
    if not VOZ:
        raise RuntimeError("a voz não foi ativada nesta sessão (marque VOZ na célula do Colab e rode de novo)")
    texto = str(texto or "").strip()
    if not texto:
        raise ValueError("texto vazio")
    if len(texto) > 5000:
        raise ValueError("texto longo demais (máximo de 5000 caracteres por chamada)")
    corpo = {"texto": texto, "lingua": "pt", "exaggeration": float(exaggeration or 0.5), "cfg_weight": float(cfg_weight or 0.5)}
    if voz:
        corpo["referencia"] = caminho_voz(nome_seguro(voz))
    jid = uuid.uuid4().hex[:12]
    os.makedirs(VOZ_SAIDAS, exist_ok=True)
    trabalho = {"pronto": threading.Event(), "erro": None, "arquivo": os.path.join(VOZ_SAIDAS, jid + ".wav")}
    trabalhos_voz[jid] = trabalho

    def rodar_fala():
        inicio = time.time()
        try:
            with trava_voz:
                preparar_gpu("voz")
                req = urllib.request.Request(f"http://127.0.0.1:{VOZ_PORT}/falar", data=json.dumps(corpo).encode(),
                                             headers={"Content-Type": "application/json"})
                dados = urllib.request.urlopen(req, timeout=1800).read()
            with open(trabalho["arquivo"], "wb") as f:
                f.write(dados)
            registrar(f"fala pronta em {time.time() - inicio:.0f}s · {jid}.wav")
        except urllib.error.HTTPError as ex:
            trabalho["erro"] = ex.read().decode("utf-8", "ignore")[:300]
        except Exception as ex:
            trabalho["erro"] = str(ex)
        if trabalho["erro"]:
            registrar(f"erro na fala · {trabalho['erro'][:80]}")
        trabalho["pronto"].set()

    threading.Thread(target=rodar_fala, daemon=True).start()
    return jid


def esperar_fala(jid, limite=80):
    """Caminho do WAV, ou None se ainda estiver gerando (mesmo limite de tempo do tunel que nas imagens)."""
    trabalho = trabalhos_voz.get(jid)
    if not trabalho:
        raise ValueError("id de fala desconhecido")
    if not trabalho["pronto"].wait(limite):
        return None
    if trabalho["erro"]:
        raise RuntimeError(trabalho["erro"])
    return trabalho["arquivo"]


def resultado_fala_mcp(jid):
    arquivo = esperar_fala(jid)
    if arquivo is None:
        return {"content": [{"type": "text", "text": f"Still rendering. Call get_audio with id={jid} in a few seconds."}]}
    return {"content": [{"type": "text", "text": "Audio ready (WAV, 24 kHz): " + com_chave(f"{base_api()}/audio/{jid}.wav")}]}


# ------------------------------------------------------------------------------
# ARQUIVOS: a IA escreve o conteudo e o Studio monta o arquivo (Excel, Word, PDF, PowerPoint, CSV, Markdown, texto)
# ------------------------------------------------------------------------------
# Tudo na CPU, com bibliotecas pequenas (openpyxl, python-docx, python-pptx, fpdf2): nao pesa na placa.
# Os arquivos vao para a galeria (e para o Drive, com SALVAR_NO_DRIVE) e voltam como link de download.

FORMATOS_ARQUIVO = {
    "xlsx": "planilha do Excel", "csv": "planilha CSV", "docx": "documento do Word", "pdf": "PDF",
    "pptx": "apresentação do PowerPoint", "md": "texto em Markdown (README)", "txt": "texto simples",
}
TIPOS_ARQUIVO = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8",
}
AZUL_DOC = "172DD7"


def _nome_arquivo(titulo, ext):
    import unicodedata
    sem_acento = unicodedata.normalize("NFKD", html.unescape(str(titulo or "arquivo"))).encode("ascii", "ignore").decode()   # "mês" -> "mes"
    base = re.sub(r"[^a-zA-Z0-9]+", "_", sem_acento).strip("_")[:40] or "arquivo"
    return f"{base.lower()}_{uuid.uuid4().hex[:6]}.{ext}"


def _linhas_markdown(texto):
    """Markdown simples -> [(tipo, nivel, texto)]: titulo (# a ###), item (- ou 1.), paragrafo. O bastante para Word e PDF."""
    blocos = []
    for linha in str(texto or "").replace("\r", "").split("\n"):
        crua = linha.strip()
        if not crua:
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", crua)
        if m:
            blocos.append(("titulo", len(m.group(1)), m.group(2)))
        elif re.match(r"^[-*•]\s+", crua):
            blocos.append(("item", 0, re.sub(r"^[-*•]\s+", "", crua)))
        elif re.match(r"^\d+[.)]\s+", crua):
            blocos.append(("numero", 0, crua))
        else:
            blocos.append(("paragrafo", 0, crua))
    return blocos


def _pedacos_negrito(texto):
    """'um **dois** tres' -> [('um ', False), ('dois', True), (' tres', False)]."""
    partes = re.split(r"(\*\*[^*]+\*\*)", texto)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in partes if p]


def _tabela(colunas, linhas):
    colunas = [str(c) for c in (colunas or [])]
    linhas = [list(l) if isinstance(l, (list, tuple)) else [l] for l in (linhas or [])]
    if not colunas and linhas:
        colunas, linhas = [str(c) for c in linhas[0]], linhas[1:]
    if not colunas:
        raise ValueError("para planilha, mande 'colunas' (os títulos) e 'linhas' (uma lista por linha)")

    def numero(v):   # "1.234,50" e "R$ 10" viram numero de verdade, para dar para somar no Excel
        if isinstance(v, (int, float)) or v is None:
            return v
        t = re.sub(r"^(R\$|US\$|\$)\s*", "", str(v).strip())
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?", t):
            return float(t.replace(".", "").replace(",", "."))
        if re.fullmatch(r"-?\d+([.,]\d+)?", t):
            return float(t.replace(",", "."))
        return v
    return colunas, [[numero(v) for v in l] + [None] * (len(colunas) - len(l)) for l in linhas]


def _celula_texto(v):
    """Valor de tabela como texto, no padrao brasileiro: 9.5 -> '9,5', 1500.0 -> '1500'."""
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.2f}".rstrip("0").replace(".", ",")
    return str(v)


def _xlsx(caminho, titulo, colunas, linhas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    colunas, linhas = _tabela(colunas, linhas)
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", str(titulo))[:31] or "Dados"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colunas))
    c = ws.cell(row=1, column=1, value=str(titulo))
    c.font, c.fill = Font(size=14, bold=True, color="FFFFFF"), PatternFill("solid", start_color=AZUL_DOC)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    borda = Border(*(Side(style="thin", color="D0D7DE"),) * 4)
    for j, nome in enumerate(colunas, 1):
        c = ws.cell(row=2, column=j, value=nome)
        c.font, c.fill, c.border = Font(bold=True), PatternFill("solid", start_color="E8EBFF"), borda
        c.alignment = Alignment(horizontal="center")
    for i, linha in enumerate(linhas, 3):
        for j, valor in enumerate(linha[:len(colunas)], 1):
            c = ws.cell(row=i, column=j, value=valor)
            c.border = borda
            if isinstance(valor, float) and not valor.is_integer():
                c.number_format = "#,##0.00"
    # linha de total nas colunas numericas
    numericas = [j for j in range(1, len(colunas) + 1) if linhas and all(isinstance(l[j - 1], (int, float)) or l[j - 1] is None for l in linhas)
                 and any(isinstance(l[j - 1], (int, float)) for l in linhas)]
    if numericas and len(linhas) > 1:
        fim = len(linhas) + 2
        ws.cell(row=fim + 1, column=1, value="Total").font = Font(bold=True)
        for j in numericas:
            if j == 1:
                continue
            letra = get_column_letter(j)
            c = ws.cell(row=fim + 1, column=j, value=f"=SUM({letra}3:{letra}{fim})")
            c.font, c.number_format = Font(bold=True), "#,##0.00"
    for j, coluna in enumerate(ws.columns, 1):
        largura = max(len(str(c.value or "")) for c in list(coluna)[1:]) if ws.max_row > 1 else 10
        ws.column_dimensions[get_column_letter(j)].width = min(60, max(12, largura + 3))
    ws.freeze_panes = "A3"
    wb.save(caminho)


def _csv(caminho, colunas, linhas):
    import csv
    colunas, linhas = _tabela(colunas, linhas)
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: o Excel abre os acentos certo
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(colunas)
        escritor.writerows([["" if v is None else str(v).replace(".", ",") if isinstance(v, float) else v for v in l] for l in linhas])


def _docx(caminho, titulo, texto, colunas, linhas):
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    cabecalho = doc.add_heading(str(titulo), level=0)
    for r in cabecalho.runs:
        r.font.color.rgb = RGBColor.from_string(AZUL_DOC)
    for tipo, nivel, conteudo in _linhas_markdown(texto):
        if tipo == "titulo":
            doc.add_heading(conteudo.replace("**", ""), level=nivel)
            continue
        estilo = {"item": "List Bullet", "numero": "List Number"}.get(tipo)
        if tipo == "numero":
            conteudo = re.sub(r"^\d+[.)]\s+", "", conteudo)
        p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
        for pedaco, negrito in _pedacos_negrito(conteudo):
            p.add_run(pedaco).bold = negrito
    if colunas or linhas:
        colunas, linhas = _tabela(colunas, linhas)
        tabela = doc.add_table(rows=1, cols=len(colunas))
        tabela.style = "Light Grid Accent 1"
        for c, nome in zip(tabela.rows[0].cells, colunas):
            c.text = nome
        for linha in linhas:
            for c, valor in zip(tabela.add_row().cells, linha):
                c.text = _celula_texto(valor)
    doc.save(caminho)


def _fonte_pdf():
    """Uma fonte com acentos e simbolos. No Colab sempre ha a DejaVu (do sistema ou do matplotlib)."""
    candidatos = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"]
    with contextlib.suppress(Exception):
        import importlib.util
        especificacao = importlib.util.find_spec("matplotlib")
        if especificacao and especificacao.submodule_search_locations:
            candidatos.insert(1, os.path.join(list(especificacao.submodule_search_locations)[0], "mpl-data", "fonts", "ttf", "DejaVuSans.ttf"))
    for c in candidatos:
        if os.path.isfile(c):
            negrito = c.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf").replace("arial.ttf", "arialbd.ttf")
            return c, negrito if os.path.isfile(negrito) else c
    return None, None


def _pdf(caminho, titulo, texto, colunas, linhas):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    normal, negrito = _fonte_pdf()
    if normal:
        pdf.add_font("Texto", "", normal)
        pdf.add_font("Texto", "B", negrito)
        familia = "Texto"
    else:   # sem fonte com acentos: Helvetica, trocando o que ela nao tem
        familia = "Helvetica"
    limpar = (lambda t: t) if normal else (lambda t: t.encode("latin-1", "replace").decode("latin-1"))
    pdf.set_font(familia, "B", 20)
    pdf.set_text_color(23, 45, 215)
    pdf.multi_cell(0, 10, limpar(str(titulo)), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_text_color(31, 35, 40)
    for tipo, nivel, conteudo in _linhas_markdown(texto):
        if tipo == "titulo":
            pdf.ln(2)
            pdf.set_font(familia, "B", {1: 16, 2: 14}.get(nivel, 12))
            pdf.multi_cell(0, 8, limpar(conteudo.replace("**", "")), new_x="LMARGIN", new_y="NEXT")
            continue
        pdf.set_font(familia, "", 11)
        prefixo = "•  " if tipo == "item" else ""
        pdf.multi_cell(0, 6, limpar(prefixo + conteudo), markdown=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)
    if colunas or linhas:
        colunas, linhas = _tabela(colunas, linhas)
        pdf.ln(3)
        pdf.set_font(familia, "", 10)
        with pdf.table(text_align="LEFT") as tabela:
            for linha in [colunas] + [[_celula_texto(v) for v in l] for l in linhas]:
                r = tabela.row()
                for v in linha:
                    r.cell(limpar(str(v)))
    pdf.output(caminho)


def _pptx(caminho, titulo, slides):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    if not slides:
        raise ValueError("para apresentação, mande 'slides': uma lista de {titulo, topicos}")
    prs = Presentation()
    capa = prs.slides.add_slide(prs.slide_layouts[0])
    capa.shapes.title.text = str(titulo)
    capa.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(AZUL_DOC)
    capa.placeholders[1].text = "Feito com o ComfyUI Studio"
    for s in slides:
        s = s if isinstance(s, dict) else {"titulo": str(s), "topicos": []}
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(s.get("titulo") or s.get("title") or "")
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(AZUL_DOC)
        corpo = slide.placeholders[1].text_frame
        corpo.word_wrap = True
        topicos = s.get("topicos") or s.get("bullets") or s.get("pontos") or []
        for i, t in enumerate(topicos if isinstance(topicos, list) else [topicos]):
            p = corpo.paragraphs[0] if i == 0 else corpo.add_paragraph()
            p.text = str(t)
            p.font.size = Pt(20)
        if s.get("notas"):
            slide.notes_slide.notes_text_frame.text = str(s["notas"])
    prs.save(caminho)


def criar_arquivo(dados):
    """{formato, titulo, texto?, colunas?, linhas?, slides?} -> (nome do arquivo na galeria, url para baixar)."""
    formato = str(dados.get("formato") or "").lower().lstrip(".")
    formato = {"excel": "xlsx", "planilha": "xlsx", "word": "docx", "documento": "docx", "powerpoint": "pptx", "slides": "pptx",
               "apresentacao": "pptx", "markdown": "md", "readme": "md", "texto": "txt"}.get(formato, formato)
    if formato not in FORMATOS_ARQUIVO:
        raise ValueError("formato precisa ser um destes: " + ", ".join(FORMATOS_ARQUIVO))
    titulo = str(dados.get("titulo") or dados.get("title") or "Arquivo").strip()[:120]
    texto = dados.get("texto") or dados.get("conteudo") or ""
    colunas, linhas = dados.get("colunas"), dados.get("linhas")
    pasta = os.path.join(COMFY, "output")
    os.makedirs(pasta, exist_ok=True)
    nome = _nome_arquivo(titulo, formato)
    caminho = os.path.join(pasta, nome)
    if formato == "xlsx":
        _xlsx(caminho, titulo, colunas, linhas)
    elif formato == "csv":
        _csv(caminho, colunas, linhas)
    elif formato == "docx":
        _docx(caminho, titulo, texto, colunas, linhas)
    elif formato == "pdf":
        _pdf(caminho, titulo, texto, colunas, linhas)
    elif formato == "pptx":
        _pptx(caminho, titulo, dados.get("slides"))
    else:
        if not str(texto).strip():
            raise ValueError("mande o conteúdo em 'texto'")
        with open(caminho, "w", encoding="utf-8") as f:
            f.write(str(texto).rstrip() + "\n")
    registrar(f"arquivo criado · {nome}")
    return nome, com_chave(f"{base_api()}/imagem/{urllib.parse.quote(nome)}")


ARQUIVO_PARAMETROS = {
    "type": "object",
    "properties": {
        "formato": {"type": "string", "enum": list(FORMATOS_ARQUIVO),
                    "description": "xlsx = Excel, csv, docx = Word, pdf, pptx = PowerPoint, md = Markdown/README, txt"},
        "titulo": {"type": "string", "description": "Título do arquivo (vira o nome e o cabeçalho)."},
        "texto": {"type": "string", "description": "Para docx, pdf, md e txt: o conteúdo completo em Markdown simples "
                                                  "(# títulos, - listas, **negrito**). Escreva o texto inteiro, não um resumo."},
        "colunas": {"type": "array", "items": {"type": "string"}, "description": "Para xlsx e csv (e opcional em docx/pdf): os títulos das colunas."},
        "linhas": {"type": "array", "items": {"type": "array", "items": {}},
                   "description": "Para xlsx e csv: uma lista por linha, na ordem das colunas. Números sem 'R$'."},
        "slides": {"type": "array", "items": {"type": "object", "properties": {
                       "titulo": {"type": "string"}, "topicos": {"type": "array", "items": {"type": "string"}},
                       "notas": {"type": "string", "description": "Opcional: o que falar neste slide."}}},
                   "description": "Para pptx: um item por slide, cada um com titulo e 3 a 6 topicos curtos."},
    },
    "required": ["formato", "titulo"],
}
ARQUIVO_DESCRICAO = ("Cria um arquivo para o usuário baixar: planilha Excel ou CSV, documento Word, PDF, apresentação PowerPoint, "
                     "Markdown (README) ou texto. Você escreve todo o conteúdo de verdade (dados, textos, slides) e a ferramenta monta o arquivo. "
                     "Use quando pedirem uma planilha, tabela, documento, relatório, currículo, contrato, apresentação, slides ou README.")


MCP_TOOLS_VOZ = [
    {
        "name": "speak",
        "description": (
            "Text-to-speech in Brazilian Portuguese with Chatterbox (running on Google Colab). Returns a URL to the WAV. "
            "Pass `voice` to clone a saved reference voice (see list_voices); omit it for the built-in voice. "
            "Long texts may still be rendering after ~80s: then it returns an id to pass to get_audio."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "text": {"type": "string", "description": "What to say (up to 5000 characters)."},
                "voice": {"type": "string", "description": "Name of a saved reference voice (optional)."},
                "exaggeration": {"type": "number", "default": 0.5, "description": "Emotion intensity, 0.25-2. Higher = more dramatic."},
                "cfg_weight": {"type": "number", "default": 0.5, "description": "0.2-1. Lower = slower, more expressive pacing."},
            },
            "required": ["text"],
        },
    },
    {
        "name": "get_audio",
        "description": "Fetch the result of a speak job that was still rendering.",
        "inputSchema": {"type": "object", "properties": {"id": {"type": "string"}}, "required": ["id"]},
    },
    {
        "name": "list_voices",
        "description": "List the saved reference voices that speak can clone. New ones are uploaded with POST /vozes/<name> (raw audio body).",
        "inputSchema": {"type": "object", "properties": {}},
    },
]


MCP_TOOLS = [
    {
        "name": "generate_image",
        "description": (
            f"Generate an image from a text prompt with {PERFIL_ATUAL['nome']} on ComfyUI (running on Google Colab). {PERFIL_ATUAL.get('nota_mcp', '')}"
            "Returns a preview plus a URL to download the full PNG. If it is still rendering after ~80s "
            "(the first image loads the models), it returns a job id to pass to get_image."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "prompt": {"type": "string", "description": "What to draw. Detailed descriptions work best."},
                "width": {"type": "integer", "default": 1024, "description": "256-2048"},
                "height": {"type": "integer", "default": 1024, "description": "256-2048"},
                "steps": {"type": "integer", "description": "optional; the default suits the loaded model"},
                "seed": {"type": "integer", "default": 0, "description": "0 = random"},
            },
            "required": ["prompt"],
        },
    },
    {
        "name": "get_image",
        "description": "Fetch the result of a generate_image job that was still rendering.",
        "inputSchema": {"type": "object", "properties": {"id": {"type": "string"}}, "required": ["id"]},
    },
]


if SO_CHAT:
    MCP_TOOLS = []   # perfil so de chat: nada de imagem para oferecer
if TEM_VIDEO:
    MCP_TOOLS.append({
        "name": "generate_video",
        "description": (
            "Generate a short video (webm, ~2 s at 832x480) with Wan 2.2 5B on ComfyUI, from text or by animating a start image. "
            "It takes several minutes: expect the 'still rendering' answer and poll get_image with the id. Describe motion and camera in the prompt."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "prompt": {"type": "string", "description": "What happens in the video."},
                "image": {"type": "string", "description": "Optional start image to animate: the file name of an image generated in this session "
                                                           "(the last part of its download URL) or the image as base64."},
                "width": {"type": "integer", "description": "Only sets the aspect ratio; the longer side is always 832. Default 832x480, or the image's ratio."},
                "height": {"type": "integer"},
                "frames": {"type": "integer", "default": 49, "description": "17-121 (24 fps by default)."},
                "fps": {"type": "integer", "default": FPS_VIDEO},
                "steps": {"type": "integer", "default": PASSOS_VIDEO},
                "seed": {"type": "integer", "default": 0, "description": "0 = random"},
            },
            "required": ["prompt"],
        },
    })
if REMOVER_FUNDO:
    MCP_TOOLS.append({
        "name": "remove_background",
        "description": "Remove the background of an image (rembg, ISNet). Returns a URL to a PNG with transparent background.",
        "inputSchema": {
            "type": "object",
            "properties": {"image": {"type": "string", "description": "The file name of an image generated in this session (the last part of its "
                                                                      "download URL), or the image as base64."}},
            "required": ["image"],
        },
    })
if MUSICA:
    MCP_TOOLS.append(MCP_TOOLS_AUDIO["musica"])
if EFEITOS:
    MCP_TOOLS.append(MCP_TOOLS_AUDIO["efeito"])
if VOZ:
    MCP_TOOLS += MCP_TOOLS_VOZ
if ARQUIVOS:
    MCP_TOOLS.append({"name": "create_file", "description": ARQUIVO_DESCRICAO + " Returns the download link: show it to the user as a Markdown link.",
                      "inputSchema": ARQUIVO_PARAMETROS})


def resultado_mcp(pid, seed=None):
    pronto = esperar(pid)
    if pronto is None:
        return {"content": [{"type": "text", "text": f"Still rendering. Call get_image with id={pid} in a few seconds."}]}
    dados, arquivo = pronto
    texto = "Ready. Download: " + com_chave(f"{base_api()}/imagem/{urllib.parse.quote(arquivo)}")
    if seed:
        texto += f" | seed: {seed}"
    conteudo = [{"type": "text", "text": texto}]
    mini = miniatura(dados)
    if mini:
        conteudo.append({"type": "image", "data": mini, "mimeType": "image/jpeg"})
    return {"content": conteudo}


def mcp_responder(msg):
    """Responde uma mensagem JSON-RPC do MCP. Notificacoes (sem id) nao tem resposta."""
    mid, metodo, params = msg.get("id"), msg.get("method", ""), msg.get("params") or {}
    if mid is None:
        return None
    if metodo == "initialize":
        resultado = {
            "protocolVersion": params.get("protocolVersion", "2025-03-26"),
            "capabilities": {"tools": {}},
            "serverInfo": {"name": "comfyui-colab", "version": "1.0"},
        }
    elif metodo == "ping":
        resultado = {}
    elif metodo == "tools/list":
        resultado = {"tools": MCP_TOOLS}
    elif metodo == "tools/call":
        nome, args = params.get("name"), params.get("arguments") or {}
        try:
            if nome == "generate_image":
                pid, seed = enfileirar(args.get("prompt"), args.get("width"), args.get("height"), args.get("steps"), args.get("seed"), args)
                registrar("MCP generate_image · " + str(args.get("prompt"))[:60])
                resultado = resultado_mcp(pid, seed)
            elif nome == "generate_video":
                pid, seed = enfileirar(args.get("prompt"), args.get("width"), args.get("height"), args.get("steps"), args.get("seed"), {**args, "video": True})
                registrar("MCP generate_video · " + str(args.get("prompt"))[:60])
                resultado = resultado_mcp(pid, seed)
            elif nome == "get_image":
                registrar("MCP get_image")
                resultado = resultado_mcp(validar_id(args.get("id")))
            elif nome == "remove_background":
                recorte, arquivo = remover_fundo(bytes_da_imagem(args.get("image") or ""))
                registrar("MCP remove_background · " + arquivo)
                resultado = {"content": [{"type": "text", "text": "Ready. Download (PNG, transparent): " + com_chave(f"{base_api()}/imagem/{arquivo}")}]}
                mini = miniatura(recorte)
                if mini:
                    resultado["content"].append({"type": "image", "data": mini, "mimeType": "image/jpeg"})
            elif nome in ("generate_music", "generate_sound_effect"):
                tipo = "musica" if nome == "generate_music" else "efeito"
                pid, seed = enfileirar_audio(tipo, args.get("prompt"), args.get("lyrics"), args.get("seconds"), args.get("steps"), args.get("seed"))
                registrar(f"MCP {nome} · " + str(args.get("prompt"))[:60])
                resultado = resultado_mcp(pid, seed)
            elif nome == "speak":
                registrar("MCP speak · " + str(args.get("text"))[:60])
                resultado = resultado_fala_mcp(iniciar_fala(args.get("text"), args.get("voice"), args.get("exaggeration"), args.get("cfg_weight")))
            elif nome == "get_audio":
                resultado = resultado_fala_mcp(nome_seguro(args.get("id")))
            elif nome == "create_file":
                arquivo, url = criar_arquivo(args)
                registrar("MCP create_file · " + arquivo)
                resultado = {"content": [{"type": "text", "text": f"Ready. Download: [{arquivo}]({url})"}]}
            elif nome == "list_voices":
                resultado = {"content": [{"type": "text", "text": "Saved voices: " + (", ".join(listar_vozes()) or "none (only the built-in voice)")}]}
            else:
                raise ValueError(f"unknown tool: {nome}")
        except Exception as ex:
            registrar(f"MCP erro · {ex}")
            resultado = {"content": [{"type": "text", "text": f"Error: {ex}"}], "isError": True}
    else:
        return {"jsonrpc": "2.0", "id": mid, "error": {"code": -32601, "message": "Method not found"}}
    return {"jsonrpc": "2.0", "id": mid, "result": resultado}


# ------------------------------------------------------------------------------
# INTERFACE WEB + CHAT COM IA (a IA chama as mesmas ferramentas da API)
# ------------------------------------------------------------------------------

trabalhos_chat = {}

CHAT_SISTEMA_SO_CHAT = "Você é um assistente direto e útil. Responda em português do Brasil, a não ser que peçam outro idioma."
CHAT_ARQUIVOS = ("Quando pedirem uma planilha, tabela, documento, relatório, PDF, currículo, apresentação, slides ou README, chame criar_arquivo "
                 "com o conteúdo completo escrito por você (dados reais e coerentes, textos inteiros, 5 a 10 slides) em vez de colar o conteúdo na conversa. ")
if ARQUIVOS:
    CHAT_SISTEMA_SO_CHAT += " " + CHAT_ARQUIVOS
CHAT_SISTEMA = CHAT_SISTEMA_SO_CHAT if SO_CHAT else (
    "Você é o assistente do ComfyUI Studio, um estúdio de criação que roda no Google Colab do usuário. "
    f"O modelo de imagem carregado é {PERFIL_ATUAL['nome']}" + (" e a voz em português (Chatterbox) está ligada. " if VOZ else ". A voz está desligada nesta sessão. ")
    + PERFIL_ATUAL.get("nota_chat", "")
    + "Responda em português do Brasil, de forma direta e curta. " + (CHAT_ARQUIVOS if ARQUIVOS else "") +
    "Quando o usuário pedir uma imagem, chame gerar_imagem com um prompt detalhado em inglês (sujeito, cenário, luz, estilo, enquadramento); "
    "se o pedido for vago, escolha você os detalhes em vez de perguntar. Cada imagem leva cerca de um minuto, então gere uma por vez, "
    "a não ser que peçam mais. A imagem e o áudio aparecem sozinhos na tela do usuário: não repita links nem descreva o arquivo, "
    "só comente em uma frase o que foi feito. Você não consegue ver a imagem gerada. "
    "Se pedirem algo que as ferramentas não fazem (" + ("" if TEM_VIDEO else "vídeo, ") + "editar imagem existente" + ("" if MUSICA else ", música") + "), diga que ainda não está disponível."
)

CHAT_TOOLS = [
    {"type": "function", "function": {
        "name": "gerar_imagem",
        "description": "Gera uma imagem a partir de um prompt e mostra ao usuário.",
        "parameters": {"type": "object", "properties": {
            "prompt": {"type": "string", "description": "Descrição detalhada em inglês."},
            "largura": {"type": "integer", "description": "256 a 2048. Padrão 1024."},
            "altura": {"type": "integer", "description": "256 a 2048. Padrão 1024."},
        }, "required": ["prompt"]},
    }},
]
if SO_VIDEO or SO_CHAT:
    CHAT_TOOLS = []   # nestes perfis nao ha modelo de imagem
if TEM_VIDEO:
    CHAT_TOOLS.append({"type": "function", "function": {
        "name": "gerar_video",
        "description": "Gera um vídeo curto (cerca de 2 segundos; leva vários minutos) e mostra ao usuário. Pode partir só do texto ou animar uma imagem já gerada.",
        "parameters": {"type": "object", "properties": {
            "prompt": {"type": "string", "description": "O que acontece no vídeo, em inglês: movimento e câmera."},
            "imagem": {"type": "string", "description": "Opcional: nome do arquivo de uma imagem já gerada nesta conversa, para animá-la."},
        }, "required": ["prompt"]},
    }})
if REMOVER_FUNDO and not SO_VIDEO:
    CHAT_TOOLS.append({"type": "function", "function": {
        "name": "remover_fundo",
        "description": "Remove o fundo de uma imagem já gerada nesta conversa e mostra o recorte (PNG transparente) ao usuário.",
        "parameters": {"type": "object", "properties": {
            "imagem": {"type": "string", "description": "Nome do arquivo da imagem, como devolvido por gerar_imagem."},
        }, "required": ["imagem"]},
    }})
if MUSICA:
    CHAT_TOOLS.append({"type": "function", "function": {
        "name": "gerar_musica",
        "description": "Gera uma música (cerca de 1 minuto de espera para 30 s de música) e toca para o usuário.",
        "parameters": {"type": "object", "properties": {
            "prompt": {"type": "string", "description": "Estilo em inglês: gênero, instrumentos, clima, andamento."},
            "letra": {"type": "string", "description": "Opcional. Letra, com marcações como [verse] e [chorus]. Sem letra, a música sai instrumental."},
            "segundos": {"type": "integer", "description": "Duração, 5 a 240. Padrão 30."},
        }, "required": ["prompt"]},
    }})
if EFEITOS:
    CHAT_TOOLS.append({"type": "function", "function": {
        "name": "gerar_efeito",
        "description": "Gera um efeito sonoro ou som ambiente curto e toca para o usuário.",
        "parameters": {"type": "object", "properties": {
            "prompt": {"type": "string", "description": "O som, em inglês. Ex.: 'thunder rumbling in the distance, light rain'."},
            "segundos": {"type": "integer", "description": "Duração, 1 a 47. Padrão 8."},
        }, "required": ["prompt"]},
    }})
if VOZ:
    CHAT_TOOLS += [
        {"type": "function", "function": {
            "name": "falar",
            "description": "Transforma um texto em fala em português e toca para o usuário.",
            "parameters": {"type": "object", "properties": {
                "texto": {"type": "string"},
                "voz": {"type": "string", "description": "Nome de uma voz salva (opcional; veja listar_vozes)."},
            }, "required": ["texto"]},
        }},
        {"type": "function", "function": {
            "name": "listar_vozes",
            "description": "Lista as vozes de referência salvas que podem ser clonadas.",
            "parameters": {"type": "object", "properties": {}},
        }},
    ]


if ARQUIVOS:
    CHAT_TOOLS.append({"type": "function", "function": {"name": "criar_arquivo", "description": ARQUIVO_DESCRICAO, "parameters": ARQUIVO_PARAMETROS}})


def chamar_llm(mensagens, com_ferramentas=True, pensar=False, ao_escrever=None, ao_pensar=None, deve_parar=None):
    """Uma resposta da IA, lida aos poucos. Devolve (mensagem, {"tokens", "segundos", "tps"}).

    ao_escrever e ao_pensar recebem o texto acumulado ate ali (resposta e raciocinio), para a tela ir mostrando.
    """
    if CHAT_MODO == "reveza":
        preparar_gpu("chat")
    # reasoning_effort "none" desliga a etapa de raciocinio do Qwen3. Medido no Colab: a mesma pergunta levou 11,2 s
    # (133 tokens, quase todos de raciocinio escondido) sem isto e 0,8 s (9 tokens) com isto.
    # O "think": false e o "/no_think" no texto, usados antes, nao tinham efeito nesta rota do Ollama
    corpo = {"model": CHAT_MODELO, "messages": mensagens, "stream": True, "stream_options": {"include_usage": True},
             "reasoning_effort": "medium" if pensar else "none", "max_tokens": CHAT_MAX_TOKENS + (1500 if pensar else 0)}
    if com_ferramentas and CHAT_TOOLS:
        corpo.update(tools=CHAT_TOOLS, tool_choice="auto")
    inicio = time.time()
    for tentativa in range(3):
        for recusado in chat_recusa:
            corpo.pop(recusado, None)
            if recusado == "tools":
                corpo.pop("tool_choice", None)
        req = urllib.request.Request(CHAT_URL, data=json.dumps(corpo).encode(), headers={
            "Content-Type": "application/json", "Authorization": "Bearer ollama", "User-Agent": "comfyui-studio/1.0",
        })
        try:
            resp = urllib.request.urlopen(req, timeout=600)
            break
        except urllib.error.HTTPError as ex:
            motivo = ex.read().decode("utf-8", "ignore")[:300]
            # modelo sem suporte a ferramentas ou a raciocinio: o Ollama responde 400 dizendo qual dos dois
            novo = "tools" if "tool" in motivo.lower() and "tools" in corpo else "reasoning_effort" if ("think" in motivo.lower() or "reason" in motivo.lower()) and "reasoning_effort" in corpo else None
            if ex.code != 400 or not novo or tentativa == 2:
                raise RuntimeError(f"a IA do chat respondeu {ex.code}: " + motivo)
            chat_recusa.add(novo)
            registrar(f"o modelo do chat não aceita {'ferramentas' if novo == 'tools' else 'o controle de raciocínio'}; seguindo sem isso")

    texto, raciocinio, chamadas, uso, primeiro, ultimo_aviso = "", "", [], {}, None, [0.0]

    def avisar_tela(funcao, valor):   # no maximo ~4 vezes por segundo, para nao encher a lista de eventos
        if funcao and time.time() - ultimo_aviso[0] > 0.25:
            ultimo_aviso[0] = time.time()
            funcao(valor)

    if "event-stream" not in (resp.headers.get("Content-Type") or ""):
        mensagem = json.loads(resp.read())["choices"][0]["message"]   # servidor que nao responde aos poucos
    else:
        for bruta in resp:
            if deve_parar and deve_parar():   # o usuario clicou em parar: larga a resposta onde esta
                resp.close()
                break
            linha = bruta.decode("utf-8", "ignore").strip()
            if not linha.startswith("data:") or linha[5:].strip() == "[DONE]":
                continue
            pedaco = json.loads(linha[5:])
            uso = pedaco.get("usage") or uso
            for escolha in pedaco.get("choices") or []:
                delta = escolha.get("delta") or {}
                primeiro = primeiro or time.time()
                pensado = delta.get("reasoning") or delta.get("reasoning_content")
                if pensado:
                    raciocinio += pensado
                    avisar_tela(ao_pensar, raciocinio)
                if delta.get("content"):
                    texto += delta["content"]
                    avisar_tela(ao_escrever, texto)
                for parte in delta.get("tool_calls") or []:   # uma chamada de ferramenta pode chegar em pedacos
                    i = parte.get("index", len(chamadas))
                    while len(chamadas) <= i:
                        chamadas.append({"id": "", "type": "function", "function": {"name": "", "arguments": ""}})
                    funcao = parte.get("function") or {}
                    argumentos = funcao.get("arguments") or ""
                    chamadas[i]["id"] = parte.get("id") or chamadas[i]["id"]
                    chamadas[i]["function"]["name"] = funcao.get("name") or chamadas[i]["function"]["name"]
                    chamadas[i]["function"]["arguments"] += argumentos if isinstance(argumentos, str) else json.dumps(argumentos)
        mensagem = {"role": "assistant", "content": texto}
        if chamadas:
            mensagem["tool_calls"] = chamadas
    segundos = time.time() - inicio
    tokens = uso.get("completion_tokens")
    velocidade = tokens / max(time.time() - (primeiro or inicio), 0.05) if tokens else None
    registrar(f"a IA do chat respondeu em {segundos:.1f}s" + (f" · {tokens} tokens · {velocidade:.0f} tokens/s" if tokens else ""))
    return mensagem, {"tokens": tokens, "segundos": round(segundos, 1), "tps": round(velocidade, 1) if velocidade else None}


# Skills e atalhos que ja vem instalados no Open WebUI: aparecem quando a pessoa digita "/" no chat.
# Skill = instrucoes que deixam a IA especialista numa tarefa. Atalho = um pedido pronto, para quem nao sabe por onde comecar.
SKILLS_WEBUI = [
    {"id": "gerador-de-prompt", "name": "Gerador de prompt de imagem",
     "description": "Transforma uma ideia simples num prompt detalhado em inglês para o Z Image e outros geradores.",
     "content": """Você transforma ideias simples em prompts de imagem excelentes.

Quando o usuário descrever o que quer (mesmo em uma frase curta e em português), responda com:
1. **Prompt** em inglês, num parágrafo só, nesta ordem: sujeito e ação, roupa e detalhes, cenário, luz, câmera e enquadramento, estilo.
   Seja concreto: "golden hour side light", "85mm portrait, shallow depth of field", "shot on iPhone, candid".
2. **Variações**: mais 2 prompts curtos com outra luz ou outro enquadramento.
3. **Dica** de uma linha em português sobre o que mudar se não ficar bom.

Regras: o Z Image não usa prompt negativo, então descreva o que quer ver (em vez de "sem pessoas", escreva "empty street").
Para fotos realistas, cite câmera e luz natural. Para personagens caricatos, exagere os traços com números e comparações
("jaw twice as wide as the forehead"). Nunca invente marcas registradas nem pessoas reais."""},
    {"id": "criador-de-influencer", "name": "Criador de influencer",
     "description": "Cria a persona completa de um influencer virtual: nome, personalidade, bio, visual, ideias de posts.",
     "content": """Você é estrategista de influenciadores virtuais criados com IA.

Quando pedirem um influencer, entregue:
- **Nome e @** (curto, fácil de lembrar, com o @ livre provável).
- **Nicho e público** em uma frase.
- **Personalidade**: 3 traços e o jeito de falar (com 2 exemplos de frases).
- **Visual fixo**: rosto, cabelo, corpo, estilo de roupa e cores, escrito como um prompt em inglês para manter o personagem sempre igual.
- **Bio do Instagram** (até 150 caracteres).
- **10 ideias de posts** variando Reels, carrossel e foto.
- **Calendário da semana**: o que postar em cada dia.

Se for um personagem engraçado ou caricato, capriche no exagero e no bordão. Personagens são sempre fictícios e adultos.
Lembre o usuário de avisar nas redes que o perfil é feito com IA."""},
    {"id": "roteiro-reels", "name": "Roteirista de Reels e TikTok",
     "description": "Roteiros curtos que prendem nos primeiros 3 segundos, prontos para gravar ou narrar com a voz do Studio.",
     "content": """Você escreve roteiros de vídeos curtos (Reels, TikTok, Shorts) em português do Brasil.

Formato da resposta:
- **Gancho (0 a 3 s)**: uma frase que faz parar de rolar a tela. Dê 3 opções.
- **Roteiro** em cenas numeradas, cada uma com: o que aparece na tela, o texto falado e o texto escrito na tela.
- **Final** com chamada para ação (seguir, comentar, salvar).
- **Narração**: o texto falado inteiro junto, pronto para colar na ferramenta de voz.
- **Legenda** com 3 a 5 hashtags.

Padrão: 30 segundos, linguagem falada e simples, frases curtas. Se o usuário der o nicho, adapte as gírias e as referências."""},
    {"id": "legendas-hashtags", "name": "Legendas e hashtags",
     "description": "Legendas para Instagram e TikTok em vários tons, com hashtags certas para o nicho.",
     "content": """Você escreve legendas para redes sociais em português do Brasil.

Para cada pedido, entregue 3 legendas em tons diferentes (divertido, inspirador, direto ao ponto). Cada uma com:
- primeira linha forte (é o que aparece antes do "mais"),
- corpo curto, com quebras de linha,
- chamada para ação no final,
- 5 a 10 hashtags: misture as grandes do nicho com outras menores e mais específicas.
Use emojis com moderação, só onde ajudam a leitura."""},
    {"id": "copy-vendas", "name": "Copy de vendas",
     "description": "Anúncios, descrição de produto e mensagens de WhatsApp que vendem.",
     "content": """Você é redator publicitário focado em vendas para pequenos negócios brasileiros.

Antes de escrever, se faltar informação essencial (produto, preço, público), faça no máximo 2 perguntas. Depois entregue:
- **Título** chamativo (3 opções).
- **Texto do anúncio** curto, usando a estrutura problema, solução, prova, oferta, chamada para ação.
- **Descrição do produto** para loja virtual ou Mercado Livre, com tópicos de benefícios.
- **Mensagem de WhatsApp** pronta para enviar aos clientes, curta e pessoal.
Sem promessas falsas ou exageradas. Linguagem simples e direta."""},
    {"id": "explica-facil", "name": "Explica fácil",
     "description": "Explica qualquer assunto como para quem nunca ouviu falar, com exemplos do dia a dia.",
     "content": """Você explica qualquer assunto para quem nunca ouviu falar dele.

Regras:
- Comece com a ideia principal em uma frase.
- Use uma comparação com algo do dia a dia brasileiro.
- Explique em passos curtos, sem jargão. Se precisar de um termo técnico, explique na hora entre parênteses.
- Termine com um exemplo prático e uma pergunta para ver se a pessoa entendeu.
Se a pessoa disser que não entendeu, explique de outro jeito, ainda mais simples."""},
]

ATALHOS_WEBUI = [   # (comando, nome, texto que entra na caixa de mensagem)
    ("planilha", "Criar planilha Excel", "Crie uma planilha Excel de "),
    ("documento", "Criar documento Word", "Crie um documento Word com "),
    ("pdf", "Criar PDF", "Crie um PDF com "),
    ("slides", "Criar apresentação", "Crie uma apresentação de 8 slides sobre "),
    ("curriculo", "Criar currículo", "Crie um currículo em PDF para mim. Meus dados: nome, cidade, experiências, formação e habilidades: "),
    ("readme", "Criar README", "Crie um README.md para o meu projeto: "),
]

SUGESTOES_WEBUI = [   # aparecem na tela inicial de cada conversa nova
    {"title": ["Criar uma planilha", "de gastos do mês"], "content": "Crie uma planilha Excel para controlar os gastos do mês de uma família, com categorias e exemplos."},
    {"title": ["Roteiro de Reels", "de 30 segundos"], "content": "Escreva um roteiro de Reels de 30 segundos sobre "},
    {"title": ["Criar um influencer", "virtual"], "content": "Crie a persona de um influencer virtual de "},
    {"title": ["Pesquisar na internet", "as notícias de hoje"], "content": "Pesquise na internet as principais notícias de hoje no Brasil e resuma em tópicos."},
]


def preparar_webui(token):
    """Instala as skills e os atalhos no Open WebUI recem-criado. Se alguma falhar, o chat continua normal."""
    base = f"http://127.0.0.1:{WEBUI_PORT}/api/v1"
    cabecalhos = {"Content-Type": "application/json", **({"Authorization": f"Bearer {token}"} if token else {})}

    def postar(caminho, dados):
        with contextlib.suppress(OSError, ValueError):
            urllib.request.urlopen(urllib.request.Request(base + caminho, data=json.dumps(dados).encode(), headers=cabecalhos), timeout=30).read()
            return True
        return False

    criadas = sum(postar("/skills/create", {**s, "meta": {}, "is_active": True}) for s in SKILLS_WEBUI)
    atalhos = sum(postar("/prompts/create", {"command": c, "name": n, "content": t}) for c, n, t in ATALHOS_WEBUI if ARQUIVOS)
    return criadas, atalhos


def ambiente_webui():
    """Variaveis que ligam o Open WebUI ao que esta sessao tem. Tudo passa pela nossa API (http://127.0.0.1), nunca pelo tunel.

    A conversa vai pela rota /v1 e nao direto ao Ollama por tres motivos medidos no Colab: direto, o Qwen3 raciocinava ~40 s
    antes de cada resposta; respondia em ingles; e as ferramentas embutidas do Open WebUI (agenda, automacoes) confundiam o modelo.
    """
    nossa_api = f"http://127.0.0.1:{API_PORT}/v1"
    chave = CHAVE_API or "sem-chave"
    amb = {
        **os.environ, "DATA_DIR": WEBUI_DADOS, "WEBUI_AUTH": "True" if CHAVE_API else "False", "DEFAULT_LOCALE": "pt-BR", "WEBUI_NAME": "ComfyUI Studio Chat",
        "ANONYMIZED_TELEMETRY": "False", "SCARF_NO_ANALYTICS": "True", "DO_NOT_TRACK": "True",
        "CUDA_VISIBLE_DEVICES": "",   # o Open WebUI nao precisa da placa; ela fica para a IA e para a imagem
        "ENABLE_OLLAMA_API": "False", "ENABLE_OPENAI_API": "True", "OPENAI_API_BASE_URL": nossa_api, "OPENAI_API_KEY": chave,
        "DEFAULT_MODELS": CHAT_MODELO, "ENABLE_EVALUATION_ARENA_MODELS": "False", "ENABLE_TAGS_GENERATION": "False", "ENABLE_FOLLOW_UP_GENERATION": "False",
        "ENABLE_TITLE_GENERATION": "False",   # cada titulo e uma chamada a mais a IA, e no Colab ela ainda falhava (KeyError 'model')
        "ENABLE_AUTOMATIONS": "False", "ENABLE_CALENDAR": "False", "ENABLE_MEMORIES": "False", "ENABLE_NOTES": "False",
        "ENABLE_CODE_INTERPRETER": "False", "ENABLE_CODE_EXECUTION": "False",
        # pesquisa na web pelo DuckDuckGo: nao precisa de chave. Sem a etapa de embeddings, que baixaria mais um modelo e pesaria na RAM
        "ENABLE_WEB_SEARCH": "True", "WEB_SEARCH_ENGINE": "duckduckgo", "WEB_SEARCH_RESULT_COUNT": "3",
        "BYPASS_WEB_SEARCH_EMBEDDING_AND_RETRIEVAL": "True",
        "DEFAULT_PROMPT_SUGGESTIONS": json.dumps(SUGESTOES_WEBUI, ensure_ascii=False),
        # as ferramentas do Studio (imagem, video, som, voz, recorte) entram como um servidor MCP
        "TOOL_SERVER_CONNECTIONS": json.dumps([{
            "url": f"http://127.0.0.1:{API_PORT}/mcp", "path": "", "type": "mcp", "auth_type": "bearer" if CHAVE_API else "none", "key": CHAVE_API,
            "config": {"enable": True}, "info": {"id": "studio", "name": "ComfyUI Studio", "description": "Imagem, vídeo, som e voz deste Colab"},
        }]),
    }
    if not (SO_CHAT or SO_VIDEO):   # botao de imagem do Open WebUI -> nosso gerador
        amb.update({"ENABLE_IMAGE_GENERATION": "True", "IMAGE_GENERATION_ENGINE": "openai", "IMAGES_OPENAI_API_BASE_URL": nossa_api,
                    "IMAGES_OPENAI_API_KEY": chave, "IMAGE_GENERATION_MODEL": "studio", "IMAGE_SIZE": "1024x1024"})
    if VOZ:      # botao de ouvir a resposta -> Chatterbox
        amb.update({"AUDIO_TTS_ENGINE": "openai", "AUDIO_TTS_OPENAI_API_BASE_URL": nossa_api, "AUDIO_TTS_OPENAI_API_KEY": chave,
                    "AUDIO_TTS_MODEL": "chatterbox", "AUDIO_TTS_VOICE": "padrao"})
    if OUVIR:    # microfone -> Whisper
        amb.update({"AUDIO_STT_ENGINE": "openai", "AUDIO_STT_OPENAI_API_BASE_URL": nossa_api, "AUDIO_STT_OPENAI_API_KEY": chave, "AUDIO_STT_MODEL": "whisper"})
    return amb


def modelos_locais():
    """Modelos de chat instalados no Ollama, com o da sessao em primeiro."""
    try:
        nomes = [m.get("name") for m in json.loads(urllib.request.urlopen(f"http://127.0.0.1:{OLLAMA_PORT}/api/tags", timeout=20).read()).get("models", [])]
    except (OSError, ValueError):
        nomes = []
    outros = [n for n in nomes if n and n not in (CHAT_MODELO, CHAT_MODELO + ":latest")]
    return ([CHAT_MODELO] if CHAT_LIGADO else []) + outros


def repassar_openai(handler, corpo):
    """Rota /v1/chat/completions: o mesmo formato da OpenAI, entregue pela IA local. Serve a qualquer programa que fale esse formato."""
    if not CHAT_LIGADO:
        raise RuntimeError("o chat com IA está desligado nesta sessão (campo CHAT_IA na célula do Colab)")
    if CHAT_MODO == "reveza":
        preparar_gpu("chat")
    mensagens = list(corpo.get("messages") or [])
    if not any(m.get("role") == "system" for m in mensagens):   # quem chamou nao deu instrucao: responde em portugues
        mensagens.insert(0, {"role": "system", "content": CHAT_EM_PORTUGUES})
    # vale o modelo pedido se ele existir no Ollama (ex.: o de visao, escolhido no Open WebUI); senao, o da sessao
    pedido = corpo.get("model")
    corpo = {**corpo, "messages": mensagens, "model": pedido if pedido in modelos_locais() else CHAT_MODELO}
    if "reasoning_effort" not in chat_recusa:
        corpo.setdefault("reasoning_effort", "none")   # sem raciocinio escondido, a nao ser que o cliente peca
    req = urllib.request.Request(CHAT_URL, data=json.dumps(corpo).encode(), headers={"Content-Type": "application/json", "Authorization": "Bearer ollama"})
    try:
        resp = urllib.request.urlopen(req, timeout=600)
    except urllib.error.HTTPError as ex:
        return handler._enviar(ex.code, ex.read(), "application/json")
    registrar("POST /v1/chat/completions" + (" (aos poucos)" if corpo.get("stream") else ""))
    if not corpo.get("stream"):
        return handler._enviar(200, resp.read(), "application/json")
    # resposta aos poucos: sem tamanho conhecido, entao a conexao e fechada no fim
    handler.send_response(200)
    handler.send_header("Content-Type", "text/event-stream")
    handler.send_header("Cache-Control", "no-cache")
    handler.send_header("Connection", "close")
    handler.send_header("Access-Control-Allow-Origin", "*")
    handler.end_headers()
    for linha in resp:
        handler.wfile.write(linha)
        handler.wfile.flush()
    handler.close_connection = True


def ferramenta_chat(nome, args, evento):
    """Executa uma ferramenta pedida pela IA. Devolve o texto que volta para ela."""
    if nome == "gerar_imagem":
        pid, seed = enfileirar(args.get("prompt"), args.get("largura"), args.get("altura"))
        pronto = esperar(pid, 1800)   # aqui nao ha o limite do tunel: quem espera e esta thread, nao o navegador
        if pronto is None:
            return "O resultado ainda não ficou pronto depois de 30 minutos; algo travou no ComfyUI."
        evento("imagem", url="/imagem/" + urllib.parse.quote(pronto[1]))
        return f"Imagem gerada e já exibida ao usuário. Arquivo: {pronto[1]} (seed {seed})."
    if nome == "gerar_video":
        extras = {"video": True}
        if args.get("imagem"):
            extras["image"] = args["imagem"]
        pid, seed = enfileirar(args.get("prompt"), extras=extras)
        pronto = esperar(pid, 1800)
        if pronto is None:
            return "O vídeo ainda não ficou pronto depois de 30 minutos; algo travou no ComfyUI."
        evento("imagem", url="/imagem/" + urllib.parse.quote(pronto[1]))   # a pagina reconhece video pela extensao
        return "Vídeo gerado e já exibido ao usuário."
    if nome == "falar":
        jid = iniciar_fala(args.get("texto"), args.get("voz"))
        if esperar_fala(jid, 900) is None:
            return "A fala ainda não ficou pronta depois de 15 minutos; algo travou."
        evento("audio", url=f"/audio/{jid}.wav")
        return "Áudio gerado e já exibido ao usuário."
    if nome == "listar_vozes":
        return "Vozes salvas: " + (", ".join(listar_vozes()) or "nenhuma (só a voz padrão)")
    if nome == "remover_fundo":
        _, arquivo = remover_fundo(bytes_da_imagem(args.get("imagem") or ""))
        evento("imagem", url="/imagem/" + urllib.parse.quote(arquivo))
        return f"Fundo removido e recorte já exibido ao usuário. Arquivo: {arquivo}."
    if nome == "criar_arquivo":
        arquivo, _ = criar_arquivo(args)
        evento("arquivo", url="/imagem/" + urllib.parse.quote(arquivo), nome=arquivo)
        return f"Arquivo criado e o link de download já aparece para o usuário: {arquivo}. Não repita o link."
    if nome in ("gerar_musica", "gerar_efeito"):
        pid, seed = enfileirar_audio("musica" if nome == "gerar_musica" else "efeito", args.get("prompt"), args.get("letra"), args.get("segundos"))
        pronto = esperar(pid, 1800)
        if pronto is None:
            return "O áudio ainda não ficou pronto depois de 30 minutos; algo travou no ComfyUI."
        evento("audio", url="/imagem/" + urllib.parse.quote(pronto[1]))
        return "Áudio gerado e já tocando para o usuário."
    return f"Ferramenta desconhecida: {nome}"


def resumo_args(nome, args):
    texto = str(args.get("prompt") or args.get("titulo") or args.get("texto") or args.get("imagem") or "")
    return '"' + texto[:90] + ('…"' if len(texto) > 90 else '"') if texto else ""


def rodar_chat(trabalho, mensagens, pensar=False):
    """Laco da conversa: a IA responde ou pede ferramentas; no maximo 6 rodadas por mensagem do usuario."""
    def evento(tipo, **dados):
        trabalho["eventos"].append({"tipo": tipo, **dados})

    novas = []
    try:
        # manda so o fim do historico, comecando numa fala do usuario (nao pode cortar no meio de uma ferramenta)
        recentes = mensagens[-40:]
        while recentes and recentes[0].get("role") != "user":
            recentes = recentes[1:]
        msgs = [{"role": "system", "content": CHAT_SISTEMA}] + recentes
        for rodada in range(6):
            resposta, medida = chamar_llm(
                msgs, com_ferramentas=rodada < 5, pensar=pensar,
                ao_escrever=lambda t: evento("parcial", texto=t), ao_pensar=lambda t: evento("pensamento", texto=t[-1500:]),
                deve_parar=lambda: trabalho.get("parar"))
            texto = re.sub(r"<think>.*?</think>", "", resposta.get("content") or "", flags=re.S).strip()
            chamadas = resposta.get("tool_calls") or []
            if trabalho.get("parar"):
                texto, chamadas = (texto + "\n\n*(interrompido)*").strip(), []   # fica o que ja tinha sido escrito, sem ferramentas
            registro = {"role": "assistant", "content": texto}
            if chamadas:
                registro["tool_calls"] = chamadas
            msgs.append(registro)
            novas.append(registro)
            if texto:
                evento("texto", texto=texto)
            if medida["tokens"]:
                evento("velocidade", **medida)
            if not chamadas:
                break
            for chamada in chamadas:
                nome = chamada.get("function", {}).get("name", "")
                try:
                    args = json.loads(chamada.get("function", {}).get("arguments") or "{}")
                    evento("ferramenta", nome=nome, resumo=resumo_args(nome, args))
                    registrar(f"chat {nome} · " + resumo_args(nome, args)[:60])
                    saida = ferramenta_chat(nome, args, evento)
                except Exception as ex:
                    saida = f"Erro: {ex}"
                    evento("erro", texto=f"{nome} falhou: {ex}")
                resultado = {"role": "tool", "tool_call_id": chamada.get("id"), "content": saida}
                msgs.append(resultado)
                novas.append(resultado)
    except Exception as ex:
        evento("erro", texto=str(ex))
    trabalho["novas"] = novas
    trabalho["fim"] = True


def iniciar_chat(mensagens, pensar=False):
    if not CHAT_LIGADO:
        raise RuntimeError("o chat com IA está desligado nesta sessão (campo CHAT_IA na célula do Colab)")
    if not isinstance(mensagens, list) or not mensagens:
        raise ValueError("envie a lista de mensagens")
    cid = uuid.uuid4().hex[:12]
    trabalhos_chat[cid] = {"eventos": [], "fim": False, "novas": []}
    threading.Thread(target=rodar_chat, args=(trabalhos_chat[cid], mensagens, pensar), daemon=True).start()
    return cid


def transcrever(dados):
    """Audio (webm, mp3, wav...) -> texto em portugues, com o Whisper na CPU."""
    if whisper_modelo is None:
        raise RuntimeError("o ouvir está desligado nesta sessão (campo WHISPER na célula do Colab)")
    with trava_ouvir:
        with tempfile.NamedTemporaryFile(suffix=".audio", delete=False) as f:
            f.write(dados)
        try:
            # 16 kHz mono em float32 e o formato que o Whisper espera; o ffmpeg le qualquer coisa que o navegador grave
            ff = subprocess.run(
                ["ffmpeg", "-nostdin", "-loglevel", "error", "-i", f.name, "-f", "f32le", "-ac", "1", "-ar", "16000", "pipe:1"],
                capture_output=True,
            )
            if ff.returncode != 0 or not ff.stdout:
                raise ValueError("não consegui ler o áudio: " + ff.stderr.decode("utf-8", "ignore").strip()[-200:])
            import numpy as np
            segmentos, _ = whisper_modelo.transcribe(np.frombuffer(ff.stdout, dtype=np.float32), language="pt", beam_size=1)
            return " ".join(seg.text.strip() for seg in segmentos).strip()
        finally:
            os.remove(f.name)


def listar_galeria(limite=60):
    pasta = os.path.join(COMFY, "output")
    try:
        arquivos = [f for f in os.listdir(pasta) if f.lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".webm", ".mp4"))]
        arquivos.sort(key=lambda f: os.path.getmtime(os.path.join(pasta, f)), reverse=True)
        return arquivos[:limite]
    except OSError:
        return []


def pagina_interface():
    return (INTERFACE_HTML.replace("__BANNER__", BANNER.replace("font-size:13px", "")).replace("__PILHA__", pilha_html())
            .replace("__CRIADOR__", html.escape(CRIADOR)))


# ------------------------------------------------------------------------------
# DOCUMENTACAO DA API (pagina /docs e /openapi.json, geradas da mesma lista de rotas)
# ------------------------------------------------------------------------------

def rotas_doc():
    """[(grupo, resumo, [rota])] so com o que esta ligado nesta sessao. Cada rota e um dict."""
    api = base_api()
    json_h = '-H "Content-Type: application/json"'
    video, combinado = SO_VIDEO, TEM_VIDEO and not SO_VIDEO
    coisa, tipo_saida, ext = ("vídeo", "video/webm", "webm") if video else ("imagem", "image/png", "png")

    def rota(metodo, caminho, titulo, descricao, resposta, exemplo="", campos=(), corpo=""):
        if CHAVE_API:   # os exemplos ja saem com a chave desta sessao
            exemplos = [exemplo] if isinstance(exemplo, str) else exemplo
            exemplo = [e.replace("curl ", f'curl -H "Authorization: Bearer {CHAVE_API}" ', 1) for e in exemplos if e]
        return {"metodo": metodo, "caminho": caminho, "titulo": titulo, "descricao": descricao, "resposta": resposta,
                "exemplo": exemplo, "campos": list(campos), "corpo": corpo}

    campos_imagem = [
        ("prompt", "texto", True, "O que desenhar. Em inglês costuma sair melhor."),
        ("width", "inteiro", False, "Largura, de 256 a 2048. Padrão 1024."),
        ("height", "inteiro", False, "Altura, de 256 a 2048. Padrão 1024."),
        ("steps", "inteiro", False, f"Passos. Padrão {PERFIL_ATUAL['passos']} neste modelo."),
        ("seed", "inteiro", False, "Repita o número para repetir o resultado. Vazio ou 0 sorteia."),
    ]
    if video:
        campos_imagem[1:3] = [
            ("width", "inteiro", False, "Só define a proporção: o lado maior é sempre ajustado para 832. Padrão: paisagem 832x480."),
            ("height", "inteiro", False, "Para retrato, mande width 720 e height 1280."),
        ]
        campos_imagem += [
            ("frames", "inteiro", False, "Quadros, de 17 a 121. Padrão 49 (cerca de 2 s). Mais quadros demoram mais."),
            ("fps", "inteiro", False, "Quadros por segundo. Padrão 24."),
            ("image", "texto", False, "Imagem inicial para animar (imagem→vídeo): o nome de um arquivo da galeria ou a imagem em base64. "
                                      "Sem width e height, o vídeo segue a proporção da imagem."),
        ]
    if combinado:
        campos_imagem += [
            ("video", "booleano", False, "true gera um vídeo curto (webm, cerca de 2 s, vários minutos de espera) em vez de uma imagem."),
            ("image", "texto", False, "Imagem inicial para animar; já liga o modo vídeo. O nome de um arquivo da galeria (o X-Arquivo de uma imagem "
                                      "gerada antes) ou a imagem em base64. Sem width e height, o vídeo segue a proporção da imagem."),
            ("frames", "inteiro", False, "Só no vídeo: quadros, de 17 a 121. Padrão 49."),
            ("fps", "inteiro", False, "Só no vídeo: quadros por segundo. Padrão 24."),
        ]
        coisa, tipo_saida = "imagem ou vídeo", "image/png (ou video/webm no modo vídeo)"
    espera = "Se demorar mais de ~80 s, responde 202 com {status, id, resultado}: busque depois no endereço de 'resultado'."
    lote_ok = not SO_CHAT and not video
    grupos = [("Imagem", f"Modelo carregado: {PERFIL_ATUAL['nome']}.", [
        rota("POST", "/gerar", "Gerar " + coisa,
             "Gera a partir de um texto e devolve o arquivo na própria resposta." + (" Um vídeo leva vários minutos: conte com a resposta 202." if video else ""),
             f"{tipo_saida}, com os cabeçalhos X-Seed e X-Arquivo. " + espera,
             [f"curl -X POST {api}/gerar {json_h} -d '{{\"prompt\": \"a futuristic cat in space\"}}' -o imagem.{ext}"]
             + ([f"curl -X POST {api}/gerar {json_h} -d '{{\"prompt\": \"the camera slowly zooms in\", \"image\": \"'$(base64 -w0 foto.png)'\"}}'"] if video else [])
             + ([f"curl -X POST {api}/gerar {json_h} -d '{{\"prompt\": \"a fox running in the snow\", \"video\": true}}'",
                 f"curl -X POST {api}/gerar {json_h} -d '{{\"prompt\": \"the camera slowly zooms in\", \"image\": \"ARQUIVO_DA_GALERIA.png\"}}'"] if combinado else []),
             campos_imagem),
        rota("GET", "/resultado/{id}", "Buscar um resultado pendente", "Para quando o /gerar respondeu 202.",
             f"{tipo_saida}, ou 202 de novo se ainda estiver gerando.", f"curl {api}/resultado/ID -o imagem.{ext}"),
        rota("GET", "/imagem/{arquivo}", "Baixar uma imagem já gerada", "O nome do arquivo vem no cabeçalho X-Arquivo ou em /galeria.",
             "O arquivo da imagem.", f"curl {api}/imagem/ARQUIVO -o imagem.{ext}"),
        rota("GET", "/galeria", "Listar as imagens da sessão", "As mais recentes primeiro.", "JSON {imagens: [arquivo, ...]}.", f"curl {api}/galeria"),
    ] + ([
        rota("POST", "/remover-fundo", "Remover o fundo de uma imagem",
             "rembg (ISNet), na CPU. Mande o arquivo no corpo, ou um JSON {\"image\": ...} com o nome de um arquivo da galeria ou a imagem em base64.",
             "image/png com fundo transparente, e o cabeçalho X-Arquivo (o recorte também entra na galeria).",
             [f"curl -X POST {api}/remover-fundo --data-binary @foto.jpg -o recorte.png",
              f"curl -X POST {api}/remover-fundo {json_h} -d '{{\"image\": \"ARQUIVO_DA_GALERIA.png\"}}' -o recorte.png"],
             corpo="O arquivo da imagem (PNG, JPG ou WebP), até 40 MB."),
    ] if fundo_ligado else []) + ([
        rota("POST", "/lote", "Gerar muitas imagens de uma vez",
             f"Cada prompt gera 'quantidade' imagens; o modelo deste perfil faz até {LOTE_MAXIMO} por vez e o resto entra na fila. Responde na hora com um id.",
             "JSON {id, total, acompanhar}.",
             f"curl -X POST {api}/lote {json_h} -d '{{\"prompts\": [\"a red fox\", \"a blue bird\"], \"quantidade\": 4}}'",
             [("prompts", "lista", True, "Lista de textos (ou mande um só em 'prompt')."),
              ("quantidade", "inteiro", False, "Imagens por prompt, de 1 a 64. Padrão 4. No total, até 128 por pedido."),
              ("width", "inteiro", False, "Largura. Padrão 1024."), ("height", "inteiro", False, "Altura. Padrão 1024."),
              ("steps", "inteiro", False, f"Passos. Padrão {PERFIL_ATUAL['passos']}."),
              ("seed", "inteiro", False, "Com um número, cada pedido usa seed, seed+1, seed+2... Vazio sorteia.")]),
        rota("GET", "/lote/{id}", "Acompanhar um lote", "Chame de tempos em tempos; as imagens aparecem conforme ficam prontas.",
             "JSON {total, prontas: [arquivo], imagens: [endereço para baixar], fim, erros}.", f"curl {api}/lote/ID"),
    ] if lote_ok else []))]
    if lote_ok:
        grupos[0][2].append(rota(
            "POST", "/v1/images/generations", "Imagem no formato da OpenAI", "Para programas que já usam a API de imagens da OpenAI (é por aqui que o Open WebUI gera imagem).",
            "JSON {created, data: [{b64_json}]}.", f"curl -X POST {api}/v1/images/generations {json_h} -d '{{\"prompt\": \"a red apple\", \"n\": 1, \"size\": \"1024x1024\"}}'",
            [("prompt", "texto", True, "O que desenhar."), ("n", "inteiro", False, f"Quantas imagens, até {LOTE_MAXIMO}. Padrão 1."), ("size", "texto", False, "Largura x altura. Padrão 1024x1024.")]))
    if SO_CHAT:
        grupos = []
    if MUSICA or EFEITOS:
        rotas_som = []
        if MUSICA:
            rotas_som.append(rota(
                "POST", "/musica", "Gerar música", "ACE-Step. Sem letra, sai instrumental.", "audio/mpeg (mp3), com os cabeçalhos X-Seed e X-Arquivo. " + espera,
                f"curl -X POST {api}/musica {json_h} -d '{{\"prompt\": \"lofi hip hop, mellow piano, vinyl crackle, 80 bpm\", \"segundos\": 30}}' -o musica.mp3",
                [("prompt", "texto", True, "Estilo, em inglês: gênero, instrumentos, clima, andamento."),
                 ("letra", "texto", False, "Letra, com marcações como [verse] e [chorus]. Vazio = instrumental."),
                 ("segundos", "inteiro", False, "Duração, de 5 a 240. Padrão 30."),
                 ("steps", "inteiro", False, "Passos. Padrão 50."),
                 ("seed", "inteiro", False, "Repita o número para repetir o resultado. Vazio ou 0 sorteia.")]))
        if EFEITOS:
            rotas_som.append(rota(
                "POST", "/efeito", "Gerar efeito sonoro", "Stable Audio Open. Serve para efeitos e sons ambientes curtos.",
                "audio/mpeg (mp3), com os cabeçalhos X-Seed e X-Arquivo. " + espera,
                f"curl -X POST {api}/efeito {json_h} -d '{{\"prompt\": \"heavy wooden door creaking open\", \"segundos\": 6}}' -o efeito.mp3",
                [("prompt", "texto", True, "O som, em inglês."),
                 ("segundos", "inteiro", False, "Duração, de 1 a 47. Padrão 8."),
                 ("steps", "inteiro", False, "Passos. Padrão 50."),
                 ("seed", "inteiro", False, "Repita o número para repetir o resultado. Vazio ou 0 sorteia.")]))
        grupos.append(("Som", "Música e efeitos sonoros, gerados no ComfyUI. Um resultado pendente (202) é buscado em /resultado/{id}.", rotas_som))
    if VOZ:
        grupos.append(("Voz", "Chatterbox pt-BR. Fala com a voz padrão ou clonando uma voz de referência.", [
            rota("POST", "/falar", "Texto para fala", "Textos longos são divididos em frases e emendados.",
                 "audio/wav (24 kHz), com o cabeçalho X-Id. " + espera,
                 f"curl -X POST {api}/falar {json_h} -d '{{\"texto\": \"Olá, tudo bem?\"}}' -o fala.wav",
                 [("texto", "texto", True, "O que falar, até 5000 caracteres."),
                  ("voz", "texto", False, "Nome de uma voz salva. Vazio usa a voz padrão do modelo."),
                  ("exaggeration", "número", False, "Intensidade da emoção, 0.25 a 2. Padrão 0.5."),
                  ("cfg_weight", "número", False, "Ritmo, 0.2 a 1. Menor fica mais lento e expressivo. Padrão 0.5.")]),
            rota("GET", "/audio/{id}.wav", "Buscar uma fala pendente", "Para quando o /falar respondeu 202.", "audio/wav, ou 202 de novo.",
                 f"curl {api}/audio/ID.wav -o fala.wav"),
            rota("POST", "/vozes/{nome}", "Salvar uma voz de referência", "Envie de 10 a 15 s de fala limpa. Depois use o nome no campo 'voz'.",
                 "JSON {ok, voz, vozes}.", f"curl -X POST {api}/vozes/minha-voz --data-binary @referencia.mp3",
                 corpo="O arquivo de áudio (wav, mp3, flac, ogg ou m4a), até 30 MB."),
            rota("GET", "/vozes", "Listar as vozes salvas", "", "JSON {vozes: [nome, ...]}.", f"curl {api}/vozes"),
        ]))
    if whisper_modelo is not None:
        grupos.append(("Ouvir", f"Whisper {WHISPER}, em português.", [
            rota("POST", "/transcrever", "Fala para texto", "", "JSON {texto}.", f"curl -X POST {api}/transcrever --data-binary @audio.mp3",
                 corpo="O arquivo de áudio (qualquer formato que o ffmpeg leia), até 30 MB."),
        ]))
    if CHAT_LIGADO:
        grupos.append(("Chat", f"{CHAT_ROTULO}, rodando no Colab. A IA pode chamar as ferramentas de imagem e voz sozinha.", [
            rota("POST", "/chat", "Enviar a conversa", "Mande o histórico inteiro a cada vez; a resposta chega aos poucos em /chat/{id}.",
                 "JSON {id}.", f"curl -X POST {api}/chat {json_h} -d '{{\"mensagens\": [{{\"role\": \"user\", \"content\": \"faz um gato astronauta\"}}]}}'",
                 [("mensagens", "lista", True, "Mensagens no formato {role, content}, com role = user ou assistant."),
                  ("pensar", "booleano", False, "true para a IA raciocinar antes de responder (mais lento; o raciocínio chega nos eventos 'pensamento').")]),
            rota("GET", "/chat/{id}", "Acompanhar a resposta", "Chame em intervalos de ~1 s, passando em 'desde' quantos eventos já leu.",
                 "JSON {eventos, fim, novas}. POST /chat/{id}/parar interrompe a resposta. Eventos: parcial e pensamento (texto acumulado até ali), texto (resposta final), velocidade, "
                 "ferramenta, imagem, audio, erro. Com fim = true, junte 'novas' ao histórico.",
                 f"curl \"{api}/chat/ID?desde=0\""),
            rota("POST", "/v1/chat/completions", "Formato da OpenAI",
                 f"Para programas que já falam com a OpenAI: use {api}/v1 como base_url e a chave desta sessão como api_key. O campo 'model' é ignorado "
                 "(responde a IA desta sessão). Com \"stream\": true a resposta vem aos poucos. Sem ferramentas do Studio: é só a conversa.",
                 "O mesmo JSON da OpenAI (choices, usage).",
                 f"curl -X POST {api}/v1/chat/completions {json_h} -d '{{\"model\": \"local\", \"messages\": [{{\"role\": \"user\", \"content\": \"Olá!\"}}]}}'",
                 [("messages", "lista", True, "Mensagens no formato {role, content}."),
                  ("stream", "booleano", False, "true para receber a resposta aos poucos."),
                  ("reasoning_effort", "texto", False, "Padrão 'none' (rápido). Mande 'medium' para a IA raciocinar antes de responder.")]),
        ]))
    grupos.append(("Sistema", "", [
        rota("GET", "/status", "O que esta sessão sabe fazer", "Para um programa ou agente descobrir sozinho as capacidades antes de pedir algo.",
             "JSON: modelo de imagem, vídeo, música, efeitos, voz, ouvir, chat, ferramentas do MCP, fila do ComfyUI e uso de RAM/GPU/disco.", f"curl {api}/status"),
        rota("GET", "/openapi.json", "Esta documentação em OpenAPI 3", "Para importar em ferramentas e agentes.", "JSON.", f"curl {api}/openapi.json"),
        rota("GET", "/log", "Ver o log do ComfyUI", "Últimas linhas do que o ComfyUI escreveu; é onde aparece o motivo quando uma geração falha. "
             "Parâmetros na URL: linhas (padrão 60) e qual (comfy, voz, chat, fundo ou webui).", "Texto puro.", f"curl \"{api}/log?linhas=40\""),
        rota("POST", "/workflow", "Rodar um workflow próprio", "Para quem monta workflows no ComfyUI: no editor, use Workflow > Export (API) e mande o JSON aqui. "
             "O primeiro arquivo que o workflow salvar é o resultado.", "JSON {id, resultado}; busque o arquivo no endereço de 'resultado' (202 enquanto gera).",
             f"curl -X POST {api}/workflow {json_h} --data-binary @workflow_api.json", corpo="O workflow no formato de API do ComfyUI (JSON)."),
    ]))
    return grupos


def openapi():
    tipos = {"texto": "string", "inteiro": "integer", "número": "number", "lista": "array", "booleano": "boolean"}
    caminhos = {}
    for grupo, _, rotas in rotas_doc():
        for r in rotas:
            op = {"tags": [grupo], "summary": r["titulo"], "description": r["descricao"],
                  "responses": {"200": {"description": r["resposta"]}}}
            parametros = [{"name": n, "in": "path", "required": True, "schema": {"type": "string"}} for n in re.findall(r"\{(\w+)\}", r["caminho"])]
            if r["metodo"] == "GET" and r["caminho"] == "/chat/{id}":
                parametros.append({"name": "desde", "in": "query", "required": False, "schema": {"type": "integer"}})
            if parametros:
                op["parameters"] = parametros
            if r["campos"]:
                propriedades = {}
                for nome, tipo, _, descricao in r["campos"]:
                    propriedades[nome] = {"type": tipos[tipo], "description": descricao}
                    if tipo == "lista":
                        propriedades[nome]["items"] = {"type": "object"}
                op["requestBody"] = {"required": True, "content": {"application/json": {"schema": {
                    "type": "object", "properties": propriedades, "required": [n for n, _, obrigatorio, _ in r["campos"] if obrigatorio]}}}}
            elif r["corpo"]:
                op["requestBody"] = {"required": True, "description": r["corpo"],
                                     "content": {"application/octet-stream": {"schema": {"type": "string", "format": "binary"}}}}
            caminhos.setdefault(r["caminho"], {})[r["metodo"].lower()] = op
    return {
        "openapi": "3.0.3",
        "info": {"title": "ComfyUI Studio API", "version": "1.0",
                 "description": f"Imagem, voz e chat rodando no Google Colab. Modelo de imagem: {PERFIL_ATUAL['nome']}. O MCP fica em POST /mcp."},
        "servers": [{"url": base_api()}],
        "paths": caminhos,
        **({"components": {"securitySchemes": {"chave": {"type": "http", "scheme": "bearer"}}}, "security": [{"chave": []}]} if CHAVE_API else {}),
    }


def pagina_docs():
    """Pagina unica com tudo que da para chamar: rotas, campos, exemplos prontos para copiar e a ligacao por MCP."""
    esc = html.escape
    api = base_api()

    def bloco(texto):
        return (f"<div class='cod'><pre>{esc(texto)}</pre>"
                f"<button data-c=\"{esc(texto)}\" onclick=\"copiar(this)\">copiar</button></div>")

    p = [
        "<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'>"
        "<title>ComfyUI Studio — API</title><style>"
        ":root{--azul:#172dd7;--verde:#1a7f37;--cinza:#6e7781;--texto:#1f2328;--borda:#57606a;--claro:#d0d7de;--fundo:#f6f8fa}"
        "*{box-sizing:border-box}body{margin:0;background:#e9ecef;color:var(--texto);padding:28px 14px 60px;"
        "font:13.5px/1.7 'Cascadia Mono','Cascadia Code',Consolas,'DejaVu Sans Mono',Menlo,monospace}"
        ".janela{background:#fff;border:1px solid var(--borda);box-shadow:0 8px 30px rgba(0,0,0,.2);max-width:980px;margin:0 auto}"
        ".titulo{background:#f3f4f6;border-bottom:1px solid var(--borda);height:32px;display:flex;align-items:center;font:12px 'Segoe UI',system-ui,sans-serif}"
        ".titulo img{width:16px;height:16px;margin:0 10px;border-radius:3px}.titulo span.n{flex:1}.titulo .b{width:44px;text-align:center}"
        ".corpo{padding:28px 32px 36px}pre{margin:0;font:inherit;white-space:pre-wrap;word-break:break-all}"
        ".banner pre{white-space:pre;word-break:normal;font-size:11px;line-height:1.15;overflow:hidden}"
        "h2{font-size:16px;color:var(--azul);margin:38px 0 4px;padding-bottom:4px;border-bottom:1px solid var(--claro)}"
        "h3{font-size:13.5px;margin:22px 0 2px;display:flex;align-items:center;gap:10px;flex-wrap:wrap}"
        ".m{font-size:11px;font-weight:bold;color:#fff;padding:0 8px;border-radius:3px}.GET{background:var(--verde)}.POST{background:var(--azul)}"
        ".cam{font-weight:bold}.tit{color:var(--cinza);font-weight:normal}.d{color:var(--cinza)}"
        "table{border-collapse:collapse;margin:6px 0;width:100%}td{border:1px solid var(--claro);padding:2px 10px;vertical-align:top}"
        "td:first-child{white-space:nowrap;font-weight:bold}td:nth-child(2){white-space:nowrap;color:var(--cinza)}"
        ".cod{display:flex;gap:8px;align-items:flex-start;background:var(--fundo);border:1px solid var(--claro);padding:6px 10px;margin:6px 0}"
        ".cod pre{flex:1;min-width:0}button{font:inherit;font-size:11px;background:#fff;border:1px solid var(--borda);padding:0 8px;cursor:pointer;white-space:nowrap}"
        "button:hover{background:var(--azul);color:#fff;border-color:var(--azul)}"
        ".nav{display:flex;gap:6px 14px;flex-wrap:wrap;margin-top:14px}.nav a,a{color:var(--azul)}.r{color:var(--cinza)}.r b{color:var(--texto)}"
        "@media(max-width:700px){body{padding:0}.janela{border:0}.corpo{padding:18px 14px 30px}.banner pre{font-size:5.5px}}"
        "</style></head><body><div class='janela'>"
        "<div class='titulo'><img src='https://github.com/Comfy-Org.png?size=64' alt=''>"
        f"<span class='n'>Prompt de Comando — comfyui-studio-api.exe · por {esc(CRIADOR)}</span>"
        "<span class='b'>─</span><span class='b'>☐</span><span class='b'>✕</span></div><div class='corpo'>"
        f"<div class='banner'>{BANNER}</div>"
        f"<div style='margin-top:10px'>{pilha_html()}</div>"
        "<p class='d' style='margin:16px 0 4px'>Tudo o que este Colab faz, num endereço só. Todas as rotas abaixo partem deste endereço, "
        "que muda a cada execução da célula e só funciona enquanto ela estiver rodando.</p>"
    ]
    p.append(bloco(api))
    if CHAVE_API:
        p.append("<p class='d' style='margin:10px 0 4px'><b style='color:var(--texto)'>🔑 Chave da API.</b> Toda chamada precisa dela, no cabeçalho "
                 "<b>Authorization: Bearer &lt;chave&gt;</b> ou na URL como <b>?chave=&lt;chave&gt;</b>. Ela muda a cada execução da célula. "
                 "Os exemplos desta página já estão com a chave desta sessão.</p>")
        p.append(bloco(CHAVE_API))
    else:
        p.append("<p class='d' style='margin:10px 0 4px'>⚠ A API está aberta (PROTEGER_API desmarcado): quem tiver o endereço consegue usar.</p>")
    grupos = rotas_doc()
    p.append("<div class='nav'>" + "".join(f"<a href='#{esc(g)}'>{esc(g)}</a>" for g, _, _ in grupos[:-1])
             + "<a href='#MCP'>MCP</a><a href='#Sistema'>Sistema</a><a href='/'>abrir o Studio →</a></div>")

    def secao(grupo, resumo, rotas):
        p.append(f"<h2 id='{esc(grupo)}'>{esc(grupo)}</h2>" + (f"<div class='d'>{esc(resumo)}</div>" if resumo else ""))
        for r in rotas:
            p.append(f"<h3><span class='m {r['metodo']}'>{r['metodo']}</span><span class='cam'>{esc(r['caminho'])}</span>"
                     f"<span class='tit'>{esc(r['titulo'])}</span></h3>")
            if r["descricao"]:
                p.append(f"<div class='d'>{esc(r['descricao'])}</div>")
            if r["campos"]:
                p.append("<table>" + "".join(
                    f"<tr><td>{esc(nome)}{'*' if obrigatorio else ''}</td><td>{tipo}</td><td>{esc(descricao)}</td></tr>"
                    for nome, tipo, obrigatorio, descricao in r["campos"]) + "</table>")
                p.append("<div class='d' style='font-size:12px'>Corpo em JSON. * = obrigatório.</div>")
            if r["corpo"]:
                p.append(f"<div class='r'><b>Corpo:</b> {esc(r['corpo'])}</div>")
            p.append(f"<div class='r'><b>Resposta:</b> {esc(r['resposta'])}</div>")
            for exemplo in ([r["exemplo"]] if isinstance(r["exemplo"], str) else r["exemplo"]):
                if exemplo:
                    p.append(bloco(exemplo))

    for grupo, resumo, rotas in grupos[:-1]:
        secao(grupo, resumo, rotas)

    p.append("<h2 id='MCP'>MCP</h2><div class='d'>Para agentes (Claude Code, Codex, Cursor e outros) usarem tudo isto como ferramentas. "
             "Transporte: Streamable HTTP, em POST /mcp.</div>")
    p.append(bloco(com_chave(f"{api}/mcp")))
    marca = "<div class='r' style='display:flex;align-items:center;gap:8px;margin-top:10px'>{}<b>{}</b></div>"
    p.append(marca.format(logo("anthropics", 18), "Claude Code") + bloco("claude mcp add --transport http comfyui " + com_chave(f"{api}/mcp")))
    p.append(marca.format(logo("openai", 18), "Codex") + bloco("codex mcp add comfyui --url " + com_chave(f"{api}/mcp")))
    p.append("<table>" + "".join(f"<tr><td>{esc(t['name'])}</td><td>ferramenta</td><td>{esc(t['description'])}</td></tr>" for t in MCP_TOOLS) + "</table>")

    secao(*grupos[-1])
    outros = [("Studio (chat, imagem, voz e galeria)", api + "/")]
    if ui["link"]:
        outros.append(("ComfyUI (editor de workflows)", ui["link"]))
    if ui["link_colab"]:
        outros.append(("ComfyUI pelo endereço do Colab (só no navegador de quem rodou a célula)", ui["link_colab"]))
    p.append("<h2>Outros endereços</h2>" + "".join(
        f"<div class='r'><b>{esc(nome)}:</b> <a href='{esc(url)}' target='_blank'>{esc(url)}</a></div>" for nome, url in outros))
    p.append(
        "</div></div><script>function copiar(b){var t=b.getAttribute('data-c');"
        "(navigator.clipboard?navigator.clipboard.writeText(t):Promise.reject()).catch(function(){"
        "var a=document.createElement('textarea');a.value=t;document.body.appendChild(a);a.select();document.execCommand('copy');a.remove()});"
        "b.textContent='copiado ✔';setTimeout(function(){b.textContent='copiar'},1500)}</script></body></html>"
    )
    return "".join(p)


class ApiHandler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, *args):
        pass

    def _enviar(self, codigo, corpo=b"", tipo="application/json", extras=None):
        if isinstance(corpo, (dict, list)):
            corpo = json.dumps(corpo, ensure_ascii=False).encode()
        self.send_response(codigo)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(corpo)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        for nome, valor in (extras or {}).items():
            self.send_header(nome, str(valor))
        self.end_headers()
        self.wfile.write(corpo)

    def _corpo(self, limite=80 * 1024 * 1024):
        """Corpo do pedido, com tamanho declarado ou em pedacos (o Open WebUI manda arquivos em pedacos, sem Content-Length)."""
        if "chunked" in (self.headers.get("Transfer-Encoding") or "").lower():
            partes, total = [], 0
            while True:
                tamanho = int(self.rfile.readline().split(b";")[0].strip() or b"0", 16)
                if tamanho == 0:
                    while self.rfile.readline().strip():   # cabecalhos finais, se houver
                        pass
                    break
                total += tamanho
                if total > limite:
                    raise ValueError("corpo grande demais")
                partes.append(self.rfile.read(tamanho))
                self.rfile.readline()   # fim de linha depois de cada pedaco
            return b"".join(partes)
        tamanho = int(self.headers.get("Content-Length") or 0)
        if tamanho > limite:
            raise ValueError("corpo grande demais (máximo de 80 MB)")
        return self.rfile.read(tamanho) if tamanho else b""

    def _json(self):
        corpo = self._corpo()
        return json.loads(corpo) if corpo else {}

    def _chave_enviada(self):
        """A chave pode vir no cabecalho (scripts e MCP), na URL (?chave=, para links) ou no cookie (o Studio no navegador)."""
        autorizacao = self.headers.get("Authorization") or ""
        if autorizacao.lower().startswith("bearer "):
            return autorizacao[7:].strip()
        if self.headers.get("X-API-Key"):
            return self.headers["X-API-Key"].strip()
        na_url = self._consulta().get("chave")
        if na_url:
            return na_url[0]
        for parte in (self.headers.get("Cookie") or "").split(";"):
            nome, _, valor = parte.strip().partition("=")
            if nome == "studio_chave":
                return valor
        return ""

    def _barrado(self):
        """Responde 401 e devolve True quando a API esta protegida e a chave nao veio (ou veio errada)."""
        if not CHAVE_API or hmac.compare_digest(self._chave_enviada().encode(), CHAVE_API.encode()):
            return False
        if self.command == "GET" and self._caminho() in ("", "/docs"):
            self._enviar(401, (
                "<!doctype html><meta charset='utf-8'><body style='font:15px/1.6 Consolas,monospace;padding:40px;max-width:640px'>"
                "<h3>🔑 Esta página é protegida</h3><p>Abra pelo link que aparece no painel do Colab (a caixa verde): ele já leva a chave desta sessão.</p>"
            ).encode(), "text/html; charset=utf-8")
        else:
            self._enviar(401, {"erro": "chave da API ausente ou incorreta. Mande o cabeçalho 'Authorization: Bearer <chave>' "
                                       "(a chave aparece no painel do Colab) ou acrescente ?chave=<chave> ao endereço."},
                         extras={"WWW-Authenticate": "Bearer"})
        return True

    def _caminho(self):
        return self.path.split("?")[0].rstrip("/")

    def _consulta(self):
        return urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)

    def do_OPTIONS(self):
        self._enviar(204)

    def do_GET(self):
        try:
            caminho = self._caminho()
            if caminho != "/health" and self._barrado():
                return
            if CHAVE_API and caminho in ("", "/docs") and self._consulta().get("chave"):
                # entrou pelo link com a chave: guarda num cookie e tira a chave da barra de enderecos.
                # Com o cookie, as imagens, os audios e as chamadas da propria pagina passam sem precisar repetir a chave
                self._enviar(302, extras={"Location": caminho or "/",
                                          "Set-Cookie": f"studio_chave={CHAVE_API}; Path=/; HttpOnly; SameSite=Lax; Max-Age=86400"})
            elif caminho == "":
                self._enviar(200, pagina_interface().encode(), "text/html; charset=utf-8")
            elif caminho == "/status":
                fila = None
                try:
                    f = json.loads(chamar_comfy("/queue"))
                    fila = len(f.get("queue_running", [])) + len(f.get("queue_pending", []))
                except (OSError, ValueError):
                    pass
                self._enviar(200, {
                    "versao": VERSAO, "comfyui": fila is not None, "fila": fila, "imagem": not SO_VIDEO and not SO_CHAT, "modelo": PERFIL_ATUAL["nome"],
                    "video": TEM_VIDEO, "musica": MUSICA, "efeitos": EFEITOS, "voz": VOZ, "remover_fundo": fundo_ligado,
                    "ouvir": WHISPER if whisper_modelo is not None else None, "chat": CHAT_ROTULO if CHAT_LIGADO else None,
                    "mcp": True, "ferramentas_mcp": [t["name"] for t in MCP_TOOLS], "maquina": [texto for texto, _ in metricas()],
                })
            elif caminho == "/docs":
                self._enviar(200, pagina_docs().encode(), "text/html; charset=utf-8")
            elif caminho == "/openapi.json":
                self._enviar(200, openapi())
            elif caminho == "/info":
                self._enviar(200, {
                    "perfil": PERFIL_ATUAL["nome"], "tipo": PERFIL_ATUAL["chave"], "lote": SO_LOTE, "lote_maximo": LOTE_MAXIMO, "video": TEM_VIDEO, "musica": MUSICA, "efeitos": EFEITOS, "fundo": fundo_ligado, "passos": PERFIL_ATUAL["passos"], "voz": VOZ,
                    "versao": VERSAO, "manutencao": MANUTENCAO, "webui": ui["link_webui"],
                    "chat": bool(CHAT_LIGADO), "chat_modelo": CHAT_ROTULO + {"cpu": " · CPU", "fixa": " · GPU fixa"}.get(CHAT_MODO, " · GPU"), "criador": CRIADOR,
                    "ouvir": WHISPER if whisper_modelo is not None else "",
                })
            elif caminho == "/galeria":
                self._enviar(200, {"imagens": listar_galeria()})
            elif caminho.startswith("/chat/"):
                trabalho = trabalhos_chat.get(nome_seguro(caminho[len("/chat/"):]))
                if not trabalho:
                    self._enviar(404, {"erro": "conversa não encontrada"})
                else:
                    desde = int((self._consulta().get("desde") or ["0"])[0])
                    fim = trabalho["fim"]   # le antes dos eventos: se terminar no meio, o proximo pedido pega o resto
                    self._enviar(200, {"eventos": trabalho["eventos"][desde:], "fim": fim, "novas": trabalho["novas"] if fim else []})
            elif caminho.startswith("/lote/"):
                self._enviar(200, andamento_do_lote(nome_seguro(caminho[len("/lote/"):])))
            elif caminho == "/v1/models":
                self._enviar(200, {"object": "list", "data": [{"id": m, "object": "model", "owned_by": "local"} for m in modelos_locais()]})
            elif caminho == "/log":
                consulta = self._consulta()
                arquivo_log = {"comfy": LOG, "voz": VOZ_LOG, "chat": OLLAMA_LOG, "fundo": FUNDO_LOG, "webui": WEBUI_LOG}.get((consulta.get("qual") or ["comfy"])[0], LOG)
                linhas = _inteiro((consulta.get("linhas") or [None])[0], 60, 1, 400)
                try:
                    with open(arquivo_log, "rb") as f:
                        f.seek(0, 2)
                        f.seek(max(0, f.tell() - 200000))
                        texto = re.sub(chr(27) + r"\[[0-9;?]*[A-Za-z]", "", f.read().decode("utf-8", "ignore"))
                except OSError:
                    texto = ""
                # barras de progresso reescrevem a mesma linha: fica so a versao mais recente de cada uma
                limpas = [l.rstrip(CR).split(CR)[-1] for l in texto.split(NL)]
                self._enviar(200, NL.join([l for l in limpas if l.strip()][-linhas:]).encode(), "text/plain; charset=utf-8")
            elif caminho == "/health":
                self._enviar(200, {
                    "servico": "ComfyUI + " + PERFIL_ATUAL["nome"],
                    "gerar": "POST /gerar  {prompt, width, height, steps, seed}  -> image/png",
                    "resultado": "GET /resultado/<id>  (quando /gerar responde 202)",
                    "imagem": "GET /imagem/<arquivo>",
                    "falar": "POST /falar  {texto, voz, exaggeration, cfg_weight}  -> audio/wav" if VOZ else "desativada (marque VOZ na celula)",
                    "vozes": "GET /vozes  |  POST /vozes/<nome>  (corpo = audio de referencia)",
                    "mcp": "POST /mcp  (Streamable HTTP)",
                    "interface": "GET /  (página com chat)",
                    "docs": "GET /docs  (documentação)  |  GET /openapi.json",
                    "transcrever": "POST /transcrever  (corpo = audio)  -> {texto}" if whisper_modelo is not None else "desligado",
                })
            elif caminho.startswith("/imagem/"):
                nome = os.path.basename(urllib.parse.unquote(caminho[len("/imagem/"):]))
                no_disco = os.path.join(COMFY, "output", nome)
                ext = nome.rsplit(".", 1)[-1].lower()
                if nome and os.path.isfile(no_disco):
                    with open(no_disco, "rb") as f:
                        conteudo = f.read()
                    if ext in TIPOS_ARQUIVO:   # documentos: o navegador baixa em vez de tentar abrir
                        self._enviar(200, conteudo, TIPOS_ARQUIVO[ext], {"Content-Disposition": f'attachment; filename="{nome}"'})
                    else:
                        self._enviar(200, conteudo, tipo_imagem(nome))
                else:
                    url = f"http://127.0.0.1:{PORT}/view?" + urllib.parse.urlencode({"filename": nome, "type": "output"})
                    try:
                        self._enviar(200, urllib.request.urlopen(url, timeout=60).read(), tipo_imagem(nome))
                    except OSError:
                        self._enviar(404, {"erro": "arquivo não encontrado"})
            elif caminho.startswith("/resultado/"):
                pid = validar_id(caminho[len("/resultado/"):])
                pronto = esperar(pid, 60)
                if pronto is None:
                    self._enviar(202, {"status": "gerando", "id": pid, "resultado": f"{base_api()}/resultado/{pid}"})
                else:
                    self._enviar(200, pronto[0], tipo_imagem(pronto[1]), {"X-Arquivo": pronto[1]})
            elif caminho == "/vozes":
                self._enviar(200, {"vozes": listar_vozes(), "ativada": VOZ})
            elif caminho.startswith("/audio/"):
                jid = nome_seguro(os.path.splitext(os.path.basename(caminho))[0])
                arquivo = esperar_fala(jid, 60)
                if arquivo is None:
                    self._enviar(202, {"status": "gerando", "id": jid, "resultado": f"{base_api()}/audio/{jid}.wav"})
                else:
                    with open(arquivo, "rb") as f:
                        self._enviar(200, f.read(), "audio/wav")
            elif caminho == "/mcp":
                self._enviar(405, {"erro": "use POST"}, extras={"Allow": "POST"})
            else:
                self._enviar(404, {"erro": "rota não encontrada"})
        except Exception as ex:
            self._enviar(500, {"erro": str(ex)})

    def do_POST(self):
        try:
            caminho = self._caminho()
            if self._barrado():
                return
            if caminho == "/gerar":
                dados = self._json()
                inicio = time.time()
                pid, seed = enfileirar(dados.get("prompt"), dados.get("width"), dados.get("height"), dados.get("steps"), dados.get("seed"), dados)
                registrar("POST /gerar · " + str(dados.get("prompt"))[:60])
                pronto = esperar(pid)
                if pronto is None:
                    self._enviar(202, {"status": "gerando", "id": pid, "seed": seed, "resultado": f"{base_api()}/resultado/{pid}"})
                else:
                    registrar(f"imagem pronta em {time.time() - inicio:.0f}s · {pronto[1]}")
                    self._enviar(200, pronto[0], tipo_imagem(pronto[1]), {"X-Seed": seed, "X-Arquivo": pronto[1]})
            elif caminho == "/chat":
                dados = self._json()
                self._enviar(200, {"id": iniciar_chat(dados.get("mensagens"), bool(dados.get("pensar")))})
            elif caminho.startswith("/chat/") and caminho.endswith("/parar"):
                trabalho = trabalhos_chat.get(nome_seguro(caminho[len("/chat/"):-len("/parar")]))
                if trabalho:
                    trabalho["parar"] = True
                self._enviar(200, {"ok": bool(trabalho)})
            elif caminho == "/v1/chat/completions":
                repassar_openai(self, self._json())
            elif caminho == "/v1/images/generations":
                # formato da OpenAI: {"prompt", "n", "size": "1024x1024"} -> {"data": [{"b64_json"}]}. E por aqui que o Open WebUI gera imagem
                dados = self._json()
                if SO_VIDEO:
                    raise RuntimeError("este perfil gera vídeo; para imagens escolha um perfil de imagem")
                largura, _, altura = str(dados.get("size") or "1024x1024").partition("x")
                pid, _ = enfileirar(dados.get("prompt"), largura, altura, extras={"batch": dados.get("n")})
                registrar("POST /v1/images/generations · " + str(dados.get("prompt"))[:60])
                fim = time.time() + 900
                arquivos = None
                while arquivos is None and time.time() < fim:
                    arquivos = arquivos_do_pedido(pid)
                    if arquivos is None:
                        time.sleep(1)
                if not arquivos:
                    raise RuntimeError("a imagem não ficou pronta em 15 minutos")
                imagens = []
                for arquivo in arquivos:
                    url = f"http://127.0.0.1:{PORT}/view?" + urllib.parse.urlencode({"filename": arquivo, "type": "output"})
                    imagens.append({"b64_json": base64.b64encode(urllib.request.urlopen(url, timeout=60).read()).decode()})
                self._enviar(200, {"created": int(time.time()), "data": imagens})
            elif caminho == "/v1/audio/speech":
                # formato da OpenAI: {"input", "voice"} -> audio. "voice" vale se for o nome de uma voz salva; senao, a voz padrao
                dados = self._json()
                voz = nome_seguro(dados.get("voice"))
                jid = iniciar_fala(dados.get("input"), voz if voz in listar_vozes() else None)
                registrar("POST /v1/audio/speech · " + str(dados.get("input"))[:60])
                arquivo = esperar_fala(jid, 900)
                if arquivo is None:
                    raise RuntimeError("a fala não ficou pronta em 15 minutos")
                with open(arquivo, "rb") as f:
                    audio = f.read()
                mp3 = subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-i", arquivo, "-f", "mp3", "-b:a", "128k", "pipe:1"], capture_output=True)
                if mp3.returncode == 0 and mp3.stdout:   # quem pede neste formato espera mp3; sem ffmpeg vai o wav mesmo
                    self._enviar(200, mp3.stdout, "audio/mpeg")
                else:
                    self._enviar(200, audio, "audio/wav")
            elif caminho == "/v1/audio/transcriptions":
                # formato da OpenAI: formulario com o campo "file" -> {"text"}
                from email.parser import BytesParser
                from email.policy import HTTP
                corpo = self._corpo()
                formulario = BytesParser(policy=HTTP).parsebytes(b"Content-Type: " + (self.headers.get("Content-Type") or "").encode() + b"\r\n\r\n" + corpo)
                arquivo_enviado = next((parte.get_payload(decode=True) for parte in formulario.iter_parts() if parte.get_filename()), None) if formulario.is_multipart() else corpo
                if not arquivo_enviado:
                    raise ValueError("mande o áudio no campo 'file' do formulário")
                inicio = time.time()
                texto = transcrever(arquivo_enviado)
                registrar(f"POST /v1/audio/transcriptions · {time.time() - inicio:.1f}s · {texto[:60]}")
                self._enviar(200, {"text": texto})
            elif caminho == "/lote":
                dados = self._json()
                lid = iniciar_lote(dados)
                registrar(f"POST /lote · {lotes[lid]['total']} imagens em {len(lotes[lid]['pedidos'])} pedidos")
                self._enviar(200, {"id": lid, "total": lotes[lid]["total"], "acompanhar": f"{base_api()}/lote/{lid}"})
            elif caminho == "/manutencao/exec":
                dados = self._json()
                senha = self.headers.get("X-Senha") or str(dados.get("senha") or "")
                if not MANUTENCAO or not hmac.compare_digest(senha.encode(), SENHA_MANUTENCAO.encode()):
                    self._enviar(403, {"erro": "manutenção desligada ou senha incorreta"})
                else:
                    codigo = str(dados.get("codigo") or "")
                    registrar(f"manutenção · código executado ({len(codigo)} caracteres)")
                    saida, erro = io.StringIO(), None
                    # roda no mesmo espaco de nomes da celula: redefinir uma funcao aqui muda o comportamento na hora
                    with trava_manutencao, contextlib.redirect_stdout(saida):
                        try:
                            exec(codigo, globals())
                        except Exception:
                            erro = traceback.format_exc()
                    self._enviar(200, {"saida": saida.getvalue()[-20000:], "erro": erro})
            elif caminho == "/workflow":
                dados = self._json()
                grafo = dados.get("workflow") if isinstance(dados.get("workflow"), dict) else dados
                preparar_gpu("imagem")
                try:
                    pid = json.loads(chamar_comfy("/prompt", {"prompt": grafo}))["prompt_id"]
                except urllib.error.HTTPError as ex:
                    raise RuntimeError("o ComfyUI recusou o workflow: " + ex.read().decode("utf-8", "ignore")[:1500])
                registrar(f"POST /workflow · {len(grafo)} nodes")
                self._enviar(200, {"id": pid, "resultado": f"{base_api()}/resultado/{pid}"})
            elif caminho == "/remover-fundo":
                tamanho = int(self.headers.get("Content-Length") or 0)
                if not 0 < tamanho <= 40 * 1024 * 1024:
                    raise ValueError("envie a imagem no corpo da requisição (até 40 MB)")
                corpo = self.rfile.read(tamanho)
                if "json" in (self.headers.get("Content-Type") or ""):
                    try:
                        corpo = bytes_da_imagem(json.loads(corpo).get("image") or "")
                    except Exception:
                        raise ValueError("no JSON, mande em 'image' o nome de um arquivo da galeria ou a imagem em base64")
                inicio = time.time()
                recorte, nome = remover_fundo(corpo)
                registrar(f"fundo removido em {time.time() - inicio:.1f}s · {nome}")
                self._enviar(200, recorte, "image/png", {"X-Arquivo": nome})
            elif caminho in ("/musica", "/efeito"):
                dados = self._json()
                inicio = time.time()
                pid, seed = enfileirar_audio(caminho[1:], dados.get("prompt"), dados.get("letra"), dados.get("segundos"), dados.get("steps"), dados.get("seed"))
                registrar(f"POST {caminho} · " + str(dados.get("prompt"))[:60])
                pronto = esperar(pid)
                if pronto is None:
                    self._enviar(202, {"status": "gerando", "id": pid, "seed": seed, "resultado": f"{base_api()}/resultado/{pid}"})
                else:
                    registrar(f"áudio pronto em {time.time() - inicio:.0f}s · {pronto[1]}")
                    self._enviar(200, pronto[0], tipo_imagem(pronto[1]), {"X-Seed": seed, "X-Arquivo": pronto[1]})
            elif caminho == "/transcrever":
                tamanho = int(self.headers.get("Content-Length") or 0)
                if not 0 < tamanho <= 30 * 1024 * 1024:
                    raise ValueError("envie o áudio no corpo da requisição (até 30 MB)")
                inicio = time.time()
                texto = transcrever(self.rfile.read(tamanho))
                registrar(f"ouvi em {time.time() - inicio:.1f}s · {texto[:60]}")
                self._enviar(200, {"texto": texto})
            elif caminho == "/falar":
                dados = self._json()
                jid = iniciar_fala(dados.get("texto"), dados.get("voz"), dados.get("exaggeration"), dados.get("cfg_weight"))
                registrar("POST /falar · " + str(dados.get("texto"))[:60])
                arquivo = esperar_fala(jid)
                if arquivo is None:
                    self._enviar(202, {"status": "gerando", "id": jid, "resultado": f"{base_api()}/audio/{jid}.wav"})
                else:
                    with open(arquivo, "rb") as f:
                        self._enviar(200, f.read(), "audio/wav", {"X-Id": jid})
            elif caminho.startswith("/vozes/"):
                if not VOZ:
                    raise RuntimeError("a voz não foi ativada nesta sessão (marque VOZ na célula do Colab e rode de novo)")
                tamanho = int(self.headers.get("Content-Length") or 0)
                if not 0 < tamanho <= 30 * 1024 * 1024:
                    raise ValueError("envie o áudio no corpo da requisição (até 30 MB)")
                nome = salvar_voz(caminho[len("/vozes/"):], self.rfile.read(tamanho))
                registrar(f"voz de referência salva · {nome}")
                self._enviar(200, {"ok": True, "voz": nome, "vozes": listar_vozes()})
            elif caminho == "/mcp":
                msg = self._json()
                if isinstance(msg, list):
                    respostas = [r for r in (mcp_responder(m) for m in msg) if r is not None]
                    if respostas:
                        self._enviar(200, respostas)
                    else:
                        self._enviar(202)
                else:
                    resposta = mcp_responder(msg)
                    if resposta is None:
                        self._enviar(202)
                    else:
                        self._enviar(200, resposta)
            else:
                self._enviar(404, {"erro": "rota não encontrada"})
        except Exception as ex:
            registrar(f"erro · {ex}")
            self._enviar(500, {"erro": str(ex)})


def liberar_disco(folga=8e9):
    """Ao trocar de pack, apaga modelos de packs anteriores que este nao usa, mas so se faltar espaco para os novos downloads.

    Assim voltar para um pack ja usado continua rapido (os arquivos ficam), e o disco do Colab (~70 GB livres) nao lota.
    """
    pasta_modelos = os.path.join(COMFY, "models")
    faltam = sum(_tamanho_previsto(url) for p, arq, url, *_ in MODELOS if not os.path.isfile(os.path.join(pasta_modelos, p, arq)))
    usados = {os.path.join(pasta_modelos, p, arq) for p, arq, *_ in MODELOS}
    sobrando = []
    for raiz, _, arquivos in os.walk(pasta_modelos):
        for nome in arquivos:
            caminho = os.path.join(raiz, nome)
            if caminho not in usados and nome.endswith((".safetensors", ".gguf", ".ckpt", ".pt", ".pth")):
                sobrando.append((os.path.getsize(caminho), caminho))
    for tamanho, caminho in sorted(sobrando, reverse=True):   # os maiores primeiro
        if shutil.disk_usage("/content").free > faltam + folga:
            break
        os.remove(caminho)


def _tamanho_previsto(url):
    """Tamanho do arquivo no Hugging Face (cabecalho da resposta), sem baixar. 0 se nao der para saber."""
    try:
        resposta = urllib.request.urlopen(urllib.request.Request(url, method="HEAD"), timeout=15)
        return int(resposta.headers.get("Content-Length") or resposta.headers.get("X-Linked-Size") or 0)
    except (OSError, ValueError):
        return 0


def abrir_tunel(porta, chave):
    """Abre um tunel Cloudflare para a porta e guarda o link em ui[chave]."""
    proc = subprocess.Popen(
        [CLOUDFLARED, "tunnel", "--url", f"http://127.0.0.1:{porta}", "--no-autoupdate", "--protocol", "http2"],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, errors="ignore",
    )

    def ler():
        # continua lendo depois de achar o link, para o cloudflared nao travar
        for linha in proc.stdout:
            m = re.search(r"https://[a-z0-9-]+[.]trycloudflare[.]com", linha)
            if m and not ui[chave]:
                ui[chave] = m.group(0)

    threading.Thread(target=ler, daemon=True).start()
    return proc


# ------------------------------------------------------------------------------
# EXECUCAO
# ------------------------------------------------------------------------------

if SALVAR_NO_DRIVE:
    from google.colab import drive
    drive.mount("/content/drive")

# para execucoes antigas (caso rode a celula de novo)
subprocess.run(["pkill", "-f", "main.py --listen"])
subprocess.run(["pkill", "-f", "cloudflared tunnel"])
subprocess.run(["pkill", "-f", "voz_server.py"])
subprocess.run(["pkill", "-f", "fundo_server.py"])
subprocess.run(["pkill", "-f", "open-webui serve"])
subprocess.run(["pkill", "-f", "ollama serve"])

# fecha a API de uma execucao anterior, para liberar a porta
if globals().get("_api_server"):
    try:
        _api_server.shutdown()
        _api_server.server_close()
    except Exception:
        pass
_api_server = None

# saida limpa (some o "Mounted at ..." do Drive) e sem limite de altura: o Colab poe uma rolagem interna em saidas altas,
# e sem ela, junto com a margem grande embaixo do painel, da para rolar ate o painel ocupar a tela inteira
clear_output()
try:
    from google.colab import output as saida_colab
    saida_colab.no_vertical_scroll()
except Exception:
    pass
tela = display(HTML(render()), display_id=True)
comfy = None
tunel = None
tunel_api = None
voz_proc = None
ollama_proc = None
fundo_proc = None
webui_proc = None
tunel_webui = None

try:
    # 1. GPU
    comecar("gpu", "nvidia-smi")
    try:
        smi = subprocess.run(
            ["nvidia-smi", "--query-gpu=name,memory.total", "--format=csv,noheader"],
            capture_output=True, text=True,
        )
        nome_gpu = smi.stdout.strip() if smi.returncode == 0 else ""
    except OSError:
        nome_gpu = ""
    PESO_FP8 = "T4" in nome_gpu or "P100" in nome_gpu or "V100" in nome_gpu   # placas sem bf16
    if not nome_gpu:
        falhar("gpu", "Sem GPU. Vá em Ambiente de execução > Alterar o tipo de ambiente de execução > GPU T4 e rode de novo.")
    concluir("gpu", nome_gpu)

    # 2. ComfyUI (so nos packs que geram imagem, video, musica ou efeitos)
    liberar_disco()
    if not USA_COMFY:
        pass
    elif os.path.isdir(os.path.join(COMFY, ".git")):
        comecar("comfy", "git pull (já instalado, procurando atualizações)")
        rodar("comfy", ["git", "-C", COMFY, "pull", "-q"], obrigatorio=False)
        concluir("comfy", "atualizado")
    else:
        comecar("comfy", "git clone https://github.com/comfyanonymous/ComfyUI")
        # uma sessao anterior sem ComfyUI (so chat ou voz) deixa a pasta so com a galeria: ela e guardada e volta depois do clone
        galeria_antiga, guardada = os.path.join(COMFY, "output"), "/content/_galeria_antiga"
        if os.path.lexists(galeria_antiga):
            shutil.move(galeria_antiga, guardada)
        shutil.rmtree(COMFY, ignore_errors=True)
        rodar("comfy", ["git", "clone", "--depth", "1", "--progress",
                        "https://github.com/comfyanonymous/ComfyUI.git", COMFY], com_pct=True)
        if os.path.lexists(guardada):
            shutil.rmtree(galeria_antiga, ignore_errors=True)
            shutil.move(guardada, galeria_antiga)
        concluir("comfy")
    os.makedirs(COMFY, exist_ok=True)   # sem ComfyUI, a pasta existe so para a galeria (recortes, vozes geradas)

    if SALVAR_NO_DRIVE:
        pasta_drive = "/content/drive/MyDrive/ComfyUI_output"
        pasta_saida = os.path.join(COMFY, "output")
        os.makedirs(pasta_drive, exist_ok=True)
        if not os.path.islink(pasta_saida):
            if os.path.isdir(pasta_saida):
                for nome in os.listdir(pasta_saida):
                    if not os.path.exists(os.path.join(pasta_drive, nome)):
                        shutil.move(os.path.join(pasta_saida, nome), pasta_drive)
                shutil.rmtree(pasta_saida)
            os.symlink(pasta_drive, pasta_saida)

    # 3. Dependencias
    if USA_COMFY:
        comecar("deps", "pip install -r requirements.txt")
        rodar("deps", [sys.executable, "-m", "pip", "install", "--progress-bar", "off", "-r", "requirements.txt"], cwd=COMFY)
        concluir("deps")

    # 4. Cloudflared
    comecar("tunel", "conectando ao GitHub...")
    if os.path.isfile(CLOUDFLARED):
        concluir("tunel", "já estava baixado")
    else:
        baixar("tunel", CLOUDFLARED_URL, CLOUDFLARED)
        concluir("tunel")
    os.chmod(CLOUDFLARED, 0o755)

    # 5. Modelos
    for i, (pasta, arquivo, url_modelo, _, _) in enumerate(MODELOS):
        k = f"m{i}"
        comecar(k, "conectando ao Hugging Face...")
        os.makedirs(os.path.join(COMFY, "models", pasta), exist_ok=True)
        destino = os.path.join(COMFY, "models", pasta, arquivo)
        if os.path.isfile(destino):
            concluir(k, "já estava baixado")
        else:
            baixar(k, url_modelo, destino)
            concluir(k, fmt_tam(os.path.getsize(destino)))

    # Nodes extras do perfil (precisam estar na pasta antes de o ComfyUI iniciar)
    if PERFIL_ATUAL.get("nodes"):
        comecar("nodes")
        for pasta_node, repo in PERFIL_ATUAL["nodes"]:
            destino = os.path.join(COMFY, "custom_nodes", pasta_node)
            if os.path.isdir(destino):
                continue
            estado["nodes"]["info"] = f"git clone {repo}"
            rodar("nodes", ["git", "clone", "--depth", "1", "--progress", repo, destino], com_pct=True)
            requisitos = os.path.join(destino, "requirements.txt")
            if os.path.isfile(requisitos):
                estado["nodes"].update(info="pip install -r requirements.txt", pct=None)
                rodar("nodes", [sys.executable, "-m", "pip", "install", "-q", "-r", requisitos])
        concluir("nodes", ", ".join(n for n, _ in PERFIL_ATUAL["nodes"]))

    # Voz: ambiente proprio + modelo
    if VOZ:
        comecar("voz_inst", "uv venv (Python 3.11) + chatterbox-tts")
        marca_voz = os.path.join(VOZ_VENV, ".instalado")
        if os.path.isfile(marca_voz):
            concluir("voz_inst", "já estava instalado")
        else:
            rodar("voz_inst", [sys.executable, "-m", "pip", "install", "-q", "uv"])
            rodar("voz_inst", [sys.executable, "-m", "uv", "venv", "--clear", VOZ_VENV, "--python", "3.11"])
            estado["voz_inst"]["info"] = "instalando torch 2.6 + chatterbox (demora uns minutos)..."
            rodar("voz_inst", [sys.executable, "-m", "uv", "pip", "install", "--python", VOZ_PY, CHATTERBOX_PIP])
            open(marca_voz, "w").close()
            concluir("voz_inst")

        comecar("voz_modelo", "conectando ao Hugging Face...")
        os.makedirs(VOZ_MODELO, exist_ok=True)
        baixou = False
        for url_voz, nome_voz in VOZ_ARQUIVOS:
            destino = os.path.join(VOZ_MODELO, nome_voz)
            if not os.path.isfile(destino):
                baixar("voz_modelo", url_voz, destino)
                baixou = True
        concluir("voz_modelo", "3.2 GB" if baixou else "já estava baixado")

    # IA de chat local (opcional: se falhar, o resto continua e so o chat fica desligado)
    if CHAT_LOCAL:
        comecar("chat_inst", "curl https://ollama.com/install.sh | sh")
        if not shutil.which("ollama"):
            rodar("chat_inst", ["bash", "-c", "apt-get install -y -qq zstd > /dev/null 2>&1; curl -fsSL https://ollama.com/install.sh | sh"], obrigatorio=False)
        chat_ok = bool(shutil.which("ollama"))
        if chat_ok:
            estado["chat_inst"]["info"] = "ollama serve"
            ollama_proc = subprocess.Popen(
                ["ollama", "serve"],
                env={**os.environ, "OLLAMA_KEEP_ALIVE": "24h" if CHAT_MODO == "fixa" else "10m", "OLLAMA_HOST": f"127.0.0.1:{OLLAMA_PORT}"},
                stdout=open(OLLAMA_LOG, "w"), stderr=subprocess.STDOUT,
            )
            inicio = time.time()
            while not porta_aberta(OLLAMA_PORT) and ollama_proc.poll() is None and time.time() - inicio < 60:
                pintar()
                time.sleep(0.3)
            chat_ok = porta_aberta(OLLAMA_PORT)
        if chat_ok:
            concluir("chat_inst")
            comecar("chat_modelo", f"ollama pull {CHAT_LOCAL_MODELO}")
            rodar("chat_modelo", ["ollama", "pull", CHAT_LOCAL_MODELO], com_pct=True, obrigatorio=False)
            lista = subprocess.run(["ollama", "list"], capture_output=True, text=True, env={**os.environ, "OLLAMA_HOST": f"127.0.0.1:{OLLAMA_PORT}"}).stdout
            chat_ok = CHAT_LOCAL_MODELO in lista
            if chat_ok and CHAT_LOCAL_CPU:
                # cria a copia que roda so na CPU; se o Ollama recusar, o modelo original segue na GPU
                with open("/content/Modelfile.chat", "w") as f:
                    f.write(f"FROM {CHAT_LOCAL_MODELO}\nPARAMETER num_gpu 0\n")
                estado["chat_modelo"].update(info="configurando para rodar na CPU", pct=None)
                rodar("chat_modelo", ["ollama", "create", CHAT_CPU_NOME, "-f", "/content/Modelfile.chat"], obrigatorio=False)
                lista = subprocess.run(["ollama", "list"], capture_output=True, text=True, env={**os.environ, "OLLAMA_HOST": f"127.0.0.1:{OLLAMA_PORT}"}).stdout
                if CHAT_CPU_NOME in lista:
                    CHAT_MODELO = CHAT_CPU_NOME
                else:
                    CHAT_LOCAL_CPU, CHAT_MODO = False, "reveza"
            if chat_ok and CHAT_MODO == "fixa":
                # ja coloca o modelo na placa agora: sem isso, a primeira mensagem esperava 20-60 s em silencio
                # (o Open WebUI nao mostra nada enquanto o Ollama carrega, e parecia que a mensagem tinha sido ignorada)
                estado["chat_modelo"].update(info="carregando o modelo na placa...", pct=None)
                pintar(True)
                with contextlib.suppress(OSError, ValueError):
                    pedido = json.dumps({"model": CHAT_MODELO, "prompt": "", "keep_alive": "24h"}).encode()
                    urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{OLLAMA_PORT}/api/generate", data=pedido,
                                                                  headers={"Content-Type": "application/json"}), timeout=300).read()
            if chat_ok:
                concluir("chat_modelo", CHAT_LOCAL_TAMANHO + {"cpu": " · CPU", "fixa": " · GPU fixa"}.get(CHAT_MODO, " · GPU"))
            else:
                avisar("chat_modelo", "não baixou; o chat fica desligado")
        else:
            avisar("chat_inst", "não instalou; o chat fica desligado")
            avisar("chat_modelo", "pulado")
        if not chat_ok:
            CHAT_LIGADO = False

    # Criador de arquivos (opcional: se falhar, so essa ferramenta some)
    if ARQUIVOS:
        comecar("arquivos", "pip install openpyxl python-docx python-pptx fpdf2")
        import importlib.metadata
        with open("/content/restricoes_pip.txt", "w") as f:
            for pacote in ("pillow", "numpy", "torch", "torchvision", "torchaudio", "lxml"):
                with contextlib.suppress(importlib.metadata.PackageNotFoundError):
                    f.write(f"{pacote}=={importlib.metadata.version(pacote)}\n")
        rodar("arquivos", [sys.executable, "-m", "pip", "install", "-q", "-c", "/content/restricoes_pip.txt",
                           "openpyxl", "python-docx", "python-pptx", "fpdf2"], obrigatorio=False)
        try:
            import openpyxl, docx, pptx, fpdf   # noqa: F401  (so confere se instalou)
            concluir("arquivos", "xlsx · docx · pdf · pptx · csv · md")
        except ImportError as ex:
            ARQUIVOS = False
            MCP_TOOLS[:] = [t for t in MCP_TOOLS if t["name"] != "create_file"]
            CHAT_TOOLS[:] = [t for t in CHAT_TOOLS if t["function"]["name"] != "criar_arquivo"]
            avisar("arquivos", f"não instalou ({ex.name}); a criação de arquivos fica desligada")

    # Open WebUI: chat completo, ligado na IA local (opcional: se falhar, o chat do Studio continua)
    if WEBUI and CHAT_LIGADO:
        comecar("webui", "uv venv (Python 3.11) + open-webui")
        marca_webui = os.path.join(WEBUI_VENV, ".instalado")
        if not os.path.isfile(marca_webui):
            rodar("webui", [sys.executable, "-m", "pip", "install", "-q", "uv"], obrigatorio=False)
            rodar("webui", [sys.executable, "-m", "uv", "venv", "--clear", WEBUI_VENV, "--python", "3.11"], obrigatorio=False)
            estado["webui"]["info"] = "instalando o Open WebUI (uns 7 GB, cerca de 1 min e meio)..."
            rodar("webui", [sys.executable, "-m", "uv", "pip", "install", "--python", WEBUI_VENV + "/bin/python", "open-webui"], obrigatorio=False)
            if os.path.isfile(WEBUI_VENV + "/bin/open-webui"):
                open(marca_webui, "w").close()
        if os.path.isfile(marca_webui):
            # dados zerados a cada execucao: o usuario administrador e criado com a chave desta sessao, que muda sempre
            shutil.rmtree(WEBUI_DADOS, ignore_errors=True)
            estado["webui"]["info"] = "iniciando..."
            webui_proc = subprocess.Popen(
                [WEBUI_VENV + "/bin/open-webui", "serve", "--host", "127.0.0.1", "--port", str(WEBUI_PORT)],
                env=ambiente_webui(),
                stdout=open(WEBUI_LOG, "w"), stderr=subprocess.STDOUT,
            )
            inicio = time.time()
            while webui_proc.poll() is None and time.time() - inicio < 240:
                try:
                    if urllib.request.urlopen(f"http://127.0.0.1:{WEBUI_PORT}/health", timeout=3).status == 200:
                        webui_ligado = True
                        break
                except OSError:
                    pass
                pintar()
                time.sleep(1)
            token_webui = None
            if webui_ligado and CHAVE_API:
                # o primeiro cadastro vira o administrador; depois dele, o proprio Open WebUI recusa cadastros novos
                try:
                    cadastro = json.dumps({"name": "Studio", "email": WEBUI_LOGIN, "password": CHAVE_API}).encode()
                    resposta = urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{WEBUI_PORT}/api/v1/auths/signup", data=cadastro,
                                                                             headers={"Content-Type": "application/json"}), timeout=60).read()
                    with contextlib.suppress(ValueError):
                        token_webui = json.loads(resposta).get("token")
                except OSError:
                    webui_ligado = False
            if webui_ligado:
                estado["webui"]["info"] = "instalando skills e atalhos..."
                pintar(True)
                skills_ok, atalhos_ok = preparar_webui(token_webui)
        if webui_ligado:
            concluir("webui", ("login: " + WEBUI_LOGIN if CHAVE_API else "sem login") + f" · {skills_ok} skills · {atalhos_ok} atalhos · pesquisa na web")
        else:
            if webui_proc is not None and webui_proc.poll() is None:
                webui_proc.terminate()
            avisar("webui", "não subiu; use o chat do Studio")

    # Remover fundo: ambiente proprio + servidor (opcional: se falhar, so essa funcao fica desligada)
    if REMOVER_FUNDO:
        comecar("fundo", "uv venv (Python 3.11) + rembg")
        marca_fundo = os.path.join(FUNDO_VENV, ".instalado")
        if not os.path.isfile(marca_fundo):
            rodar("fundo", [sys.executable, "-m", "pip", "install", "-q", "uv"], obrigatorio=False)
            rodar("fundo", [sys.executable, "-m", "uv", "venv", "--clear", FUNDO_VENV, "--python", "3.11"], obrigatorio=False)
            rodar("fundo", [sys.executable, "-m", "uv", "pip", "install", "--python", FUNDO_PY, "rembg", "onnxruntime"], obrigatorio=False)
            if os.path.isfile(FUNDO_PY) and subprocess.run([FUNDO_PY, "-c", "import rembg"], capture_output=True).returncode == 0:
                open(marca_fundo, "w").close()
        if os.path.isfile(marca_fundo):
            with open("/content/fundo_server.py", "w", encoding="utf-8") as f:
                f.write(FUNDO_SERVER_PY)
            estado["fundo"]["info"] = f"iniciando (na primeira vez baixa o modelo {MODELO_FUNDO}, 170 MB)..."
            fundo_proc = subprocess.Popen(
                [FUNDO_PY, "-u", "/content/fundo_server.py", "--porta", str(FUNDO_PORT), "--modelo", MODELO_FUNDO],
                stdout=open(FUNDO_LOG, "w"), stderr=subprocess.STDOUT,
            )
            inicio = time.time()
            while not porta_aberta(FUNDO_PORT) and fundo_proc.poll() is None and time.time() - inicio < 300:
                pintar()
                time.sleep(0.3)
            fundo_ligado = porta_aberta(FUNDO_PORT)
        if fundo_ligado:
            concluir("fundo", "CPU")
        else:
            avisar("fundo", "não instalou; o recorte fica desligado")

    # Ouvir: Whisper na CPU (opcional: se falhar, so o microfone fica desligado)
    if OUVIR:
        comecar("ouvir", "pip install faster-whisper")
        # o Pillow, o numpy e o torch ja estao em uso nesta sessao: se o pip trocar a versao no disco, o que esta
        # na memoria deixa de combinar com o que e importado depois (foi assim que o GIF da animacao quebrou)
        import importlib.metadata
        with open("/content/restricoes_pip.txt", "w") as f:
            for pacote in ("pillow", "numpy", "torch", "torchvision", "torchaudio"):
                try:
                    f.write(f"{pacote}=={importlib.metadata.version(pacote)}\n")
                except importlib.metadata.PackageNotFoundError:
                    pass
        rodar("ouvir", [sys.executable, "-m", "pip", "install", "-q", "-c", "/content/restricoes_pip.txt", "faster-whisper"], obrigatorio=False)
        try:
            estado["ouvir"]["info"] = f"baixando o modelo {WHISPER}..."
            pintar(True)
            from faster_whisper import WhisperModel
            whisper_modelo = WhisperModel(WHISPER_REPOS.get(WHISPER, WHISPER), device="cpu", compute_type="int8")
            concluir("ouvir", "CPU")
        except Exception as ex:
            whisper_modelo = None
            avisar("ouvir", f"não instalou ({type(ex).__name__}); o microfone fica desligado")

    # 6. Iniciar ComfyUI
    if USA_COMFY:
        comecar("start", "python main.py --listen 0.0.0.0 --port 8188")
        comfy = subprocess.Popen(
            [sys.executable, "-u", "main.py", "--listen", "0.0.0.0", "--port", str(PORT), "--preview-method", "auto"],
            cwd=COMFY, stdout=open(LOG, "w"), stderr=subprocess.STDOUT,
        )
        inicio = time.time()
        while not porta_aberta():
            if comfy.poll() is not None or time.time() - inicio > 300:
                falhar("start", "O ComfyUI não conseguiu iniciar. Veja o motivo abaixo.", ultimas_linhas(25))
            ultima = ultimas_linhas(1)
            if ultima:
                estado["start"]["info"] = ultima
            pintar()
            time.sleep(0.2)
        concluir("start")
    else:
        with contextlib.suppress(OSError):
            os.remove(LOG)   # sem ComfyUI, o painel nao deve mostrar o log de uma sessao anterior

    # Voz: servidor (o modelo so vai para a GPU na primeira fala)
    if VOZ:
        comecar("voz_start", f"python voz_server.py --porta {VOZ_PORT}")
        os.makedirs(VOZ_REFS, exist_ok=True)
        with open("/content/voz_server.py", "w", encoding="utf-8") as f:
            f.write(VOZ_SERVER_PY)
        voz_proc = subprocess.Popen(
            [VOZ_PY, "-u", "/content/voz_server.py", "--porta", str(VOZ_PORT), "--modelo", VOZ_MODELO],
            stdout=open(VOZ_LOG, "w"), stderr=subprocess.STDOUT,
        )
        inicio = time.time()
        while not porta_aberta(VOZ_PORT):
            if voz_proc.poll() is not None or time.time() - inicio > 240:
                falhar("voz_start", "O servidor de voz não conseguiu iniciar. Veja o motivo abaixo.", ultimas_linhas(25, VOZ_LOG))
            estado["voz_start"]["info"] = ultimas_linhas(1, VOZ_LOG) or "carregando o torch..."
            pintar()
            time.sleep(0.3)
        concluir("voz_start")

    # 7. API + MCP
    comecar("api", f"http.server 127.0.0.1:{API_PORT}  /gerar  /mcp")
    try:
        _api_server = ThreadingHTTPServer(("127.0.0.1", API_PORT), ApiHandler)
    except OSError as ex:
        falhar("api", f"Não consegui abrir a porta {API_PORT} da API: {ex}")
    threading.Thread(target=_api_server.serve_forever, daemon=True).start()
    concluir("api", "/gerar · /mcp")

    # 8. Links (um tunel para o ComfyUI, outro para a API)
    comecar("link", "cloudflared tunnel --protocol http2 (" + ("ComfyUI + " if USA_COMFY else "") + "API)")
    if USA_COMFY:
        tunel = abrir_tunel(PORT, "link")
    tunel_api = abrir_tunel(API_PORT, "link_api")
    if webui_ligado:
        tunel_webui = abrir_tunel(WEBUI_PORT, "link_webui")
    inicio = time.time()
    while not ((ui["link"] or not USA_COMFY) and ui["link_api"] and (ui["link_webui"] or not webui_ligado)) and time.time() - inicio < 60:
        pintar()
        time.sleep(0.2)

    # link do proprio Colab, como alternativa
    try:
        from google.colab.output import eval_js
        ui["link_colab"] = eval_js(f"google.colab.kernel.proxyPort({PORT if USA_COMFY else API_PORT})")
    except Exception:
        pass

    if not ui["link"] and not ui["link_api"] and not ui["link_colab"]:
        falhar("link", "Não consegui criar nenhum link. Rode a célula de novo.")
    concluir("link", "" if (ui["link"] or not USA_COMFY) and ui["link_api"] else "o Cloudflare não devolveu todos os links")

    ui["t_fim"] = time.time()
    ui["online"] = time.time()

    # fica rodando: o log do ComfyUI aparece ao vivo no painel
    ultima_tela = None
    while comfy is None or comfy.poll() is None:   # sem ComfyUI, fica no ar ate a celula ser parada
        tela_agora = (ultimas_linhas(), len(ui["api_log"]), ui["api_log"][-1:], int((time.time() - ui["online"]) // 60), metricas())
        if tela_agora != ultima_tela:
            ultima_tela = tela_agora
            pintar(True)
        time.sleep(0.5)
    ui["erro"] = "O ComfyUI parou sozinho (provavelmente faltou memória). Rode a célula de novo."
    ui["erro_log"] = ultimas_linhas(25)

except Falha:
    pass
except KeyboardInterrupt:
    ui["parado"] = True
    for e in estado.values():
        if e["status"] == "rodando":
            e.update(status="espera", info="", pct=None, detalhe="")
finally:
    ui["t_fim"] = ui["t_fim"] or time.time()
    for proc in (comfy, tunel, tunel_api, voz_proc, ollama_proc, fundo_proc, webui_proc, tunel_webui):
        if proc is not None and proc.poll() is None:
            proc.terminate()
    if _api_server is not None:
        _api_server.shutdown()
        _api_server.server_close()
        _api_server = None
    pintar(True)
